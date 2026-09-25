import openpyxl, json, re, html, base64, os

print("=== Compiling Master ODL Second Shift Monitoring & Teaching Loads System ===")

REPO_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
DATA_FILE = os.path.join(REPO_DIR, 'data', 'schedule_odl_2nd_shift.xlsx')
AMIS_LOGO = os.path.join(REPO_DIR, 'public', 'amis_logo_opt.png')
DEPED_LOGO = os.path.join(REPO_DIR, 'public', 'deped_logo_opt.png')
OUTPUT_FILE = os.path.join(REPO_DIR, 'public', 'odl-second-shift.html')
ROOT_OUTPUT_FILE = os.path.join(REPO_DIR, '..', 'odl-second-shift-monitoring-sheet.html')

wb = openpyxl.load_workbook(DATA_FILE, data_only=True)

with open(AMIS_LOGO, 'rb') as f:
    amis_b64 = "data:image/png;base64," + base64.b64encode(f.read()).decode('utf-8')

with open(DEPED_LOGO, 'rb') as f:
    deped_b64 = "data:image/png;base64," + base64.b64encode(f.read()).decode('utf-8')

days = ['Sunday', 'Monday', 'Tuesday', 'Wednesday', 'Thursday']
day_abbr_map = {
    'Sunday': 'SUN',
    'Monday': 'MON',
    'Tuesday': 'TUE',
    'Wednesday': 'WED',
    'Thursday': 'THU',
    'Friday': 'FRI',
    'Saturday': 'SAT'
}
day_order = {'Sunday': 0, 'Monday': 1, 'Tuesday': 2, 'Wednesday': 3, 'Thursday': 4}

def get_grid(sheet):
    grid = {}
    for r in range(1, sheet.max_row + 1):
        for c in range(1, sheet.max_column + 1):
            val = sheet.cell(r, c).value
            grid[(r, c)] = str(val).strip() if val is not None else ''
            
    for rng in sheet.merged_cells.ranges:
        top_val = grid.get((rng.min_row, rng.min_col), '')
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                grid[(r, c)] = top_val
    return grid

elem_grid = get_grid(wb['ELEM'])
hs_new_grid = get_grid(wb['HS SCHED (NEW)'])
hs_grid = get_grid(wb['HS SCHED'])

def clean_time(t_str):
    if not t_str: return ''
    t = str(t_str).strip()
    if any(kw in t.upper() for kw in ['GRADE', 'USTADH', 'TEACHER', 'LEVEL', 'ADVISER']): return ''
    m_odl = re.search(r'([\d:apm\.\s-]+?)\s*(?:\(ODL\)|ODL)', t, re.IGNORECASE)
    if m_odl: t = m_odl.group(1).strip()
    else: t = re.sub(r'\(.*?\)', '', t).strip()
    t = re.sub(r'(\d{1,2}:\d{2}):+(\d{1,2}:\d{2})', r'\1 - \2', t)
    t = re.sub(r'(\d{1,2}:\d{2}):00', r'\1', t)
    is_pm = bool(re.search(r'(?:p\.?m\.?|pm)', t, re.IGNORECASE))
    is_am = bool(re.search(r'(?:a\.?m\.?|am)', t, re.IGNORECASE))
    clean = re.sub(r'(?:a\.?m\.?|p\.?m\.?|am|pm)', '', t, flags=re.IGNORECASE).strip()
    parts = re.split(r'\s*[-–]\s*', clean)
    if len(parts) == 2:
        start, end = parts[0].strip(), parts[1].strip()
        start = re.sub(r'^0(\d:)', r'\1', start)
        end = re.sub(r'^0(\d:)', r'\1', end)
        try:
            start_hour = int(start.split(':')[0]) if ':' in start else 0
            end_hour = int(end.split(':')[0]) if ':' in end else 0
        except ValueError:
            return t
        if is_pm:
            if start_hour == 11 and end_hour == 12: return f'{start} AM - {end} PM'
            return f'{start} - {end} PM'
        elif is_am: return f'{start} - {end} AM'
        else:
            if start_hour in [7, 8, 9, 10, 11] and end_hour in [7, 8, 9, 10, 11]: return f'{start} - {end} AM'
            elif start_hour == 11 and (end_hour == 12 or end_hour <= 1): return f'{start} AM - {end} PM'
            elif start_hour == 12 or start_hour in [1, 2, 3, 4, 5, 6]: return f'{start} - {end} PM'
            else: return f'{start} - {end} PM'
    elif len(parts) == 1:
        val = parts[0].strip()
        val = re.sub(r'^0(\d:)', r'\1', val)
        if is_pm or val.startswith(('12:', '1:', '2:', '3:', '4:', '5:')): return f'{val} PM'
        else: return f'{val} AM'
    return t

def parse_mins(m_str):
    if not m_str: return '40'
    m_str = str(m_str).strip()
    m = re.search(r'(\d+)', m_str)
    return str(m.group(1)) if m else '40'

KNOWN_TEACHERS = {
    'Saimona': 'Tchr. Saimona', 'Jerlyn': 'Tchr. Jerlyn', 'Keychell': 'Tchr. Keychell',
    'Faidh': 'Ustadh Faidh', 'Saliha': 'Ustadha Saliha', 'Obaydah': 'Ust. Obaydah',
    'Obayda': 'Ust. Obaydah', 'Marham': 'Tchr. Marham', 'Jenny': 'Tchr. Jenny',
    'Jessa': 'Tchr. Jessa', 'Laurel': 'Ust. Laurel', 'Laiurel': 'Ust. Laurel',
    'Silfah': 'Ustadha Silfa', 'Silfa': 'Ustadha Silfa', 'Ersahad': 'Ustadh Ersahad',
    'Hainur': 'Ustadh Hainur', 'Bustamante': 'Alim Bustamante', 'Wendy': 'Tchr. Wendy',
    'Sahdia': 'Tchr. Sahdia', 'Norhydie': 'Tchr. Norhydie', 'Joanna': 'Tchr. Joanna',
    'Katrina': 'Tchr. Katrina', 'Sitti': 'Tchr. Sitti', 'Monisa': 'Tchr. Monisa',
    'Zara': 'Tchr. Franchette', 'Anna': 'Tchr. Anna', 'Arvin': 'Tchr. Arvin',
    'Fhairudz': 'Tchr. Fhairudz', 'Erica': 'Tchr. Erica', 'Junaisah': 'Tchr. Junaisah',
    'Junaisa': 'Tchr. Junaisah', 'Ayah': 'Tchr. Ayah', 'Jayra': 'Tchr. Jayra',
    'Jairah': 'Tchr. Jayra', 'Shi': 'Tchr. Shirehan', 'Shirehan': 'Tchr. Shirehan',
    'Sophia': 'Tchr. Sophia', 'Aniah': 'Tchr. Aniah', 'Halnaisa': 'Tchr. Halnaisa',
    'Franchette': 'Tchr. Franchette', 'Ethel': 'Tchr. Ethel', 'Ali': 'Ustadh Muh Ali',
    'Radzmia': 'Tchr. Radzmia', 'Wardah': 'Tchr. Wardah', 'Jaisam': 'Ustadh Jaisam',
    'Rowena': 'Tchr. Rowena', 'Norhaima': 'Tchr. Norhaima', 'Jhelyn': 'Tchr. Jhelyn',
    'Nof': 'Tchr. Nof', 'Angeleni': 'Tchr. Angeleni', 'Nadzra': 'Tchr. Nadzra',
    'Samsuddin': 'Alim Samsuddin', 'Abdulwahab': 'Alim Abdulwahab', 'Mamonas': 'Alim Mamonas',
    'Moh': 'Sir Moh', 'Mohaymen': 'Sir Moh', 'Thea': 'Tchr. Thea', 'Abegail': 'Tchr. Abegail',
    'Marie': 'Tchr. Marie', 'Ahmad': 'Tchr. Ahmad', 'Dipatuan': 'Alim Dipatuan',
    'Abdul Karim': 'Alim Abdul Karim', 'Abdiraheem': 'Ust. Abdiraheem', 'Abdi': 'Ust. Abdiraheem',
}

def normalize_teacher(name):
    name = re.sub(r'\s+', ' ', name.strip()).strip('-– ')
    if re.match(r'^Teacher\b', name, re.IGNORECASE): name = 'Tchr. ' + name[7:].strip()
    elif re.match(r'^Tr\.?\b', name, re.IGNORECASE): name = 'Tchr. ' + name[3:].strip()
    elif re.match(r'^Tchr\.?\b', name, re.IGNORECASE): name = 'Tchr. ' + re.sub(r'^Tchr\.?\s*', '', name, flags=re.IGNORECASE)
    elif re.match(r'^(?:Ustadha|Ustadza)\b', name, re.IGNORECASE): name = 'Ustadha ' + re.sub(r'^(?:Ustadha|Ustadza)\s*', '', name, flags=re.IGNORECASE)
    elif re.match(r'^Ustadh\b', name, re.IGNORECASE): name = 'Ustadh ' + name[6:].strip()
    elif re.match(r'^Ust\.?\b', name, re.IGNORECASE): name = 'Ust. ' + re.sub(r'^Ust\.?\s*', '', name, flags=re.IGNORECASE)
    name = name.strip()
    if name in ['Tchr.', 'Tchr', 'Ust.', 'Ust', 'Teacher', 'Tr.', 'Tr', '']: return 'TBA'
    if name.upper() in ['TCHR. AHMAD', 'SIR AHMAD']: return 'Tchr. Ahmad'
    if name in ['Ustadh Jaisam', 'Ust. Jaisam']: return 'Ustadh Jaisam'
    if name in ['Tchr. Kat', 'Tchr. Katrina']: return 'Tchr. Katrina'
    if name in ['Ust. Ubaydah', 'Ust. Obaydah', 'Ustadh Obayda', 'Ust. Obayda', 'Obaydah']: return 'Ust. Obaydah'
    if name in ['Ust. Silfah', 'Ustadha Silfa', 'Ust. Silfa', 'Ustadza Silfah', 'Ust. Silfah']: return 'Ustadha Silfa'
    if name in ['Ustadha Saliha', 'Ust. Saliha', 'Saliha']: return 'Ustadha Saliha'
    if name in ['Tchr. Junaisa', 'Tchr. Junaisah']: return 'Tchr. Junaisah'
    if name in ['Tchr. Jairah', 'Tchr. Jayra']: return 'Tchr. Jayra'
    if name in ['Tchr. Moh', 'Sir Mohaymen', 'Sir Moh']: return 'Sir Moh'
    if name in ['Tchr. Shi', 'Tchr. Shirehan']: return 'Tchr. Shirehan'
    if name in ['Tchr. Zara', 'Tchr. Franchette', 'Franchette']: return 'Tchr. Franchette'
    if name in ['Ust. Abdi', 'Ust. Abdiraheem', 'Ustadh Abdi', 'Ustadh Abdiraheem', 'Ustd. Abdi', 'Ustd. Abdiraheem']: return 'Ust. Abdiraheem'
    if name in ['Ust. Ali', 'Ustadh Ali', 'Ustadh Muh Ali', 'Ustdh ali', 'Ustdh. Ali', 'Ust. Muh Ali', 'Ustadh Muh. Ali']: return 'Ustadh Muh Ali'
    if name in ['Ust. Faidh', 'Ustadh Faidh', 'Faidh', 'Us. Faidh']: return 'Ustadh Faidh'
    if name in ['Ust. Ersahad', 'Ustadh Ersahad', 'Ersahad']: return 'Ustadh Ersahad'
    if name in ['Ust. Hainur', 'Ustadh Hainur', 'Hainur']: return 'Ustadh Hainur'
    if name in ['Alim Abdul Karim', 'Ust. Abdul Karim', 'Abdul Karim']: return 'Alim Abdul Karim'
    if name in ['Alim Samsuddin', 'Ust. Samsuddin', 'Samsuddin']: return 'Alim Samsuddin'
    if name in ['Alim Abdulwahab', 'Ust. Abdulwahab']: return 'Alim Abdulwahab'
    if name in ['Alim Mamonas', 'Ust. Mamonas']: return 'Alim Mamonas'
    if name in ['Alim Bustamante', 'Ust. Bustamante']: return 'Alim Bustamante'
    if name in ['Laiurel', 'Ust. Laurel', 'Laurel']: return 'Ust. Laurel'
    if name in ['Tr Nof', 'Tchr. Nof', 'Nof']: return 'Tchr. Nof'
    return name

def clean_parse(text):
    text = text.strip()
    if not text: return '', '', 'EMPTY'
    upper = text.upper()
    for r in ['GENERAL ASSEMBLY', 'LUNCH AND SALAH', 'SALAH & DEPARTURE', 'DEPARTURE', 'RECESS', 'SHORT BREAK', 'TRANSITION', 'BREAK', 'HOMEROOM', 'ENTRANCE EXAM REVIEW', 'RESEARCH CONSULTATION']:
        if upper == r or (upper.startswith(r) and not any(k in upper for k in ['TCHR', 'UST', 'SIR', 'ALIM', 'TEACHER', 'TR'])):
            return text, '', 'ROUTINE'
    t_match = re.search(r'[-–]?\s*(Tchr\.?|Ust\.?|Sir|Alim|Teacher|Ustadha|Ustadza|Ustadh|Tr\.?)\s+(.*)$', text, re.IGNORECASE)
    if t_match:
        tchr_full = normalize_teacher((t_match.group(1) + ' ' + t_match.group(2)).strip())
        subj_part = text[:t_match.start()].strip().rstrip('-–').strip()
        return subj_part or text, tchr_full, 'CLASS'
    for k, v in KNOWN_TEACHERS.items():
        m = re.search(r'[-–]?\s*' + re.escape(k) + r'\b', text, re.IGNORECASE)
        if m:
            subj = text[:m.start()].strip().rstrip('-–').strip()
            subj = re.sub(r'([A-Za-z]+)(\d+)$', r'\1 \2', subj)
            return subj or text, v, 'CLASS'
    for p in ['Tchr.', 'Teacher', 'Sir', 'Ustadh', 'Ustadha', 'Ustadza', 'Ust.', 'Alim', 'Tr.']:
        if p.lower() in text.lower():
            idx = text.lower().find(p.lower())
            tchr_full = normalize_teacher(text[idx:].strip())
            subj_part = text[:idx].strip().rstrip('-–').strip()
            return subj_part or text, tchr_full, 'CLASS'
    return text, 'TBA', 'CLASS'

def time_to_sort_key(t_str):
    m = re.search(r'(\d+):(\d+)', t_str)
    if not m: return 9999
    h, mins = int(m.group(1)), int(m.group(2))
    if 'PM' in t_str.upper():
        if h < 12: h += 12
    elif 'AM' in t_str.upper():
        if h == 12: h = 0
    else:
        if h in [1, 2, 3, 4, 5, 6, 12]: h = (h % 12) + 12
    return h * 60 + mins

sections = [
    # Kindergarten & Elementary (20 sections)
    {'id': 'k2_abdullah', 'code': 'K2-ABD', 'sec_short': 'K2 - Abdullah', 'name': "Kindergarten 2 - Abdullah Ibn Mas'ud (2nd Shift)", 'dept_code': 'elem', 'dept_label': 'Kindergarten & Elementary Department', 'sheet': 'ELEM', 'c_start': 1, 'r_start': 5, 'r_end': 11, 'has_mins': True, 'room': 'ODL Room 1', 'adviser': 'Tchr. Ayah'},
    {'id': 'k2_umar', 'code': 'K2-UMA', 'sec_short': 'K2 - Umar', 'name': 'Kindergarten 2 - Umar Ibn Al-Khattab (2nd Shift)', 'dept_code': 'elem', 'dept_label': 'Kindergarten & Elementary Department', 'sheet': 'ELEM', 'c_start': 9, 'r_start': 5, 'r_end': 11, 'has_mins': True, 'room': 'ODL Room 2', 'adviser': 'Tchr. Joanna'},
    {'id': 'k1_husain', 'code': 'K1-HUS', 'sec_short': 'K1 - Husain', 'name': 'Kindergarten 1 - Husain Ibn Ali (2nd Shift)', 'dept_code': 'elem', 'dept_label': 'Kindergarten & Elementary Department', 'sheet': 'ELEM', 'c_start': 1, 'r_start': 15, 'r_end': 21, 'has_mins': True, 'room': 'ODL Room 3', 'adviser': 'Tchr. Wendy'},
    {'id': 'k2_khabaab', 'code': 'K2-KHA', 'sec_short': 'K2 - Khabaab', 'name': 'Kindergarten 2 - Khabaab Ibn Arat (2nd Shift)', 'dept_code': 'elem', 'dept_label': 'Kindergarten & Elementary Department', 'sheet': 'ELEM', 'c_start': 9, 'r_start': 15, 'r_end': 21, 'has_mins': True, 'room': 'ODL Room 4', 'adviser': 'Tchr. Keychell'},
    {'id': 'g1_suhayb', 'code': 'G1-SUH', 'sec_short': 'G1 - Suhayb', 'name': 'Grade 1 - Suhayb Ar-Rumi (2nd Shift)', 'dept_code': 'elem', 'dept_label': 'Kindergarten & Elementary Department', 'sheet': 'ELEM', 'c_start': 1, 'r_start': 26, 'r_end': 32, 'has_mins': True, 'room': 'ODL Room 5', 'adviser': 'Tchr. Katrina'},
    {'id': 'g1_saad', 'code': 'G1-SAA', 'sec_short': "G1 - Sa'ad", 'name': "Grade 1 - Sa'ad Ibn Abi Waqqaas (2nd Shift)", 'dept_code': 'elem', 'dept_label': 'Kindergarten & Elementary Department', 'sheet': 'ELEM', 'c_start': 9, 'r_start': 26, 'r_end': 32, 'has_mins': True, 'room': 'ODL Room 6', 'adviser': 'Tchr. Sahdia'},
    {'id': 'g2_saeed', 'code': 'G2-SAE', 'sec_short': 'G2 - Saeed', 'name': 'Grade 2 - Saeed Ibn Zayd (2nd Shift)', 'dept_code': 'elem', 'dept_label': 'Kindergarten & Elementary Department', 'sheet': 'ELEM', 'c_start': 1, 'r_start': 41, 'r_end': 47, 'has_mins': True, 'room': 'ODL Room 7', 'adviser': 'Tchr. Sitti'},
    {'id': 'g2_aasim', 'code': 'G2-AAS', 'sec_short': 'G2 - Aasim', 'name': 'Grade 2 - Aasim Ibn Thabit (2nd Shift)', 'dept_code': 'elem', 'dept_label': 'Kindergarten & Elementary Department', 'sheet': 'ELEM', 'c_start': 9, 'r_start': 41, 'r_end': 47, 'has_mins': True, 'room': 'ODL Room 8', 'adviser': 'Tchr. Marham'},
    {'id': 'g3_zayd', 'code': 'G3-ZAY', 'sec_short': 'G3 - Zayd (Girls)', 'name': 'Grade 3 - Zayd Ibn Haritha (2nd Shift) - Girls', 'dept_code': 'elem', 'dept_label': 'Kindergarten & Elementary Department', 'sheet': 'ELEM', 'c_start': 1, 'r_start': 56, 'r_end': 62, 'has_mins': True, 'room': 'ODL Room 9', 'adviser': 'Tchr. Jerlyn'},
    {'id': 'g3_thabit', 'code': 'G3-THA', 'sec_short': 'G3 - Thabit (Boys)', 'name': 'Grade 3 - Thabit Ibn Qays (2nd Shift) - Boys', 'dept_code': 'elem', 'dept_label': 'Kindergarten & Elementary Department', 'sheet': 'ELEM', 'c_start': 9, 'r_start': 56, 'r_end': 62, 'has_mins': True, 'room': 'ODL Room 10', 'adviser': 'Tchr. Jerlyn'},
    {'id': 'g3_asad', 'code': 'G3-ASA', 'sec_short': "G3 - As'ad (Mix)", 'name': "Grade 3 - As'ad Ibn Zurarah (2nd Shift) Mix", 'dept_code': 'elem', 'dept_label': 'Kindergarten & Elementary Department', 'sheet': 'ELEM', 'c_start': 9, 'r_start': 66, 'r_end': 72, 'has_mins': True, 'room': 'ODL Room 11', 'adviser': 'Tchr. Jenny'},
    {'id': 'g4_zubair', 'code': 'G4-ZUB', 'sec_short': 'G4 - Az Zubair', 'name': 'Grade 4 - Az Zubair Ibn Al Awwaam (2nd Shift)', 'dept_code': 'elem', 'dept_label': 'Kindergarten & Elementary Department', 'sheet': 'ELEM', 'c_start': 1, 'r_start': 76, 'r_end': 82, 'has_mins': True, 'room': 'ODL Room 12', 'adviser': 'Tchr. Arvin'},
    {'id': 'g4_ikrimah', 'code': 'G4-IKR', 'sec_short': 'G4 - Ikrimah', 'name': 'Grade 4 - Ikrimah Ibn Abi Jahl (2nd Shift)', 'dept_code': 'elem', 'dept_label': 'Kindergarten & Elementary Department', 'sheet': 'ELEM', 'c_start': 9, 'r_start': 76, 'r_end': 82, 'has_mins': True, 'room': 'ODL Room 13', 'adviser': 'Tchr. Anna'},
    {'id': 'g4_hassan', 'code': 'G4-HAS', 'sec_short': 'G4 - Hassan (Mix)', 'name': 'Grade 4 - Hassan Ibn Thabit (2nd Shift) - Mix', 'dept_code': 'elem', 'dept_label': 'Kindergarten & Elementary Department', 'sheet': 'ELEM', 'c_start': 17, 'r_start': 76, 'r_end': 82, 'has_mins': True, 'room': 'ODL Room 14', 'adviser': 'Tchr. Monisa'},
    {'id': 'g5_ayyash', 'code': 'G5-AYY', 'sec_short': 'G5 - Ayyash', 'name': "Grade 5 - Ayyash Ibn Abi Rabi'ah (1st Shift Masterlist)", 'dept_code': 'elem', 'dept_label': 'Kindergarten & Elementary Department', 'sheet': 'ELEM', 'c_start': 1, 'r_start': 91, 'r_end': 97, 'has_mins': True, 'room': 'ODL Room 15', 'adviser': 'Tchr. Saimona'},
    {'id': 'g5_musab', 'code': 'G5-MUS', 'sec_short': "G5 - Mus'ab", 'name': "Grade 5 - Mus'ab Ibn Abdul Mutalib (2nd Shift)", 'dept_code': 'elem', 'dept_label': 'Kindergarten & Elementary Department', 'sheet': 'ELEM', 'c_start': 9, 'r_start': 91, 'r_end': 97, 'has_mins': True, 'room': 'ODL Room 16', 'adviser': 'Tchr. Jessa'},
    {'id': 'g5_harith', 'code': 'G5-HAR', 'sec_short': 'G5 - Al Harith', 'name': 'Grade 5 - Al Harith Bin Awf (2nd Shift)', 'dept_code': 'elem', 'dept_label': 'Kindergarten & Elementary Department', 'sheet': 'ELEM', 'c_start': 17, 'r_start': 91, 'r_end': 97, 'has_mins': True, 'room': 'ODL Room 17', 'adviser': 'Tchr. Anna'},
    {'id': 'g5_jafar', 'code': 'G5-JAF', 'sec_short': "G5 - Ja'far (Mix)", 'name': "Grade 5 - Ja'far Ibn Abi Talib (2nd Shift) - Mix", 'dept_code': 'elem', 'dept_label': 'Kindergarten & Elementary Department', 'sheet': 'ELEM', 'c_start': 25, 'r_start': 91, 'r_end': 97, 'has_mins': True, 'room': 'ODL Room 18', 'adviser': 'Tchr. Saimona'},
    {'id': 'g6_khaleed', 'code': 'G6-KHA', 'sec_short': 'G6 - Khaleed', 'name': 'Grade 6 - Khaleed Ibn Waleed (2nd Shift)', 'dept_code': 'elem', 'dept_label': 'Kindergarten & Elementary Department', 'sheet': 'ELEM', 'c_start': 1, 'r_start': 106, 'r_end': 112, 'has_mins': True, 'room': 'ODL Room 19', 'adviser': 'Tchr. Jenny'},
    {'id': 'g6_dihya', 'code': 'G6-DIH', 'sec_short': 'G6 - Dihya (Girls)', 'name': 'Grade 6 - Dihya Ibn Khalifah (2nd Shift) Girls', 'dept_code': 'elem', 'dept_label': 'Kindergarten & Elementary Department', 'sheet': 'ELEM', 'c_start': 9, 'r_start': 106, 'r_end': 112, 'has_mins': True, 'room': 'ODL Room 20', 'adviser': 'Tchr. Fhairudz'},

    # Junior High School (6 sections)
    {'id': 'g7_anas', 'code': 'G7-ANA', 'sec_short': 'G7 - Anas (Mix)', 'name': 'Grade 7 - Anas Ibn Malik (2nd Shift) - Mix', 'dept_code': 'jhs', 'dept_label': 'Junior High School Department', 'sheet': 'HS SCHED (NEW)', 'c_start': 1, 'r_start': 7, 'r_end': 13, 'has_mins': True, 'room': 'ODL JHS Room 1', 'adviser': 'Tchr. Jayra'},
    {'id': 'g8_muadh', 'code': 'G8-MUA', 'sec_short': "G8 - Mu'adh (Boys)", 'name': "Grade 8 - Mu'adh Ibn Jabal (2nd Shift) - Boys", 'dept_code': 'jhs', 'dept_label': 'Junior High School Department', 'sheet': 'HS SCHED (NEW)', 'c_start': 1, 'r_start': 18, 'r_end': 24, 'has_mins': True, 'room': 'ODL JHS Room 2', 'adviser': 'Tchr. Radzmia'},
    {'id': 'g8_nuaym', 'code': 'G8-NUA', 'sec_short': 'G8 - Nuaym (Mix)', 'name': "Grade 8 - Nuaym Ibn Mas'ud (2nd Shift) Mix", 'dept_code': 'jhs', 'dept_label': 'Junior High School Department', 'sheet': 'HS SCHED (NEW)', 'c_start': 1, 'r_start': 30, 'r_end': 36, 'has_mins': True, 'room': 'ODL JHS Room 3', 'adviser': 'Tchr. Jayra'},
    {'id': 'g9_abudharr', 'code': 'G9-ABD', 'sec_short': 'G9 - Abu Dharr (Boys)', 'name': 'Grade 9 - Abu Dharr Al Ghifarri (2nd Shift) Boys', 'dept_code': 'jhs', 'dept_label': 'Junior High School Department', 'sheet': 'HS SCHED (NEW)', 'c_start': 1, 'r_start': 42, 'r_end': 48, 'has_mins': True, 'room': 'ODL JHS Room 4', 'adviser': 'Tchr. Rowena'},
    {'id': 'g9_abujandal', 'code': 'G9-ABJ', 'sec_short': 'G9 - Abu Jandal (Girls)', 'name': 'Grade 9 - Abu Jandal Ibn Suhayl (2nd Shift) Girls', 'dept_code': 'jhs', 'dept_label': 'Junior High School Department', 'sheet': 'HS SCHED (NEW)', 'c_start': 1, 'r_start': 53, 'r_end': 59, 'has_mins': True, 'room': 'ODL JHS Room 5', 'adviser': 'Tchr. Nadzra'},
    {'id': 'g10_abuayyub', 'code': 'G10-AYY', 'sec_short': 'G10 - Abu Ayyub (Boys)', 'name': 'Grade 10 - Abu Ayyub Al-Ansari (2nd Shift) Boys', 'dept_code': 'jhs', 'dept_label': 'Junior High School Department', 'sheet': 'HS SCHED (NEW)', 'c_start': 1, 'r_start': 64, 'r_end': 70, 'has_mins': True, 'room': 'ODL JHS Room 6', 'adviser': 'Tchr. Jhelyn'},

    # Senior High School (3 sections)
    {'id': 'g11_b_term2', 'code': 'G11-B2', 'sec_short': 'G11 Boys (Term 2)', 'name': 'Grade 11 - 2nd Shift Boys (2nd Term / Active)', 'dept_code': 'shs', 'dept_label': 'Senior High School Department', 'sheet': 'HS SCHED (NEW)', 'c_start': 1, 'r_start': 78, 'r_end': 85, 'has_mins': True, 'room': 'ODL SHS Room 1', 'adviser': 'Tchr. Nadzra'},
    {'id': 'g11_b_sem1', 'code': 'G11-B1', 'sec_short': 'G11 Boys (Sem 1)', 'name': 'Grade 11 - 2nd Shift Boys (1st Semester)', 'dept_code': 'shs', 'dept_label': 'Senior High School Department', 'sheet': 'HS SCHED', 'c_start': 26, 'r_start': 67, 'r_end': 73, 'has_mins': True, 'room': 'ODL SHS Room 1', 'adviser': 'Tchr. Thea'},
    {'id': 'g12_b_sem1', 'code': 'G12-B1', 'sec_short': 'G12 Suhayb (Boys)', 'name': 'Grade 12 - Suhayb Ar-Rumi (Boys) - 2nd Shift', 'dept_code': 'shs', 'dept_label': 'Senior High School Department', 'sheet': 'HS SCHED', 'c_start': 18, 'r_start': 83, 'r_end': 89, 'has_mins': True, 'room': 'ODL SHS Room 2', 'adviser': 'Tchr. Thea'},
]

print(f"Configured {len(sections)} sections for ODL 2nd Shift.")

def build_header_html(form_title, form_subtitle=None):
    sub_markup = f'<div class="form-sub">{form_subtitle}</div>' if form_subtitle else '<div class="form-sub">Online Distance Learning (Second Shift) Modality &bull; School Year 2026 - 2027</div>'
    return f'''
        <div class="sheet-header">
          <div class="header-logo-side">
            <img class="header-logo deped-img" alt="Department of Education Logo">
          </div>
          <div class="header-center-text">
            <div class="arabic-header" dir="rtl" lang="ar">المدرسة المنورة الإسلامية</div>
            <div class="school-name">AL MUNAWWARA ISLAMIC SCHOOL</div>
            <div class="form-title">{form_title}</div>
            {sub_markup}
          </div>
          <div class="header-logo-side">
            <img class="header-logo amis-img" alt="Al Munawwara Islamic School Official Seal">
          </div>
        </div>
    '''

all_master_attendance_rows = []
teacher_monitor_map = {}

for sec in sections:
    s_grid = elem_grid if sec['sheet'] == 'ELEM' else (hs_new_grid if sec['sheet'] == 'HS SCHED (NEW)' else hs_grid)
    c_start = sec['c_start']
    
    for r in range(sec['r_start'], sec['r_end'] + 1):
        raw_t = s_grid.get((r, c_start), '')
        t_slot = clean_time(raw_t)
        mins_val = parse_mins(s_grid.get((r, c_start + 1), ''))
        if not t_slot: continue
        
        for d_idx, day in enumerate(days):
            day_abbr = day_abbr_map.get(day, day[:3].upper())
            c_val = s_grid.get((r, c_start + 2 + d_idx), '').strip()
            if not c_val: continue
            
            subj, tchr, kind = clean_parse(c_val)
            if kind == 'CLASS' and tchr and tchr != 'TBA':
                # Add to master attendance
                all_master_attendance_rows.append({
                    'day': day,
                    'day_abbr': day_abbr,
                    'time': t_slot,
                    'sort_time': time_to_sort_key(t_slot),
                    'mins': mins_val,
                    'section': sec['name'],
                    'sec_code': sec['code'],
                    'subject': subj,
                    'teacher': tchr,
                    'dept_code': sec['dept_code'],
                    'dept_label': sec['dept_label'],
                    'room': sec['room']
                })
                # Add to teacher attendance map
                if tchr not in teacher_monitor_map:
                    teacher_monitor_map[tchr] = {
                        'name': tchr,
                        'dept': sec['dept_label'],
                        'dept_code': sec['dept_code'],
                        'items': []
                    }
                teacher_monitor_map[tchr]['items'].append({
                    'day': day,
                    'day_abbr': day_abbr,
                    'day_idx': day_order.get(day, 0),
                    'time': t_slot,
                    'sort_time': time_to_sort_key(t_slot),
                    'mins': mins_val,
                    'section': sec['sec_short'],
                    'subject': subj.upper(),
                    'room': sec['room']
                })

for t_name, t_data in teacher_monitor_map.items():
    if any(t_name.startswith(p) for p in ['Ust.', 'Ustadh', 'Ustadha', 'Alim']):
        t_data['cat'] = 'isal'
        t_data['dept'] = 'ISAL & Islamic Studies Department'
    elif any('G11' in it['section'] or 'G12' in it['section'] for it in t_data['items']):
        t_data['cat'] = 'shs'
        t_data['dept'] = 'Senior High School Faculty'
    elif any('G7' in it['section'] or 'G8' in it['section'] or 'G9' in it['section'] or 'G10' in it['section'] for it in t_data['items']):
        t_data['cat'] = 'jhs'
        t_data['dept'] = 'Junior High School Faculty'
    else:
        t_data['cat'] = 'elem'
        t_data['dept'] = 'Elementary Faculty'
    t_data['items'].sort(key=lambda x: (x['day_idx'], x['sort_time'], x['section']))

teachers_list = sorted(teacher_monitor_map.values(), key=lambda x: x['name'])
print(f"Compiled {len(all_master_attendance_rows)} master attendance rows and {len(teachers_list)} teachers for ODL 2nd Shift.")

html_out = []
html_out.append('''<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <link rel="preconnect" href="https://fonts.googleapis.com">
  <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
  <link href="https://fonts.googleapis.com/css2?family=Amiri:wght@700&display=swap" rel="stylesheet">
  <title>ODL Second Shift Instructional Monitoring & Section Schedules - Al Munawwara Islamic School</title>
  <style>
    @page {
      size: A4 portrait;
      margin: 5mm 6mm 5mm 6mm;
    }
    * {
      box-sizing: border-box;
      font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Arial, sans-serif;
    }
    body {
      margin: 0;
      padding: 0;
      background: #f1f5f9;
      color: #0f172a;
      -webkit-print-color-adjust: exact;
      print-color-adjust: exact;
    }
    .toolbar {
      background: #ffffff;
      color: #0f172a;
      padding: 12px 24px;
      position: sticky;
      top: 0;
      z-index: 999;
      box-shadow: 0 2px 10px rgba(0,0,0,0.06);
      border-bottom: 2px solid #059669;
    }
    .toolbar-header {
      display: flex;
      justify-content: space-between;
      align-items: center;
      margin-bottom: 12px;
      flex-wrap: wrap;
      gap: 12px;
    }
    .toolbar-brand {
      display: flex;
      align-items: center;
      gap: 12px;
    }
    .toolbar-logo {
      width: 44px;
      height: 44px;
      object-fit: contain;
    }
    .toolbar-title {
      font-size: 16px;
      font-weight: 800;
      color: #064e3b;
      letter-spacing: 0.3px;
      display: flex;
      align-items: center;
      gap: 10px;
    }
    .badge-shift {
      background: #dbeafe;
      color: #1e40af;
      border: 1px solid #3b82f6;
      font-size: 11px;
      padding: 2px 8px;
      border-radius: 9999px;
      font-weight: 800;
      text-transform: uppercase;
    }
    .modality-bar {
      display: flex;
      align-items: center;
      gap: 8px;
      margin-bottom: 12px;
      padding-bottom: 10px;
      border-bottom: 1px solid #e2e8f0;
      flex-wrap: wrap;
    }
    .modality-label {
      font-size: 11px;
      font-weight: 800;
      color: #64748b;
      text-transform: uppercase;
      letter-spacing: 0.4px;
      margin-right: 4px;
    }
    .modality-pill {
      display: inline-flex;
      align-items: center;
      gap: 6px;
      padding: 5px 12px;
      border-radius: 6px;
      font-size: 11.5px;
      font-weight: 700;
      text-decoration: none;
      transition: all 0.15s ease;
      border: 1px solid #cbd5e1;
      background: #f8fafc;
      color: #334155;
    }
    .modality-pill:hover {
      background: #f1f5f9;
      color: #0f172a;
      border-color: #94a3b8;
    }
    .modality-pill.active {
      background: #064e3b;
      color: #ffffff;
      border-color: #064e3b;
    }
    .modality-dot {
      width: 7px;
      height: 7px;
      border-radius: 50%;
      display: inline-block;
    }
    .dot-active {
      background: #10b981;
      box-shadow: 0 0 0 2px rgba(16, 185, 129, 0.3);
    }
    .toolbar-actions {
      display: flex;
      gap: 8px;
      align-items: center;
      flex-wrap: wrap;
    }
    .btn {
      padding: 7px 14px;
      font-size: 12px;
      font-weight: 700;
      border-radius: 6px;
      cursor: pointer;
      display: inline-flex;
      align-items: center;
      gap: 6px;
      border: 1px solid transparent;
      transition: all 0.15s ease;
      text-decoration: none;
    }
    .btn-primary { background: #059669; color: #ffffff; }
    .btn-primary:hover { background: #047857; }
    .btn-blue { background: #0284c7; color: #ffffff; }
    .btn-blue:hover { background: #0369a1; }
    .btn-purple { background: #7c3aed; color: #ffffff; }
    .btn-purple:hover { background: #6d28d9; }
    .btn-outline { background: #ffffff; color: #334155; border-color: #cbd5e1; }
    .btn-outline:hover { background: #f8fafc; color: #0f172a; border-color: #94a3b8; }

    .view-tabs {
      display: flex;
      gap: 8px;
      margin-bottom: 12px;
      border-bottom: 1px solid #e2e8f0;
      padding-bottom: 8px;
      overflow-x: auto;
    }
    .view-tab {
      background: #f8fafc;
      color: #475569;
      border: 1px solid #cbd5e1;
      padding: 7px 16px;
      border-radius: 6px;
      font-size: 12px;
      font-weight: 600;
      cursor: pointer;
      white-space: nowrap;
      transition: all 0.15s;
    }
    .view-tab:hover {
      background: #f1f5f9;
      color: #0f172a;
    }
    .view-tab.active {
      background: #059669;
      color: #ffffff;
      border-color: #059669;
      font-weight: 700;
      box-shadow: 0 2px 4px rgba(5, 150, 105, 0.2);
    }
    .filter-grid {
      display: grid;
      grid-template-columns: repeat(auto-fit, minmax(180px, 1fr));
      gap: 12px;
      align-items: center;
    }
    .filter-group {
      display: flex;
      flex-direction: column;
      gap: 4px;
    }
    .filter-group label {
      font-size: 11px;
      font-weight: 700;
      color: #475569;
      text-transform: uppercase;
      letter-spacing: 0.3px;
    }
    .filter-group select, .filter-group input {
      background: #ffffff;
      border: 1px solid #cbd5e1;
      color: #0f172a;
      padding: 7px 10px;
      font-size: 12px;
      border-radius: 6px;
      outline: none;
      transition: border-color 0.15s;
    }
    .filter-group select:focus, .filter-group input:focus {
      border-color: #059669;
      box-shadow: 0 0 0 2px rgba(5, 150, 105, 0.15);
    }
    .status-summary {
      margin-top: 10px;
      font-size: 11.5px;
      color: #64748b;
      display: flex;
      justify-content: space-between;
      align-items: center;
      flex-wrap: wrap;
      gap: 8px;
    }
    .sheet-wrapper {
      padding: 24px 10px 60px 10px;
      display: flex;
      flex-direction: column;
      align-items: center;
      gap: 24px;
    }
    .page-sheet {
      width: 210mm;
      min-height: 290mm;
      padding: 5mm 6mm;
      background: #ffffff;
      box-shadow: 0 4px 15px rgba(0,0,0,0.08);
      border: 1px solid #e2e8f0;
      position: relative;
      display: flex;
      flex-direction: column;
      justify-content: space-between;
      box-sizing: border-box;
      page-break-after: always;
      break-after: page;
    }
    .sheet-content {
      flex: 1;
      display: flex;
      flex-direction: column;
    }
    .sheet-header {
      display: flex;
      align-items: center;
      justify-content: space-between;
      border-bottom: 2px solid #0f172a;
      padding-bottom: 4px;
      margin-bottom: 4px;
      gap: 10px;
    }
    .header-logo-side {
      width: 48px;
      display: flex;
      justify-content: center;
      align-items: center;
      flex-shrink: 0;
    }
    .header-logo {
      width: 48px;
      height: 48px;
      object-fit: contain;
    }
    .header-center-text {
      flex: 1;
      text-align: center;
    }
    .arabic-header {
      font-family: 'Amiri', 'Traditional Arabic', 'Times New Roman', serif;
      font-size: 14pt;
      font-weight: 700;
      color: #064e3b;
      direction: rtl;
      line-height: 1.15;
      margin-bottom: 1px;
      text-align: center;
    }
    .school-name {
      font-size: 11pt;
      font-weight: 900;
      color: #0f172a;
      letter-spacing: 0.5px;
      margin: 1px 0 2px 0;
      text-transform: uppercase;
      line-height: 1.15;
    }
    .form-title {
      font-size: 9pt;
      font-weight: 800;
      color: #064e3b;
      text-transform: uppercase;
      letter-spacing: 0.5px;
      line-height: 1.15;
    }
    .form-sub {
      font-size: 7pt;
      color: #475569;
      font-weight: 600;
      margin-top: 1px;
    }
    .meta-box {
      border: 1px solid #94a3b8;
      background: #f8fafc;
      border-radius: 3px;
      padding: 4px 8px;
      margin-bottom: 5px;
      display: grid;
      grid-template-columns: 1.3fr 1.2fr 1fr;
      gap: 3px 12px;
      font-size: 7.5pt;
    }
    .meta-row {
      display: flex;
      align-items: center;
      gap: 4px;
    }
    .meta-lbl {
      font-weight: 700;
      color: #334155;
      white-space: nowrap;
    }
    .meta-val {
      font-weight: 600;
      color: #0f172a;
      border-bottom: 1px dotted #94a3b8;
      flex: 1;
      min-height: 12px;
    }
    .td-bold {
      font-weight: 800;
      color: #064e3b;
    }
    .section-title {
      font-size: 7.6pt;
      font-weight: 800;
      color: #064e3b;
      text-transform: uppercase;
      letter-spacing: 0.3px;
      margin: 3px 0 2px 0;
      display: flex;
      align-items: center;
      gap: 6px;
    }
    .section-title::after {
      content: '';
      flex: 1;
      height: 1px;
      background: #cbd5e1;
    }
    table.schedule-table {
      width: 100%;
      border-collapse: collapse;
      font-size: 6.9pt;
      line-height: 1.15;
      margin-bottom: 5px;
    }
    table.schedule-table th, table.schedule-table td {
      border: 1px solid #94a3b8;
      padding: 2px 4px;
      vertical-align: middle;
    }
    table.schedule-table thead th {
      background: #e2e8f0;
      color: #0f172a;
      font-weight: 800;
      text-align: center;
      font-size: 7pt;
      padding: 3px 4px;
    }
    .time-slot-col {
      width: 17%;
      text-align: center;
      font-weight: 700;
      background: #f8fafc;
      white-space: nowrap;
    }
    .mins-col {
      width: 5%;
      text-align: center;
      font-weight: 700;
      background: #f8fafc;
    }
    .day-col {
      width: 15.6%;
      text-align: center;
    }
    .routine-row-cell {
      background: #f1f5f9;
      color: #475569;
      font-weight: 700;
      text-align: center;
      font-size: 6.8pt;
      letter-spacing: 0.4px;
      padding: 2px 4px;
    }
    .class-cell {
      background: #ecfdf5;
      border: 1px solid #6ee7b7 !important;
      text-align: center;
      padding: 2px 3px;
    }
    .cell-subj {
      font-weight: 800;
      color: #064e3b;
      font-size: 6.9pt;
      line-height: 1.1;
    }
    .cell-tchr {
      font-weight: 600;
      color: #0284c7;
      font-size: 6.5pt;
      margin-top: 1px;
    }

    /* TEACHER SHEET TABLE */
    table.teacher-record-table {
      width: 100%;
      border-collapse: collapse;
      font-size: 7.5pt;
      line-height: 1.15;
    }
    table.teacher-record-table th, table.teacher-record-table td {
      border: 1px solid #94a3b8;
      padding: 3.5px 4.5px;
      vertical-align: middle;
    }
    table.teacher-record-table thead th {
      background: #f1f5f9;
      color: #0f172a;
      font-weight: 800;
      text-align: center;
      font-size: 7.5pt;
      padding: 4px 4px;
      text-transform: uppercase;
      letter-spacing: 0.3px;
    }
    .th-num { width: 3%; }
    .th-day { width: 6.5%; }
    .th-time { width: 16.5%; }
    .th-mins { width: 4.5%; }
    .th-in { width: 8%; }
    .th-out { width: 8%; }
    .th-grade { width: 13.5%; }
    .th-subject { width: 18.5%; }
    .th-room { width: 5%; }
    .th-status { width: 11.5%; }
    .th-remarks { width: 5%; }
    .td-center { text-align: center; }

    .day-cell {
      font-weight: 800;
      color: #064e3b;
      text-align: center;
      text-transform: uppercase;
      letter-spacing: 0.3px;
      font-size: 7.8pt;
      white-space: nowrap;
    }
    .time-slot {
      font-weight: 800;
      color: #0f172a;
      text-align: center;
      white-space: nowrap;
      font-size: 7.8pt;
      letter-spacing: 0.2px;
      font-variant-numeric: tabular-nums;
    }
    .subject-cell {
      font-weight: 900;
      color: #064e3b;
      text-transform: uppercase;
      letter-spacing: 0.3px;
      font-size: 8pt;
    }
    .grade-cell {
      font-weight: 700;
      color: #1e293b;
      font-size: 7.8pt;
      white-space: nowrap;
    }

    /* 2x2 STATUS CHECKLIST */
    .status-grid {
      display: flex;
      flex-direction: column;
      gap: 1.5px;
      font-size: 6.8pt;
      font-weight: 700;
    }
    .status-cell-row {
      display: flex;
      justify-content: space-between;
      gap: 3px;
    }
    .status-chk {
      display: inline-flex;
      align-items: center;
      gap: 2px;
      cursor: pointer;
      user-select: none;
      white-space: nowrap;
    }
    .status-chk input {
      margin: 0;
      width: 10px;
      height: 10px;
      cursor: pointer;
    }

    /* SIGNATURE FOOTER */
    .sign-row {
      display: grid;
      grid-template-columns: repeat(3, 1fr);
      gap: 16px;
      margin-top: 8px;
      padding-top: 6px;
      border-top: 1px solid #cbd5e1;
      font-size: 7.2pt;
    }
    .sign-col {
      text-align: center;
    }
    .sign-line {
      border-bottom: 1px solid #0f172a;
      height: 22px;
      margin-bottom: 3px;
    }
    .sign-label {
      font-weight: 800;
      color: #1e293b;
      text-transform: uppercase;
      font-size: 6.8pt;
    }
    .sign-title {
      font-size: 6.5pt;
      color: #64748b;
    }

    /* ATTENDANCE TABLE VIEW */
    .attendance-view-container {
      width: 100%;
      max-width: 1200px;
      margin: 0 auto;
      background: #ffffff;
      border-radius: 8px;
      border: 1px solid #cbd5e1;
      box-shadow: 0 4px 15px rgba(0,0,0,0.05);
      padding: 20px;
      display: none;
    }
    .attendance-table {
      width: 100%;
      border-collapse: collapse;
      font-size: 11.5px;
    }
    .attendance-table th, .attendance-table td {
      border: 1px solid #cbd5e1;
      padding: 6px 8px;
      vertical-align: middle;
    }
    .attendance-table thead th {
      background: #064e3b;
      color: #ffffff;
      font-weight: 700;
      text-align: left;
    }
    .day-matrix-tabs {
      display: flex;
      gap: 6px;
      margin-bottom: 16px;
      border-bottom: 1px solid #e2e8f0;
      padding-bottom: 8px;
    }
    .day-matrix-tab {
      padding: 6px 14px;
      border-radius: 6px;
      font-size: 12px;
      font-weight: 700;
      border: 1px solid #cbd5e1;
      background: #f8fafc;
      color: #475569;
      cursor: pointer;
    }
    .day-matrix-tab.active {
      background: #059669;
      color: #ffffff;
      border-color: #059669;
    }

    @media print {
      body {
        background: #ffffff !important;
        margin: 0;
        padding: 0;
      }
      .toolbar, .no-print {
        display: none !important;
      }
      .sheet-wrapper {
        padding: 0 !important;
        gap: 0 !important;
      }
      .page-sheet {
        box-shadow: none !important;
        border: none !important;
        margin: 0 !important;
        padding: 4mm 6mm !important;
        width: 100% !important;
        min-height: auto !important;
        page-break-after: always !important;
        break-after: page !important;
      }
      .page-sheet.hide-in-print {
        display: none !important;
      }
      .attendance-view-container {
        display: block !important;
        box-shadow: none !important;
        border: none !important;
        padding: 0 !important;
        max-width: 100% !important;
      }
    }
  </style>
</head>
<body>

  <div class="toolbar no-print">
    <div class="toolbar-header">
      <div class="toolbar-brand">
        <img class="toolbar-logo deped-img" alt="DepEd Logo">
        <img class="toolbar-logo amis-img" alt="AMIS Logo">
        <div>
          <div class="toolbar-title">
            AL MUNAWWARA ISLAMIC SCHOOL
            <span class="badge-shift">ODL Second Shift</span>
          </div>
          <div style="font-size: 11px; color: #64748b; font-weight: 600;">Instructional Monitoring & Daily Per-Teacher Loads</div>
        </div>
      </div>
      <div class="toolbar-actions">
        <a href="/odl-second-shift-teacher.html" class="btn btn-purple">Open Standalone Teacher Portal</a>
        <button class="btn btn-primary" onclick="window.print()">Print Current View</button>
        <button class="btn btn-outline" id="btnPrintAll" onclick="handlePrintAll()">Print All Sheets</button>
      </div>
    </div>

    <!-- Modality Switcher -->
    <div class="modality-bar">
      <span class="modality-label">Learning Modality:</span>
      <a href="/f2f-monitoring.html" class="modality-pill"><span class="modality-dot"></span>Face-to-Face</a>
      <a href="/odl-first-shift.html" class="modality-pill"><span class="modality-dot"></span>ODL First Shift</a>
      <a href="/odl-second-shift.html" class="modality-pill active"><span class="modality-dot dot-active"></span>ODL Second Shift</a>
      <a href="/all-teachers-monitoring.html" class="modality-pill" style="border-color:#3b82f6;color:#2563eb;background:#eff6ff;"><span class="modality-dot" style="background:#2563eb;"></span>All Modalities (152 Teachers)</a>
    </div>

    <!-- View Switcher Tabs -->
    <div class="view-tabs">
      <button class="view-tab active" id="tabSections" onclick="switchView('sections')">Section Classrooms (29 Sections)</button>
      <button class="view-tab" id="tabTeachers" onclick="switchView('teachers')">Daily Per-Teacher Attendance Sheets (54 Teachers)</button>
      <button class="view-tab" id="tabMasterAttendance" onclick="switchView('master')">Master Daily Attendance Matrix</button>
    </div>

    <!-- Section Filter Controls -->
    <div id="controlsSections" class="filter-grid">
      <div class="filter-group">
        <label>Filter Department</label>
        <select id="deptFilter" onchange="applySectionFilter()">
          <option value="all">All Departments</option>
          <option value="elem">Kindergarten & Elementary (20 Sections)</option>
          <option value="jhs">Junior High School (6 Sections)</option>
          <option value="shs">Senior High School (3 Sections)</option>
        </select>
      </div>
      <div class="filter-group">
        <label>Quick Jump Section</label>
        <select id="sectionJump" onchange="jumpToSection()">
          <option value="all">-- Show All Matching Sections --</option>
''')

for sec in sections:
    html_out.append(f'          <option value="{sec["id"]}">{sec["name"]}</option>\n')

html_out.append('''        </select>
      </div>
      <div class="filter-group">
        <label>Search Teacher or Subject</label>
        <input type="text" id="secSearch" placeholder="Type teacher or subject..." onkeyup="searchSectionClasses()">
      </div>
    </div>

    <!-- Teacher Filter Controls -->
    <div id="controlsTeachers" class="filter-grid" style="display: none;">
      <div class="filter-group">
        <label>Faculty Department</label>
        <select id="tchrDeptFilter" onchange="onTchrDeptChange()">
          <option value="all">All Faculty (54 Teachers)</option>
          <option value="isal">ISAL & Islamic Studies Department</option>
          <option value="elem">Elementary Faculty</option>
          <option value="jhs">Junior High School Faculty</option>
          <option value="shs">Senior High School Faculty</option>
        </select>
      </div>
      <div class="filter-group">
        <label>Select Teacher</label>
        <select id="teacherSelect" onchange="onTeacherSelectChange()">
''')

for t in teachers_list:
    cat_code = t.get('cat', 'elem')
    html_out.append(f'          <option value="{html.escape(t["name"])}" data-cat="{cat_code}">{html.escape(t["name"])} ({html.escape(t["dept"])})</option>\n')

html_out.append('''        </select>
      </div>
      <div class="filter-group">
        <label>Search Teacher</label>
        <input type="text" id="tchrSearch" placeholder="Type teacher name..." onkeyup="searchTeacherByName()">
      </div>
    </div>

    <!-- Master Attendance Matrix Filter Controls -->
    <div id="controlsMaster" class="filter-grid" style="display: none;">
      <div class="filter-group">
        <label>Filter By Day</label>
        <select id="matrixDaySelect" onchange="renderMatrixTable()">
          <option value="Sunday">Sunday (SUN)</option>
          <option value="Monday">Monday (MON)</option>
          <option value="Tuesday">Tuesday (TUE)</option>
          <option value="Wednesday">Wednesday (WED)</option>
          <option value="Thursday">Thursday (THU)</option>
        </select>
      </div>
      <div class="filter-group">
        <label>Department</label>
        <select id="matrixDeptFilter" onchange="renderMatrixTable()">
          <option value="all">All Departments</option>
          <option value="elem">Kindergarten & Elementary</option>
          <option value="jhs">Junior High School</option>
          <option value="shs">Senior High School</option>
        </select>
      </div>
      <div class="filter-group">
        <label>Search Class</label>
        <input type="text" id="matrixSearch" placeholder="Search teacher, subject, or section..." onkeyup="renderMatrixTable()">
      </div>
    </div>

    <div class="status-summary">
      <span id="resultCountInfo">Displaying 29 Section Monitoring Sheets (ODL Second Shift)</span>
      <span>Official School Year 2026 - 2027</span>
    </div>
  </div>

  <!-- SECTION SHEETS CONTAINER -->
  <div class="sheet-wrapper" id="sectionsContainer">
''')

for s_idx, sec in enumerate(sections):
    s_grid = elem_grid if sec['sheet'] == 'ELEM' else (hs_new_grid if sec['sheet'] == 'HS SCHED (NEW)' else hs_grid)
    c_start = sec['c_start']
    
    html_out.append(f'''
    <!-- SECTION SHEET: {sec["name"]} -->
    <div class="page-sheet section-sheet dept-target-{sec["dept_code"]}" id="sheet-{sec["id"]}" data-id="{sec["id"]}" data-dept="{sec["dept_code"]}">
      <div class="sheet-content">
        {build_header_html("CLASSROOM INSTRUCTIONAL MONITORING & SCHEDULE SHEET", "Online Distance Learning (Second Shift) &bull; SY 2026 - 2027")}
        
        <div class="meta-box">
          <div class="meta-row"><span class="meta-lbl">Grade & Section:</span><span class="meta-val td-bold">{sec["name"]}</span></div>
          <div class="meta-row"><span class="meta-lbl">Department:</span><span class="meta-val">{sec["dept_label"]}</span></div>
          <div class="meta-row"><span class="meta-lbl">Assigned Room:</span><span class="meta-val">{sec["room"]}</span></div>
          <div class="meta-row"><span class="meta-lbl">Class Adviser:</span><span class="meta-val">{sec["adviser"]}</span></div>
          <div class="meta-row"><span class="meta-lbl">Shift Modality:</span><span class="meta-val td-bold">Second Shift (Afternoon/Evening)</span></div>
          <div class="meta-row"><span class="meta-lbl">School Year:</span><span class="meta-val">2026 - 2027</span></div>
        </div>

        <div class="section-title">Weekly Instructional Timetable</div>
        <table class="schedule-table">
          <thead>
            <tr>
              <th class="time-slot-col">Time Slot</th>
              <th class="mins-col">Mins</th>
              <th class="day-col">SUN</th>
              <th class="day-col">MON</th>
              <th class="day-col">TUE</th>
              <th class="day-col">WED</th>
              <th class="day-col">THU</th>
            </tr>
          </thead>
          <tbody>
    ''')

    for r in range(sec['r_start'], sec['r_end'] + 1):
        raw_t = s_grid.get((r, c_start), '')
        t_slot = clean_time(raw_t)
        mins_val = parse_mins(s_grid.get((r, c_start + 1), ''))
        if not t_slot: continue
        
        is_routine = False
        sample_cell = s_grid.get((r, c_start + 2), '')
        _, _, s_type = clean_parse(sample_cell)
        if s_type == 'ROUTINE': is_routine = True
        
        html_out.append('            <tr>\n')
        html_out.append(f'              <td class="time-slot-col">{html.escape(t_slot)}</td>\n')
        html_out.append(f'              <td class="mins-col">{mins_val}</td>\n')
        
        if is_routine:
            html_out.append(f'              <td colspan="5" class="routine-row-cell">{html.escape(sample_cell)}</td>\n')
        else:
            for d_idx, day in enumerate(days):
                c_val = s_grid.get((r, c_start + 2 + d_idx), '').strip()
                if not c_val:
                    html_out.append('              <td class="day-col">&mdash;</td>\n')
                    continue
                subj, tchr, kind = clean_parse(c_val)
                if kind == 'ROUTINE':
                    html_out.append(f'              <td class="day-col routine-row-cell">{html.escape(c_val)}</td>\n')
                elif kind == 'CLASS':
                    tchr_clean = tchr if tchr != 'TBA' else 'TBA'
                    html_out.append(f'''              <td class="day-col class-cell">
                    <div class="cell-subj">{html.escape(subj.upper())}</div>
                    <div class="cell-tchr">{html.escape(tchr_clean)}</div>
                  </td>\n''')
                else:
                    html_out.append(f'              <td class="day-col">{html.escape(c_val)}</td>\n')
        html_out.append('            </tr>\n')

    html_out.append(f'''          </tbody>
        </table>

        <div class="section-title">Instructional Attendance & Delivery Sign-Off</div>
        <div style="font-size: 6.8pt; color: #475569; margin-bottom: 4px;">
          This official monitoring record certifies that instructional periods have been conducted according to the DepEd MATATAG / DepEd-approved curriculum and schedule.
        </div>
      </div>

      <div class="sign-row">
        <div class="sign-col">
          <div class="sign-line"></div>
          <div class="sign-label">Class Adviser / Homeroom Teacher</div>
          <div class="sign-title">Prepared & Certified Correct</div>
        </div>
        <div class="sign-col">
          <div class="sign-line"></div>
          <div class="sign-label">Academic Coordinator / Department Head</div>
          <div class="sign-title">Verified & Monitored</div>
        </div>
        <div class="sign-col">
          <div class="sign-line"></div>
          <div class="sign-label">School Principal / Directress</div>
          <div class="sign-title">Noted & Approved</div>
        </div>
      </div>
    </div>
    ''')

html_out.append('''  </div>

  <!-- TEACHER SHEETS CONTAINER -->
  <div class="sheet-wrapper" id="teachersContainer" style="display: none;">
''')

ROWS_PER_PAGE = 14

for t_idx, t in enumerate(teachers_list):
    t_name = t['name']
    t_dept = t['dept']
    t_cat = t.get('cat', 'elem')
    items = t['items']
    
    total_pages = max(1, (len(items) + ROWS_PER_PAGE - 1) // ROWS_PER_PAGE)
    
    for page in range(total_pages):
        page_items = items[page * ROWS_PER_PAGE : (page + 1) * ROWS_PER_PAGE]
        page_sheet_id = f"sheet-tchr-{re.sub(r'[^a-zA-Z0-9]', '_', t_name)}-p{page+1}"
        
        display_style = "display: flex;" if t_idx == 0 else "display: none;"
        print_class = "" if t_idx == 0 else "hide-in-print"
        page_note = f" (Page {page+1} of {total_pages})" if total_pages > 1 else ""

        html_out.append(f'''
    <!-- TEACHER RECORD SHEET: {html.escape(t_name)}{page_note} -->
    <div class="page-sheet teacher-sheet dept-target-{t_cat} {print_class}" id="{page_sheet_id}" data-teacher="{html.escape(t_name)}" data-cat="{t_cat}" style="{display_style}">
      <div class="sheet-content">
        {build_header_html("TEACHER INSTRUCTIONAL ATTENDANCE & LOAD MONITORING RECORD", "Online Distance Learning (Second Shift) &bull; Official Monitoring Record &bull; SY 2026 - 2027")}

        <div class="meta-box">
          <div class="meta-row"><span class="meta-lbl">Teacher's Name:</span><span class="meta-val td-bold">{html.escape(t_name)}</span></div>
          <div class="meta-row"><span class="meta-lbl">Department / Faculty:</span><span class="meta-val">{html.escape(t_dept)}</span></div>
          <div class="meta-row"><span class="meta-lbl">Shift Modality:</span><span class="meta-val td-bold">ODL Second Shift</span></div>
          <div class="meta-row"><span class="meta-lbl">School Year:</span><span class="meta-val">2026 - 2027</span></div>
          <div class="meta-row"><span class="meta-lbl">Weekly Teaching Load:</span><span class="meta-val td-bold">{len(items)} Official Periods</span></div>
          <div class="meta-row"><span class="meta-lbl">Form Status:</span><span class="meta-val">Active Record{page_note}</span></div>
        </div>

        <table class="teacher-record-table">
          <thead>
            <tr>
              <th class="th-num">#</th>
              <th class="th-day">DAY</th>
              <th class="th-time">SCHEDULED TIME</th>
              <th class="th-mins">MINS</th>
              <th class="th-in">ACTUAL TIME IN</th>
              <th class="th-out">ACTUAL TIME OUT</th>
              <th class="th-grade">GRADE / SEC</th>
              <th class="th-subject">SUBJECT / LEARNING AREA</th>
              <th class="th-room">ROOM</th>
              <th class="th-status">INSTRUCTION STATUS</th>
              <th class="th-remarks">SIGN</th>
            </tr>
          </thead>
          <tbody>
        ''')

        for row_idx in range(ROWS_PER_PAGE):
            row_num = page * ROWS_PER_PAGE + row_idx + 1
            if row_idx < len(page_items):
                it = page_items[row_idx]
                html_out.append(f'''            <tr>
              <td class="td-center" style="font-weight:700; color:#64748b;">{row_num}</td>
              <td class="day-cell">{html.escape(it['day_abbr'])}</td>
              <td class="time-slot">{html.escape(it['time'])}</td>
              <td class="td-center" style="font-weight:700;">{html.escape(it['mins'])}</td>
              <td class="td-center">&nbsp;</td>
              <td class="td-center">&nbsp;</td>
              <td class="grade-cell">{html.escape(it['section'])}</td>
              <td class="subject-cell">{html.escape(it['subject'])}</td>
              <td class="td-center" style="font-size: 7.2pt; color: #475569;">{html.escape(it.get('room', 'ODL'))}</td>
              <td>
                <div class="status-grid">
                  <div class="status-cell-row">
                    <label class="status-chk"><input type="checkbox"> In</label>
                    <label class="status-chk"><input type="checkbox"> Late</label>
                  </div>
                  <div class="status-cell-row">
                    <label class="status-chk"><input type="checkbox"> Abs</label>
                    <label class="status-chk"><input type="checkbox"> Sub</label>
                  </div>
                </div>
              </td>
              <td class="td-center">&nbsp;</td>
            </tr>
''')
            else:
                html_out.append(f'''            <tr>
              <td class="td-center" style="font-weight:600; color:#cbd5e1;">{row_num}</td>
              <td class="td-center">&nbsp;</td>
              <td class="td-center">&nbsp;</td>
              <td class="td-center">&nbsp;</td>
              <td class="td-center">&nbsp;</td>
              <td class="td-center">&nbsp;</td>
              <td class="td-center">&nbsp;</td>
              <td class="td-center">&nbsp;</td>
              <td class="td-center">&nbsp;</td>
              <td>
                <div class="status-grid">
                  <div class="status-cell-row">
                    <label class="status-chk"><input type="checkbox"> In</label>
                    <label class="status-chk"><input type="checkbox"> Late</label>
                  </div>
                  <div class="status-cell-row">
                    <label class="status-chk"><input type="checkbox"> Abs</label>
                    <label class="status-chk"><input type="checkbox"> Sub</label>
                  </div>
                </div>
              </td>
              <td class="td-center">&nbsp;</td>
            </tr>
''')

        html_out.append(f'''          </tbody>
        </table>

        <!-- SIGNATURE BLOCK -->
        <div class="sign-row" style="margin-top: auto; padding-top: 6px; border-top: 1.5px solid #cbd5e1;">
          <div class="sign-col">
            <div class="sign-line"></div>
            <div class="sign-label">{html.escape(t_name)}</div>
            <div class="sign-title">Teacher's Signature over Printed Name</div>
          </div>
          <div class="sign-col">
            <div class="sign-line"></div>
            <div class="sign-label">Academic Coordinator / Department Head</div>
            <div class="sign-title">Verified & Monitored By</div>
          </div>
          <div class="sign-col">
            <div class="sign-line"></div>
            <div class="sign-label">School Principal / Directress</div>
            <div class="sign-title">Approved By</div>
          </div>
        </div>
      </div>
    </div>
''')

html_out.append('''  </div>

  <!-- MASTER DAILY ATTENDANCE MATRIX CONTAINER -->
  <div id="masterAttendanceContainer" class="attendance-view-container">
    <div style="display: flex; justify-content: space-between; align-items: center; margin-bottom: 16px;">
      <div>
        <h2 style="margin: 0; font-size: 16px; color: #064e3b; font-weight: 800; text-transform: uppercase;">
          Master Daily Attendance & Instructional Matrix
        </h2>
        <div style="font-size: 12px; color: #64748b; margin-top: 2px;">
          Online Distance Learning (Second Shift) &bull; SY 2026 - 2027
        </div>
      </div>
      <div>
        <button class="btn btn-primary" onclick="window.print()">Print Matrix</button>
      </div>
    </div>

    <!-- Day Selector Tabs -->
    <div class="day-matrix-tabs">
      <button class="day-matrix-tab active" data-day="Sunday" onclick="setMatrixDay('Sunday', this)">Sunday (SUN)</button>
      <button class="day-matrix-tab" data-day="Monday" onclick="setMatrixDay('Monday', this)">Monday (MON)</button>
      <button class="day-matrix-tab" data-day="Tuesday" onclick="setMatrixDay('Tuesday', this)">Tuesday (TUE)</button>
      <button class="day-matrix-tab" data-day="Wednesday" onclick="setMatrixDay('Wednesday', this)">Wednesday (WED)</button>
      <button class="day-matrix-tab" data-day="Thursday" onclick="setMatrixDay('Thursday', this)">Thursday (THU)</button>
    </div>

    <table class="attendance-table">
      <thead>
        <tr>
          <th style="width: 50px; text-align: center;">#</th>
          <th style="width: 140px;">Scheduled Time</th>
          <th style="width: 50px; text-align: center;">Mins</th>
          <th style="width: 200px;">Grade & Section</th>
          <th style="width: 180px;">Learning Area / Subject</th>
          <th style="width: 180px;">Assigned Teacher</th>
          <th style="width: 100px;">Room</th>
          <th style="width: 110px; text-align: center;">Status</th>
        </tr>
      </thead>
      <tbody id="matrixTableBody">
        <!-- Populated via JavaScript -->
      </tbody>
    </table>
  </div>

  <script>
    const AMIS_LOGO_B64 = ''' + json.dumps(amis_b64) + ''';
    const DEPED_LOGO_B64 = ''' + json.dumps(deped_b64) + ''';
    const MASTER_ATTENDANCE = ''' + json.dumps(all_master_attendance_rows) + ''';

    document.addEventListener("DOMContentLoaded", () => {
      document.querySelectorAll(".deped-img").forEach(img => img.src = DEPED_LOGO_B64);
      document.querySelectorAll(".amis-img").forEach(img => img.src = AMIS_LOGO_B64);
      renderMatrixTable();
    });

    let currentView = 'sections';
    let currentMatrixDay = 'Sunday';
    let currentTchrDept = 'all';

    function switchView(view) {
      currentView = view;
      const tabSections = document.getElementById("tabSections");
      const tabTeachers = document.getElementById("tabTeachers");
      const tabMaster = document.getElementById("tabMasterAttendance");

      const ctrlSections = document.getElementById("controlsSections");
      const ctrlTeachers = document.getElementById("controlsTeachers");
      const ctrlMaster = document.getElementById("controlsMaster");

      const secContainer = document.getElementById("sectionsContainer");
      const tchrContainer = document.getElementById("teachersContainer");
      const masterContainer = document.getElementById("masterAttendanceContainer");

      // Reset tabs
      tabSections.classList.remove("active");
      tabTeachers.classList.remove("active");
      tabMaster.classList.remove("active");

      // Hide all controls and containers
      ctrlSections.style.display = "none";
      ctrlTeachers.style.display = "none";
      ctrlMaster.style.display = "none";

      secContainer.style.display = "none";
      tchrContainer.style.display = "none";
      masterContainer.style.display = "none";

      if (view === 'sections') {
        tabSections.classList.add("active");
        ctrlSections.style.display = "grid";
        secContainer.style.display = "flex";
        applySectionFilter();
      } else if (view === 'teachers') {
        tabTeachers.classList.add("active");
        ctrlTeachers.style.display = "grid";
        tchrContainer.style.display = "flex";
        onTeacherSelectChange();
      } else if (view === 'master') {
        tabMaster.classList.add("active");
        ctrlMaster.style.display = "grid";
        masterContainer.style.display = "block";
        renderMatrixTable();
      }
    }

    function applySectionFilter() {
      const dept = document.getElementById("deptFilter").value;
      const sheets = document.querySelectorAll(".section-sheet");
      let visibleCount = 0;
      sheets.forEach(sheet => {
        const sDept = sheet.getAttribute("data-dept");
        if (dept === "all" || dept === sDept) {
          sheet.style.display = "flex";
          sheet.classList.remove("hide-in-print");
          visibleCount++;
        } else {
          sheet.style.display = "none";
          sheet.classList.add("hide-in-print");
        }
      });
      document.getElementById("sectionJump").value = "all";
      document.getElementById("resultCountInfo").textContent = `Displaying ${visibleCount} Section Monitoring Sheets (ODL Second Shift)`;
    }

    function jumpToSection() {
      const secId = document.getElementById("sectionJump").value;
      const sheets = document.querySelectorAll(".section-sheet");
      if (secId === "all") {
        applySectionFilter();
        return;
      }
      let visibleCount = 0;
      sheets.forEach(sheet => {
        if (sheet.getAttribute("data-id") === secId) {
          sheet.style.display = "flex";
          sheet.classList.remove("hide-in-print");
          sheet.scrollIntoView({ behavior: 'smooth', block: 'start' });
          visibleCount++;
        } else {
          sheet.style.display = "none";
          sheet.classList.add("hide-in-print");
        }
      });
      document.getElementById("resultCountInfo").textContent = `Displaying Section: ${secId.toUpperCase()}`;
    }

    function searchSectionClasses() {
      const query = document.getElementById("secSearch").value.toLowerCase().trim();
      const sheets = document.querySelectorAll(".section-sheet");
      let visibleCount = 0;
      sheets.forEach(sheet => {
        const text = sheet.innerText.toLowerCase();
        if (!query || text.includes(query)) {
          sheet.style.display = "flex";
          sheet.classList.remove("hide-in-print");
          visibleCount++;
        } else {
          sheet.style.display = "none";
          sheet.classList.add("hide-in-print");
        }
      });
      document.getElementById("resultCountInfo").textContent = `Found ${visibleCount} Section Monitoring Sheets matching "${query}"`;
    }

    function onTchrDeptChange() {
      const dept = document.getElementById("tchrDeptFilter").value;
      currentTchrDept = dept;
      const select = document.getElementById("teacherSelect");
      let firstVisible = null;

      for (let i = 0; i < select.options.length; i++) {
        const opt = select.options[i];
        const cat = opt.getAttribute("data-cat");
        if (dept === 'all' || cat === dept) {
          opt.style.display = "block";
          if (!firstVisible) firstVisible = opt.value;
        } else {
          opt.style.display = "none";
        }
      }

      if (firstVisible) {
        select.value = firstVisible;
        onTeacherSelectChange();
      }
    }

    function onTeacherSelectChange() {
      const selectedTeacher = document.getElementById("teacherSelect").value;
      const sheets = document.querySelectorAll(".teacher-sheet");
      let visibleCount = 0;

      sheets.forEach(sheet => {
        if (sheet.getAttribute("data-teacher") === selectedTeacher) {
          sheet.style.display = "flex";
          sheet.classList.remove("hide-in-print");
          visibleCount++;
        } else {
          sheet.style.display = "none";
          sheet.classList.add("hide-in-print");
        }
      });
      document.getElementById("resultCountInfo").textContent = `Displaying Record for: ${selectedTeacher}`;
    }

    function searchTeacherByName() {
      const query = document.getElementById("tchrSearch").value.toLowerCase().trim();
      const select = document.getElementById("teacherSelect");
      let firstMatch = null;
      for (let i = 0; i < select.options.length; i++) {
        const opt = select.options[i];
        const text = opt.text.toLowerCase();
        if (!query || text.includes(query)) {
          opt.style.display = "block";
          if (!firstMatch) firstMatch = opt.value;
        } else {
          opt.style.display = "none";
        }
      }
      if (firstMatch) {
        select.value = firstMatch;
        onTeacherSelectChange();
      }
    }

    function handlePrintAll() {
      if (currentView === 'sections') {
        document.getElementById("deptFilter").value = "all";
        document.getElementById("sectionJump").value = "all";
        document.getElementById("secSearch").value = "";
        applySectionFilter();
        window.print();
      } else if (currentView === 'teachers') {
        const sheets = document.querySelectorAll(".teacher-sheet");
        sheets.forEach(sheet => {
          const cat = sheet.getAttribute("data-cat");
          if (currentTchrDept === 'all' || cat === currentTchrDept) {
            sheet.style.display = "flex";
            sheet.classList.remove("hide-in-print");
          } else {
            sheet.style.display = "none";
            sheet.classList.add("hide-in-print");
          }
        });
        window.print();
      } else {
        window.print();
      }
    }

    function setMatrixDay(day, btn) {
      currentMatrixDay = day;
      document.querySelectorAll(".day-matrix-tab").forEach(b => b.classList.remove("active"));
      btn.classList.add("active");
      document.getElementById("matrixDaySelect").value = day;
      renderMatrixTable();
    }

    function renderMatrixTable() {
      const day = document.getElementById("matrixDaySelect").value || currentMatrixDay;
      const dept = document.getElementById("matrixDeptFilter").value;
      const query = document.getElementById("matrixSearch").value.toLowerCase().trim();

      const tbody = document.getElementById("matrixTableBody");
      tbody.innerHTML = "";

      const filtered = MASTER_ATTENDANCE.filter(r => {
        if (r.day !== day) return false;
        if (dept !== "all" && r.dept_code !== dept) return false;
        if (query) {
          const matchStr = `${r.section} ${r.subject} ${r.teacher} ${r.room}`.toLowerCase();
          if (!matchStr.includes(query)) return false;
        }
        return true;
      });

      filtered.sort((a, b) => a.sort_time - b.sort_time || a.section.localeCompare(b.section));

      if (filtered.length === 0) {
        tbody.innerHTML = `<tr><td colspan="8" style="text-align: center; padding: 20px; color: #94a3b8;">No scheduled instructional classes found matching the criteria.</td></tr>`;
        document.getElementById("resultCountInfo").textContent = `Displaying 0 classes for ${day}`;
        return;
      }

      filtered.forEach((row, i) => {
        const tr = document.createElement("tr");
        tr.innerHTML = `
          <td style="text-align: center; font-weight: 700; color: #64748b;">${i + 1}</td>
          <td style="font-weight: 700; color: #064e3b;">${row.time}</td>
          <td style="text-align: center;">${row.mins}m</td>
          <td style="font-weight: 600;">${row.section}</td>
          <td style="font-weight: 800; color: #065f46;">${row.subject.toUpperCase()}</td>
          <td style="font-weight: 700; color: #0284c7;">${row.teacher}</td>
          <td style="color: #475569;">${row.room}</td>
          <td style="text-align: center;">
            <span style="display: inline-block; padding: 2px 6px; border-radius: 4px; font-size: 10px; font-weight: 700; background: #ecfdf5; color: #065f46; border: 1px solid #a7f3d0;">SCHEDULED</span>
          </td>
        `;
        tbody.appendChild(tr);
      });

      document.getElementById("resultCountInfo").textContent = `Displaying ${filtered.length} classes for ${day} (ODL Second Shift)`;
    }
  </script>
</body>
</html>
''')

with open(OUTPUT_FILE, 'w', encoding='utf-8') as f:
    f.write(''.join(html_out))

with open(ROOT_OUTPUT_FILE, 'w', encoding='utf-8') as f:
    f.write(''.join(html_out))

print(f"Successfully generated: {OUTPUT_FILE} ({os.path.getsize(OUTPUT_FILE):,} bytes)")
print(f"Successfully mirrored: {ROOT_OUTPUT_FILE}")
