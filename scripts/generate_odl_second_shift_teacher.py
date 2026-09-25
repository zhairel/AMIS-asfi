import openpyxl, json, re, html, base64, os

print("=== Generating ODL Second Shift Per-Teacher Monitoring Portal ===")

REPO_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
DATA_FILE = os.path.join(REPO_DIR, 'data', 'schedule_odl_2nd_shift.xlsx')
AMIS_LOGO = os.path.join(REPO_DIR, 'public', 'amis_logo_opt.png')
DEPED_LOGO = os.path.join(REPO_DIR, 'public', 'deped_logo_opt.png')
OUTPUT_FILE = os.path.join(REPO_DIR, 'public', 'odl-second-shift-teacher.html')
ROOT_OUTPUT_FILE = os.path.join(REPO_DIR, '..', 'odl-second-shift-teacher.html')

wb = openpyxl.load_workbook(DATA_FILE, data_only=True)

with open(AMIS_LOGO, 'rb') as f:
    amis_b64 = "data:image/png;base64," + base64.b64encode(f.read()).decode('utf-8')

with open(DEPED_LOGO, 'rb') as f:
    deped_b64 = "data:image/png;base64," + base64.b64encode(f.read()).decode('utf-8')

days = ['Sunday', 'Monday', 'Tuesday', 'Wednesday', 'Thursday']
day_abbr_map = {
    'Sunday': 'SUN',
    'Monday': 'MON',
    'Tuesday': 'TUE',
    'Wednesday': 'WED',
    'Thursday': 'THU'
}

def get_grid(sheet):
    grid = {}
    for r in range(1, sheet.max_row + 1):
        for c in range(1, sheet.max_column + 1):
            val = sheet.cell(r, c).value
            grid[(r, c)] = str(val).strip() if val is not None else ''
            
    for rng in sheet.merged_cells.ranges:
        top_val = grid.get((rng.min_row, rng.min_col), '')
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                grid[(r, c)] = top_val
    return grid

elem_grid = get_grid(wb['ELEM'])
hs_new_grid = get_grid(wb['HS SCHED (NEW)'])
hs_grid = get_grid(wb['HS SCHED'])

def clean_time(t_str):
    if not t_str: return ''
    t = str(t_str).strip()
    if any(kw in t.upper() for kw in ['GRADE', 'USTADH', 'TEACHER', 'LEVEL', 'ADVISER']): return ''
    m_odl = re.search(r'([\d:apm\.\s-]+?)\s*(?:\(ODL\)|ODL)', t, re.IGNORECASE)
    if m_odl: t = m_odl.group(1).strip()
    else: t = re.sub(r'\(.*?\)', '', t).strip()
    t = re.sub(r'(\d{1,2}:\d{2}):+(\d{1,2}:\d{2})', r'\1 - \2', t)
    t = re.sub(r'(\d{1,2}:\d{2}):00', r'\1', t)
    is_pm = bool(re.search(r'(?:p\.?m\.?|pm)', t, re.IGNORECASE))
    is_am = bool(re.search(r'(?:a\.?m\.?|am)', t, re.IGNORECASE))
    clean = re.sub(r'(?:a\.?m\.?|p\.?m\.?|am|pm)', '', t, flags=re.IGNORECASE).strip()
    parts = re.split(r'\s*[-–]\s*', clean)
    if len(parts) == 2:
        start, end = parts[0].strip(), parts[1].strip()
        start = re.sub(r'^0(\d:)', r'\1', start)
        end = re.sub(r'^0(\d:)', r'\1', end)
        try:
            start_hour = int(start.split(':')[0]) if ':' in start else 0
            end_hour = int(end.split(':')[0]) if ':' in end else 0
        except ValueError:
            return t
        if is_pm:
            if start_hour == 11 and end_hour == 12: return f'{start} AM - {end} PM'
            return f'{start} - {end} PM'
        elif is_am: return f'{start} - {end} AM'
        else:
            if start_hour in [7, 8, 9, 10, 11] and end_hour in [7, 8, 9, 10, 11]: return f'{start} - {end} AM'
            elif start_hour == 11 and (end_hour == 12 or end_hour <= 1): return f'{start} AM - {end} PM'
            elif start_hour == 12 or start_hour in [1, 2, 3, 4, 5, 6]: return f'{start} - {end} PM'
            else: return f'{start} - {end} PM'
    elif len(parts) == 1:
        val = parts[0].strip()
        val = re.sub(r'^0(\d:)', r'\1', val)
        if is_pm or val.startswith(('12:', '1:', '2:', '3:', '4:', '5:')): return f'{val} PM'
        else: return f'{val} AM'
    return t

def parse_mins(m_str):
    if not m_str: return '40'
    m_str = str(m_str).strip()
    m = re.search(r'(\d+)', m_str)
    return str(m.group(1)) if m else '40'

KNOWN_TEACHERS = {
    'Saimona': 'Tchr. Saimona',
    'Jerlyn': 'Tchr. Jerlyn',
    'Keychell': 'Tchr. Keychell',
    'Faidh': 'Ustadh Faidh',
    'Saliha': 'Ustadha Saliha',
    'Obaydah': 'Ust. Obaydah',
    'Obayda': 'Ust. Obaydah',
    'Marham': 'Tchr. Marham',
    'Jenny': 'Tchr. Jenny',
    'Jessa': 'Tchr. Jessa',
    'Laurel': 'Ust. Laurel',
    'Laiurel': 'Ust. Laurel',
    'Silfah': 'Ustadha Silfa',
    'Silfa': 'Ustadha Silfa',
    'Ersahad': 'Ustadh Ersahad',
    'Hainur': 'Ustadh Hainur',
    'Bustamante': 'Alim Bustamante',
    'Wendy': 'Tchr. Wendy',
    'Sahdia': 'Tchr. Sahdia',
    'Norhydie': 'Tchr. Norhydie',
    'Joanna': 'Tchr. Joanna',
    'Katrina': 'Tchr. Katrina',
    'Sitti': 'Tchr. Sitti',
    'Monisa': 'Tchr. Monisa',
    'Zara': 'Tchr. Franchette',
    'Anna': 'Tchr. Anna',
    'Arvin': 'Tchr. Arvin',
    'Fhairudz': 'Tchr. Fhairudz',
    'Erica': 'Tchr. Erica',
    'Junaisah': 'Tchr. Junaisah',
    'Junaisa': 'Tchr. Junaisah',
    'Ayah': 'Tchr. Ayah',
    'Jayra': 'Tchr. Jayra',
    'Jairah': 'Tchr. Jayra',
    'Shi': 'Tchr. Shirehan',
    'Shirehan': 'Tchr. Shirehan',
    'Sophia': 'Tchr. Sophia',
    'Aniah': 'Tchr. Aniah',
    'Halnaisa': 'Tchr. Halnaisa',
    'Franchette': 'Tchr. Franchette',
    'Ethel': 'Tchr. Ethel',
    'Ali': 'Ustadh Muh Ali',
    'Radzmia': 'Tchr. Radzmia',
    'Wardah': 'Tchr. Wardah',
    'Jaisam': 'Ustadh Jaisam',
    'Rowena': 'Tchr. Rowena',
    'Norhaima': 'Tchr. Norhaima',
    'Jhelyn': 'Tchr. Jhelyn',
    'Nof': 'Tchr. Nof',
    'Angeleni': 'Tchr. Angeleni',
    'Nadzra': 'Tchr. Nadzra',
    'Samsuddin': 'Alim Samsuddin',
    'Abdulwahab': 'Alim Abdulwahab',
    'Mamonas': 'Alim Mamonas',
    'Moh': 'Sir Moh',
    'Mohaymen': 'Sir Moh',
    'Thea': 'Tchr. Thea',
    'Abegail': 'Tchr. Abegail',
    'Marie': 'Tchr. Marie',
    'Ahmad': 'Tchr. Ahmad',
    'Dipatuan': 'Alim Dipatuan',
    'Abdul Karim': 'Alim Abdul Karim',
    'Abdiraheem': 'Ust. Abdiraheem',
    'Abdi': 'Ust. Abdiraheem',
}

def normalize_teacher(name):
    name = re.sub(r'\s+', ' ', name.strip()).strip('-– ')
    if re.match(r'^Teacher\b', name, re.IGNORECASE): name = 'Tchr. ' + name[7:].strip()
    elif re.match(r'^Tr\.?\b', name, re.IGNORECASE): name = 'Tchr. ' + name[3:].strip()
    elif re.match(r'^Tchr\.?\b', name, re.IGNORECASE): name = 'Tchr. ' + re.sub(r'^Tchr\.?\s*', '', name, flags=re.IGNORECASE)
    elif re.match(r'^(?:Ustadha|Ustadza)\b', name, re.IGNORECASE): name = 'Ustadha ' + re.sub(r'^(?:Ustadha|Ustadza)\s*', '', name, flags=re.IGNORECASE)
    elif re.match(r'^Ustadh\b', name, re.IGNORECASE): name = 'Ustadh ' + name[6:].strip()
    elif re.match(r'^Ust\.?\b', name, re.IGNORECASE): name = 'Ust. ' + re.sub(r'^Ust\.?\s*', '', name, flags=re.IGNORECASE)
    name = name.strip()
    if name in ['Tchr.', 'Tchr', 'Ust.', 'Ust', 'Teacher', 'Tr.', 'Tr', '']: return 'TBA'
    if name.upper() in ['TCHR. AHMAD', 'SIR AHMAD']: return 'Tchr. Ahmad'
    if name in ['Ustadh Jaisam', 'Ust. Jaisam']: return 'Ustadh Jaisam'
    if name in ['Tchr. Kat', 'Tchr. Katrina']: return 'Tchr. Katrina'
    if name in ['Ust. Ubaydah', 'Ust. Obaydah', 'Ustadh Obayda', 'Ust. Obayda', 'Obaydah']: return 'Ust. Obaydah'
    if name in ['Ust. Silfah', 'Ustadha Silfa', 'Ust. Silfa', 'Ustadza Silfah', 'Ust. Silfah']: return 'Ustadha Silfa'
    if name in ['Ustadha Saliha', 'Ust. Saliha', 'Saliha']: return 'Ustadha Saliha'
    if name in ['Tchr. Junaisa', 'Tchr. Junaisah']: return 'Tchr. Junaisah'
    if name in ['Tchr. Jairah', 'Tchr. Jayra']: return 'Tchr. Jayra'
    if name in ['Tchr. Moh', 'Sir Mohaymen', 'Sir Moh']: return 'Sir Moh'
    if name in ['Tchr. Shi', 'Tchr. Shirehan']: return 'Tchr. Shirehan'
    if name in ['Tchr. Zara', 'Tchr. Franchette', 'Franchette']: return 'Tchr. Franchette'
    if name in ['Ust. Abdi', 'Ust. Abdiraheem', 'Ustadh Abdi', 'Ustadh Abdiraheem', 'Ustd. Abdi', 'Ustd. Abdiraheem']: return 'Ust. Abdiraheem'
    if name in ['Ust. Ali', 'Ustadh Ali', 'Ustadh Muh Ali', 'Ustdh ali', 'Ustdh. Ali', 'Ust. Muh Ali', 'Ustadh Muh. Ali']: return 'Ustadh Muh Ali'
    if name in ['Ust. Faidh', 'Ustadh Faidh', 'Faidh', 'Us. Faidh']: return 'Ustadh Faidh'
    if name in ['Ust. Ersahad', 'Ustadh Ersahad', 'Ersahad']: return 'Ustadh Ersahad'
    if name in ['Ust. Hainur', 'Ustadh Hainur', 'Hainur']: return 'Ustadh Hainur'
    if name in ['Alim Abdul Karim', 'Ust. Abdul Karim', 'Abdul Karim']: return 'Alim Abdul Karim'
    if name in ['Alim Samsuddin', 'Ust. Samsuddin', 'Samsuddin']: return 'Alim Samsuddin'
    if name in ['Alim Abdulwahab', 'Ust. Abdulwahab']: return 'Alim Abdulwahab'
    if name in ['Alim Mamonas', 'Ust. Mamonas']: return 'Alim Mamonas'
    if name in ['Alim Bustamante', 'Ust. Bustamante']: return 'Alim Bustamante'
    if name in ['Laiurel', 'Ust. Laurel', 'Laurel']: return 'Ust. Laurel'
    if name in ['Tr Nof', 'Tchr. Nof', 'Nof']: return 'Tchr. Nof'
    return name

def clean_parse(text):
    text = text.strip()
    if not text: return '', '', 'EMPTY'
    upper = text.upper()
    for r in ['GENERAL ASSEMBLY', 'LUNCH AND SALAH', 'SALAH & DEPARTURE', 'DEPARTURE', 'RECESS', 'SHORT BREAK', 'TRANSITION', 'BREAK', 'HOMEROOM', 'ENTRANCE EXAM REVIEW', 'RESEARCH CONSULTATION']:
        if upper == r or (upper.startswith(r) and not any(k in upper for k in ['TCHR', 'UST', 'SIR', 'ALIM', 'TEACHER', 'TR'])):
            return text, '', 'ROUTINE'
    t_match = re.search(r'[-–]?\s*(Tchr\.?|Ust\.?|Sir|Alim|Teacher|Ustadha|Ustadza|Ustadh|Tr\.?)\s+(.*)$', text, re.IGNORECASE)
    if t_match:
        tchr_full = normalize_teacher((t_match.group(1) + ' ' + t_match.group(2)).strip())
        subj_part = text[:t_match.start()].strip().rstrip('-–').strip()
        return subj_part or text, tchr_full, 'CLASS'
    for k, v in KNOWN_TEACHERS.items():
        m = re.search(r'[-–]?\s*' + re.escape(k) + r'\b', text, re.IGNORECASE)
        if m:
            subj = text[:m.start()].strip().rstrip('-–').strip()
            subj = re.sub(r'([A-Za-z]+)(\d+)$', r'\1 \2', subj)
            return subj or text, v, 'CLASS'
    for p in ['Tchr.', 'Teacher', 'Sir', 'Ustadh', 'Ustadha', 'Ustadza', 'Ust.', 'Alim', 'Tr.']:
        if p.lower() in text.lower():
            idx = text.lower().find(p.lower())
            tchr_full = normalize_teacher(text[idx:].strip())
            subj_part = text[:idx].strip().rstrip('-–').strip()
            return subj_part or text, tchr_full, 'CLASS'
    return text, 'TBA', 'CLASS'

def time_to_sort_key(t_str):
    m = re.search(r'(\d+):(\d+)', t_str)
    if not m: return 9999
    h, mins = int(m.group(1)), int(m.group(2))
    if 'PM' in t_str.upper():
        if h < 12: h += 12
    elif 'AM' in t_str.upper():
        if h == 12: h = 0
    else:
        if h in [1, 2, 3, 4, 5, 6, 12]: h = (h % 12) + 12
    return h * 60 + mins

day_order = {'Sunday': 0, 'Monday': 1, 'Tuesday': 2, 'Wednesday': 3, 'Thursday': 4}

sections = [
    # Kindergarten & Elementary (20 sections)
    {'id': 'k2_abdullah', 'sec_short': 'K2 - Abdullah', 'dept': 'Elementary', 'dept_code': 'elem', 'sheet': 'ELEM', 'grid': elem_grid, 'c_start': 1, 'r_start': 5, 'r_end': 11, 'room': 'ODL Room 1'},
    {'id': 'k2_umar', 'sec_short': 'K2 - Umar', 'dept': 'Elementary', 'dept_code': 'elem', 'sheet': 'ELEM', 'grid': elem_grid, 'c_start': 9, 'r_start': 5, 'r_end': 11, 'room': 'ODL Room 2'},
    {'id': 'k1_husain', 'sec_short': 'K1 - Husain', 'dept': 'Elementary', 'dept_code': 'elem', 'sheet': 'ELEM', 'grid': elem_grid, 'c_start': 1, 'r_start': 15, 'r_end': 21, 'room': 'ODL Room 3'},
    {'id': 'k2_khabaab', 'sec_short': 'K2 - Khabaab', 'dept': 'Elementary', 'dept_code': 'elem', 'sheet': 'ELEM', 'grid': elem_grid, 'c_start': 9, 'r_start': 15, 'r_end': 21, 'room': 'ODL Room 4'},
    {'id': 'g1_suhayb', 'sec_short': 'G1 - Suhayb', 'dept': 'Elementary', 'dept_code': 'elem', 'sheet': 'ELEM', 'grid': elem_grid, 'c_start': 1, 'r_start': 26, 'r_end': 32, 'room': 'ODL Room 5'},
    {'id': 'g1_saad', 'sec_short': "G1 - Sa'ad", 'dept': 'Elementary', 'dept_code': 'elem', 'sheet': 'ELEM', 'grid': elem_grid, 'c_start': 9, 'r_start': 26, 'r_end': 32, 'room': 'ODL Room 6'},
    {'id': 'g2_saeed', 'sec_short': 'G2 - Saeed', 'dept': 'Elementary', 'dept_code': 'elem', 'sheet': 'ELEM', 'grid': elem_grid, 'c_start': 1, 'r_start': 41, 'r_end': 47, 'room': 'ODL Room 7'},
    {'id': 'g2_aasim', 'sec_short': 'G2 - Aasim', 'dept': 'Elementary', 'dept_code': 'elem', 'sheet': 'ELEM', 'grid': elem_grid, 'c_start': 9, 'r_start': 41, 'r_end': 47, 'room': 'ODL Room 8'},
    {'id': 'g3_zayd', 'sec_short': 'G3 - Zayd (Girls)', 'dept': 'Elementary', 'dept_code': 'elem', 'sheet': 'ELEM', 'grid': elem_grid, 'c_start': 1, 'r_start': 56, 'r_end': 62, 'room': 'ODL Room 9'},
    {'id': 'g3_thabit', 'sec_short': 'G3 - Thabit (Boys)', 'dept': 'Elementary', 'dept_code': 'elem', 'sheet': 'ELEM', 'grid': elem_grid, 'c_start': 9, 'r_start': 56, 'r_end': 62, 'room': 'ODL Room 10'},
    {'id': 'g3_asad', 'sec_short': "G3 - As'ad (Mix)", 'dept': 'Elementary', 'dept_code': 'elem', 'sheet': 'ELEM', 'grid': elem_grid, 'c_start': 9, 'r_start': 66, 'r_end': 72, 'room': 'ODL Room 11'},
    {'id': 'g4_zubair', 'sec_short': 'G4 - Az Zubair', 'dept': 'Elementary', 'dept_code': 'elem', 'sheet': 'ELEM', 'grid': elem_grid, 'c_start': 1, 'r_start': 76, 'r_end': 82, 'room': 'ODL Room 12'},
    {'id': 'g4_ikrimah', 'sec_short': 'G4 - Ikrimah', 'dept': 'Elementary', 'dept_code': 'elem', 'sheet': 'ELEM', 'grid': elem_grid, 'c_start': 9, 'r_start': 76, 'r_end': 82, 'room': 'ODL Room 13'},
    {'id': 'g4_hassan', 'sec_short': 'G4 - Hassan (Mix)', 'dept': 'Elementary', 'dept_code': 'elem', 'sheet': 'ELEM', 'grid': elem_grid, 'c_start': 17, 'r_start': 76, 'r_end': 82, 'room': 'ODL Room 14'},
    {'id': 'g5_ayyash', 'sec_short': 'G5 - Ayyash', 'dept': 'Elementary', 'dept_code': 'elem', 'sheet': 'ELEM', 'grid': elem_grid, 'c_start': 1, 'r_start': 91, 'r_end': 97, 'room': 'ODL Room 15'},
    {'id': 'g5_musab', 'sec_short': "G5 - Mus'ab", 'dept': 'Elementary', 'dept_code': 'elem', 'sheet': 'ELEM', 'grid': elem_grid, 'c_start': 9, 'r_start': 91, 'r_end': 97, 'room': 'ODL Room 16'},
    {'id': 'g5_harith', 'sec_short': 'G5 - Al Harith', 'dept': 'Elementary', 'dept_code': 'elem', 'sheet': 'ELEM', 'grid': elem_grid, 'c_start': 17, 'r_start': 91, 'r_end': 97, 'room': 'ODL Room 17'},
    {'id': 'g5_jafar', 'sec_short': "G5 - Ja'far (Mix)", 'dept': 'Elementary', 'dept_code': 'elem', 'sheet': 'ELEM', 'grid': elem_grid, 'c_start': 25, 'r_start': 91, 'r_end': 97, 'room': 'ODL Room 18'},
    {'id': 'g6_khaleed', 'sec_short': 'G6 - Khaleed', 'dept': 'Elementary', 'dept_code': 'elem', 'sheet': 'ELEM', 'grid': elem_grid, 'c_start': 1, 'r_start': 106, 'r_end': 112, 'room': 'ODL Room 19'},
    {'id': 'g6_dihya', 'sec_short': 'G6 - Dihya (Girls)', 'dept': 'Elementary', 'dept_code': 'elem', 'sheet': 'ELEM', 'grid': elem_grid, 'c_start': 9, 'r_start': 106, 'r_end': 112, 'room': 'ODL Room 20'},

    # Junior High School (6 sections)
    {'id': 'g7_anas', 'sec_short': 'G7 - Anas (Mix)', 'dept': 'Junior High School', 'dept_code': 'jhs', 'sheet': 'HS SCHED (NEW)', 'grid': hs_new_grid, 'c_start': 1, 'r_start': 7, 'r_end': 13, 'room': 'ODL JHS Room 1'},
    {'id': 'g8_muadh', 'sec_short': "G8 - Mu'adh (Boys)", 'dept': 'Junior High School', 'dept_code': 'jhs', 'sheet': 'HS SCHED (NEW)', 'grid': hs_new_grid, 'c_start': 1, 'r_start': 18, 'r_end': 24, 'room': 'ODL JHS Room 2'},
    {'id': 'g8_nuaym', 'sec_short': 'G8 - Nuaym (Mix)', 'dept': 'Junior High School', 'dept_code': 'jhs', 'sheet': 'HS SCHED (NEW)', 'grid': hs_new_grid, 'c_start': 1, 'r_start': 30, 'r_end': 36, 'room': 'ODL JHS Room 3'},
    {'id': 'g9_abudharr', 'sec_short': 'G9 - Abu Dharr (Boys)', 'dept': 'Junior High School', 'dept_code': 'jhs', 'sheet': 'HS SCHED (NEW)', 'grid': hs_new_grid, 'c_start': 1, 'r_start': 42, 'r_end': 48, 'room': 'ODL JHS Room 4'},
    {'id': 'g9_abujandal', 'sec_short': 'G9 - Abu Jandal (Girls)', 'dept': 'Junior High School', 'dept_code': 'jhs', 'sheet': 'HS SCHED (NEW)', 'grid': hs_new_grid, 'c_start': 1, 'r_start': 53, 'r_end': 59, 'room': 'ODL JHS Room 5'},
    {'id': 'g10_abuayyub', 'sec_short': 'G10 - Abu Ayyub (Boys)', 'dept': 'Junior High School', 'dept_code': 'jhs', 'sheet': 'HS SCHED (NEW)', 'grid': hs_new_grid, 'c_start': 1, 'r_start': 64, 'r_end': 70, 'room': 'ODL JHS Room 6'},

    # Senior High School (3 sections)
    {'id': 'g11_b_term2', 'sec_short': 'G11 Boys (Term 2)', 'dept': 'Senior High School', 'dept_code': 'shs', 'sheet': 'HS SCHED (NEW)', 'grid': hs_new_grid, 'c_start': 1, 'r_start': 78, 'r_end': 85, 'room': 'ODL SHS Room 1'},
    {'id': 'g11_b_sem1', 'sec_short': 'G11 Boys (Sem 1)', 'dept': 'Senior High School', 'dept_code': 'shs', 'sheet': 'HS SCHED', 'grid': hs_grid, 'c_start': 26, 'r_start': 67, 'r_end': 73, 'room': 'ODL SHS Room 1'},
    {'id': 'g12_b_sem1', 'sec_short': 'G12 Suhayb (Boys)', 'dept': 'Senior High School', 'dept_code': 'shs', 'sheet': 'HS SCHED', 'grid': hs_grid, 'c_start': 18, 'r_start': 83, 'r_end': 89, 'room': 'ODL SHS Room 2'},
]

teacher_monitor_map = {}

for sec in sections:
    grid = sec['grid']
    c_start = sec['c_start']
    for r in range(sec['r_start'], sec['r_end'] + 1):
        raw_t = grid.get((r, c_start), '').strip()
        t_slot = clean_time(raw_t)
        m_slot = parse_mins(grid.get((r, c_start + 1), ''))
        if not t_slot: continue
        for d_idx, day in enumerate(days):
            c_val = grid.get((r, c_start + 2 + d_idx), '').strip()
            if not c_val: continue
            subj, tchr, kind = clean_parse(c_val)
            if kind == 'CLASS' and tchr and tchr != 'TBA':
                if tchr not in teacher_monitor_map:
                    teacher_monitor_map[tchr] = {
                        'name': tchr,
                        'dept': sec['dept'],
                        'dept_code': sec['dept_code'],
                        'items': []
                    }
                teacher_monitor_map[tchr]['items'].append({
                    'day': day,
                    'day_abbr': day_abbr_map.get(day, day[:3].upper()),
                    'day_idx': day_order.get(day, 0),
                    'time': t_slot,
                    'sort_time': time_to_sort_key(t_slot),
                    'mins': m_slot,
                    'section': sec['sec_short'],
                    'subject': subj.upper(),
                    'room': sec['room']
                })

for t_name, t_data in teacher_monitor_map.items():
    if any(t_name.startswith(p) for p in ['Ust.', 'Ustadh', 'Ustadha', 'Alim']):
        t_data['cat'] = 'isal'
        t_data['dept'] = 'ISAL & Islamic Studies Department'
    elif any('G11' in it['section'] or 'G12' in it['section'] for it in t_data['items']):
        t_data['cat'] = 'shs'
        t_data['dept'] = 'Senior High School Faculty'
    elif any('G7' in it['section'] or 'G8' in it['section'] or 'G9' in it['section'] or 'G10' in it['section'] for it in t_data['items']):
        t_data['cat'] = 'jhs'
        t_data['dept'] = 'Junior High School Faculty'
    else:
        t_data['cat'] = 'elem'
        t_data['dept'] = 'Elementary Faculty'
    
    t_data['items'].sort(key=lambda x: (x['day_idx'], x['sort_time'], x['section']))

teachers_list = sorted(teacher_monitor_map.values(), key=lambda x: x['name'])
print(f"Extracted {len(teachers_list)} teachers for ODL 2nd Shift.")

html_out = []
html_out.append('''<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <link rel="preconnect" href="https://fonts.googleapis.com">
  <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
  <link href="https://fonts.googleapis.com/css2?family=Amiri:wght@700&family=Inter:wght@400;500;600;700;800;900&display=swap" rel="stylesheet">
  <title>Per-Teacher Instructional Attendance & Load Monitoring Record (ODL Second Shift) - Al Munawwara Islamic School</title>
  <style>
    * {
      box-sizing: border-box;
      margin: 0;
      padding: 0;
    }

    body {
      font-family: 'Inter', -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, sans-serif;
      background: #f1f5f9;
      color: #0f172a;
      -webkit-font-smoothing: antialiased;
      line-height: 1.3;
    }

    /* TOOLBAR */
    .toolbar {
      background: #ffffff;
      color: #0f172a;
      padding: 10px 20px;
      position: sticky;
      top: 0;
      z-index: 999;
      box-shadow: 0 2px 10px rgba(0,0,0,0.06);
      border-bottom: 2px solid #059669;
    }
    .modality-bar {
      display: flex;
      align-items: center;
      gap: 8px;
      margin-bottom: 8px;
      padding-bottom: 8px;
      border-bottom: 1px solid #e2e8f0;
      flex-wrap: wrap;
    }
    .modality-label {
      font-size: 11px;
      font-weight: 800;
      color: #64748b;
      text-transform: uppercase;
      letter-spacing: 0.4px;
      margin-right: 4px;
    }
    .modality-pill {
      display: inline-flex;
      align-items: center;
      gap: 6px;
      padding: 4px 10px;
      border-radius: 6px;
      font-size: 11px;
      font-weight: 700;
      text-decoration: none;
      color: #475569;
      background: #f8fafc;
      border: 1px solid #cbd5e1;
      transition: all 0.15s ease;
    }
    .modality-pill:hover {
      background: #f1f5f9;
      border-color: #94a3b8;
      color: #0f172a;
    }
    .modality-pill.active {
      background: #ecfdf5;
      color: #065f46;
      border-color: #10b981;
    }
    .modality-dot {
      width: 7px;
      height: 7px;
      border-radius: 50%;
      background: #94a3b8;
    }
    .dot-active {
      background: #10b981;
      box-shadow: 0 0 0 2px #d1fae5;
    }
    .dot-dev {
      background: #f59e0b;
    }
    .badge-dev {
      background: #fef3c7;
      color: #92400e;
      border: 1px solid #fde68a;
      font-size: 9px;
      padding: 1px 5px;
      border-radius: 9999px;
      font-weight: 800;
      letter-spacing: 0.2px;
      text-transform: uppercase;
    }

    .toolbar-header {
      display: flex;
      justify-content: space-between;
      align-items: center;
      flex-wrap: wrap;
      gap: 10px;
      margin-bottom: 8px;
    }
    .toolbar-brand {
      display: flex;
      align-items: center;
      gap: 10px;
    }
    .toolbar-logo {
      width: 38px;
      height: 38px;
      object-fit: contain;
    }
    .toolbar-title {
      font-size: 15px;
      font-weight: 800;
      color: #064e3b;
      letter-spacing: 0.3px;
      display: flex;
      align-items: center;
      gap: 8px;
    }
    .badge-f2f {
      background: #d1fae5;
      color: #065f46;
      border: 1px solid #10b981;
      font-size: 10px;
      padding: 2px 7px;
      border-radius: 9999px;
      font-weight: 800;
      text-transform: uppercase;
    }
    .toolbar-actions {
      display: flex;
      gap: 6px;
      align-items: center;
      flex-wrap: wrap;
    }
    .btn {
      padding: 6px 12px;
      font-size: 11.5px;
      font-weight: 700;
      border-radius: 6px;
      cursor: pointer;
      display: inline-flex;
      align-items: center;
      gap: 5px;
      border: 1px solid transparent;
      transition: all 0.15s ease;
      text-decoration: none;
    }
    .btn-primary { background: #059669; color: #ffffff; }
    .btn-primary:hover { background: #047857; }
    .btn-blue { background: #0284c7; color: #ffffff; }
    .btn-blue:hover { background: #0369a1; }
    .btn-outline { background: #ffffff; color: #334155; border-color: #cbd5e1; }
    .btn-outline:hover { background: #f8fafc; color: #0f172a; border-color: #94a3b8; }

    /* FILTER BAR */
    .filter-bar {
      display: flex;
      align-items: center;
      gap: 8px;
      flex-wrap: wrap;
      background: #f8fafc;
      padding: 6px 12px;
      border-radius: 6px;
      border: 1px solid #e2e8f0;
    }
    .filter-label {
      font-size: 11px;
      font-weight: 800;
      color: #475569;
      text-transform: uppercase;
      letter-spacing: 0.3px;
    }
    .dept-pill {
      background: #ffffff;
      color: #334155;
      border: 1px solid #cbd5e1;
      padding: 4px 10px;
      border-radius: 5px;
      font-size: 11px;
      font-weight: 600;
      cursor: pointer;
      transition: all 0.15s ease;
    }
    .dept-pill:hover {
      background: #f1f5f9;
      color: #0f172a;
    }
    .dept-pill.active {
      background: #059669;
      color: #ffffff;
      border-color: #059669;
      font-weight: 700;
    }
    .teacher-select-box {
      margin-left: auto;
      display: flex;
      align-items: center;
      gap: 6px;
    }
    .teacher-select-box select {
      background: #ffffff;
      border: 1.5px solid #059669;
      color: #064e3b;
      padding: 4px 8px;
      border-radius: 6px;
      font-size: 11.5px;
      font-weight: 700;
      outline: none;
      min-width: 260px;
    }

    /* SHEET LAYOUT */
    .sheet-wrapper {
      padding: 24px 20px;
      display: flex;
      flex-direction: column;
      align-items: center;
      gap: 24px;
    }
    .page-sheet {
      background: #ffffff;
      width: 210mm;
      min-height: 297mm;
      box-shadow: 0 4px 15px rgba(0,0,0,0.06);
      border-radius: 2px;
      padding: 7mm 9mm 6mm 9mm;
      display: flex;
      flex-direction: column;
      position: relative;
    }
    .page-sheet.hidden-sheet {
      display: none !important;
    }
    .sheet-content {
      display: flex;
      flex-direction: column;
      flex: 1;
    }

    /* HEADER */
    .sheet-header {
      display: flex;
      align-items: center;
      justify-content: space-between;
      border-bottom: 2px solid #064e3b;
      padding-bottom: 4px;
      margin-bottom: 4px;
    }
    .header-logo-side {
      width: 52px;
      height: 52px;
      display: flex;
      align-items: center;
      justify-content: center;
    }
    .header-logo {
      max-width: 100%;
      max-height: 100%;
      object-fit: contain;
    }
    .header-center-text {
      text-align: center;
      flex: 1;
      padding: 0 8px;
    }
    .arabic-header {
      font-family: 'Amiri', serif;
      font-size: 11.5pt;
      color: #064e3b;
      font-weight: 700;
      line-height: 1.1;
      direction: rtl;
    }
    .school-name {
      font-size: 9.5pt;
      font-weight: 900;
      color: #064e3b;
      letter-spacing: 0.8px;
    }
    .form-title {
      font-size: 8.5pt;
      font-weight: 800;
      color: #0f172a;
      text-transform: uppercase;
      letter-spacing: 0.4px;
      margin-top: 1px;
    }
    .form-sub {
      font-size: 7.2pt;
      color: #475569;
      font-weight: 600;
    }

    /* META BOX */
    .meta-box {
      border: 1px solid #94a3b8;
      background: #f8fafc;
      border-radius: 3px;
      padding: 4px 8px;
      margin-bottom: 5px;
      display: grid;
      grid-template-columns: 1.4fr 1.2fr 1fr;
      gap: 3px 12px;
      font-size: 7.5pt;
    }
    .meta-row {
      display: flex;
      align-items: center;
      gap: 4px;
    }
    .meta-lbl {
      font-weight: 700;
      color: #334155;
      white-space: nowrap;
    }
    .meta-val {
      font-weight: 600;
      color: #0f172a;
      border-bottom: 1px dotted #94a3b8;
      flex: 1;
      min-height: 12px;
    }
    .td-bold {
      font-weight: 800;
      color: #064e3b;
    }

    /* TABLE */
    table.sheet-table {
      width: 100%;
      border-collapse: collapse;
      font-size: 7.4pt;
      line-height: 1.15;
    }
    table.sheet-table th, table.sheet-table td {
      border: 1px solid #64748b;
      padding: 3px 4px;
    }
    table.sheet-table th {
      background: #f1f5f9;
      color: #0f172a;
      font-weight: 800;
      font-size: 7.2pt;
      padding: 4px 3px;
      text-transform: uppercase;
      letter-spacing: 0.2px;
      text-align: center;
    }
    .th-num { width: 3%; }
    .th-day { width: 6.5%; }
    .th-time { width: 16.5%; }
    .th-mins { width: 4.5%; }
    .th-in { width: 8%; }
    .th-out { width: 8%; }
    .th-grade { width: 13.5%; }
    .th-subject { width: 18.5%; }
    .th-room { width: 5%; }
    .th-status { width: 11.5%; }
    .th-remarks { width: 5%; }
    .td-center { text-align: center; }

    .day-cell {
      font-weight: 800;
      color: #064e3b;
      text-align: center;
      text-transform: uppercase;
      letter-spacing: 0.3px;
      font-size: 7.8pt;
      white-space: nowrap;
    }
    .time-slot {
      font-weight: 800;
      color: #0f172a;
      text-align: center;
      white-space: nowrap;
      font-size: 7.8pt;
      letter-spacing: 0.2px;
      font-variant-numeric: tabular-nums;
    }
    .subject-cell {
      font-weight: 900;
      color: #064e3b;
      text-transform: uppercase;
      letter-spacing: 0.3px;
      font-size: 8pt;
    }
    .grade-cell {
      font-weight: 700;
      color: #1e293b;
      font-size: 7.8pt;
      white-space: nowrap;
    }

    /* 2x2 STATUS GRID CHECKLIST */
    .status-grid {
      display: flex;
      flex-direction: column;
      gap: 1.5px;
      font-size: 6.8pt;
      font-weight: 700;
    }
    .status-cell-row {
      display: flex;
      justify-content: space-between;
      gap: 3px;
    }
    .status-chk {
      display: inline-flex;
      align-items: center;
      gap: 2px;
      cursor: pointer;
      user-select: none;
      white-space: nowrap;
    }
    .status-chk input {
      margin: 0;
      width: 10px;
      height: 10px;
      cursor: pointer;
    }

    /* SIGNATURE BLOCK */
    .sign-row {
      display: grid;
      grid-template-columns: repeat(3, 1fr);
      gap: 16px;
      font-size: 7.5pt;
    }
    .sign-col {
      text-align: center;
    }
    .sign-line {
      border-bottom: 1.5px solid #0f172a;
      height: 24px;
      margin-bottom: 4px;
    }
    .sign-label {
      font-weight: 800;
      color: #1e293b;
      text-transform: uppercase;
      font-size: 7pt;
      letter-spacing: 0.3px;
    }
    .sign-title {
      font-size: 6.8pt;
      color: #64748b;
      font-weight: 600;
    }

    @media print {
      body {
        background: #ffffff !important;
        margin: 0;
        padding: 0;
      }
      .toolbar, .no-print {
        display: none !important;
      }
      .sheet-wrapper {
        padding: 0 !important;
        gap: 0 !important;
      }
      .page-sheet {
        box-shadow: none !important;
        border: none !important;
        margin: 0 !important;
        padding: 4mm 6mm !important;
        page-break-after: always !important;
        break-after: page !important;
        width: 100% !important;
        height: 297mm !important;
      }
      .page-sheet.hidden-sheet {
        display: none !important;
      }
    }
  </style>
</head>
<body>

  <!-- TOOLBAR -->
  <div class="toolbar">
    <!-- MODALITY NAVIGATION -->
    <div class="modality-bar">
      <span class="modality-label">Learning Modality:</span>
      <a href="/teacher-monitoring.html" class="modality-pill" title="Face-to-Face Portal">
        <span class="modality-dot"></span>
        Face-to-Face (45 Teachers)
      </a>
      <a href="/odl-teacher-monitoring.html" class="modality-pill" title="Online Distance Learning First Shift Portal">
        <span class="modality-dot"></span>
        ODL First Shift (53 Teachers)
      </a>
      <a href="/odl-second-shift-teacher.html" class="modality-pill active" title="Online Distance Learning Second Shift Portal">
        <span class="modality-dot dot-active"></span>
        ODL Second Shift (54 Teachers)
      </a>
      <a href="/all-teachers-monitoring.html" class="modality-pill" title="Unified Multi-Modality Portal" style="border-color:#3b82f6;color:#2563eb;background:#eff6ff;">
        <span class="modality-dot" style="background:#2563eb;"></span>
        All Modalities (152 Teachers)
      </a>
    </div>

    <div class="toolbar-header">
      <div class="toolbar-brand">
        <img class="toolbar-logo amis-img" alt="AMIS Logo">
        <div>
          <div class="toolbar-title">
            AL MUNAWWARA ISLAMIC SCHOOL
            <span class="badge-f2f">ODL 2nd Shift Teacher Portal</span>
          </div>
          <div style="font-size:11.5px;color:#475569;font-weight:500;">
            Daily Instructional Attendance & Load Monitoring Record (Online Distance Learning &bull; 2nd Shift &bull; S.Y. 2026 - 2027)
          </div>
        </div>
      </div>
      <div class="toolbar-actions">
        <a href="/odl-second-shift.html" class="btn btn-outline" title="Return to Classroom Section Monitoring">
          Classroom Monitoring Forms
        </a>
        <button class="btn btn-outline" onclick="navigateTeacher(-1)" title="Previous Teacher">
          Prev
        </button>
        <button class="btn btn-outline" onclick="navigateTeacher(1)" title="Next Teacher">
          Next
        </button>
        <button class="btn btn-blue" onclick="printActiveTeacher()">
          Print Active Teacher
        </button>
        <button class="btn btn-primary" onclick="printAllTeachers()">
          Print All Teachers (A4)
        </button>
      </div>
    </div>

    <!-- FILTER BAR -->
    <div class="filter-bar">
      <span class="filter-label">Filter Faculty:</span>
      <button class="dept-pill active" onclick="filterDept('all')">All Faculty (''' + str(len(teachers_list)) + ''')</button>
      <button class="dept-pill" onclick="filterDept('jhs')">Junior High School</button>
      <button class="dept-pill" onclick="filterDept('shs')">Senior High School</button>
      <button class="dept-pill" onclick="filterDept('elem')">Elementary Faculty</button>
      <button class="dept-pill" onclick="filterDept('isal')">ISAL Department</button>

      <div class="teacher-select-box">
        <label for="teacher-select" class="filter-label">Select Teacher:</label>
        <select id="teacher-select" onchange="onTeacherSelectChange()">
          <option value="all">All Faculty (Show All ''' + str(len(teachers_list)) + ''' Teachers)</option>
''')

for idx, t in enumerate(teachers_list):
    html_out.append(f'          <option value="{idx}">{html.escape(t["name"])} ({len(t["items"])} ODL classes)</option>\n')

html_out.append('''        </select>
      </div>
    </div>
  </div>

  <!-- SHEETS CONTAINER -->
  <div class="sheet-wrapper" id="sheet-wrapper">
''')

for idx, t in enumerate(teachers_list):
    t_name_esc = html.escape(t['name'])
    dept_esc = html.escape(t['dept'])
    cat_esc = html.escape(t['cat'])
    items = t['items']
    items.sort(key=lambda x: (days.index(x['day']), time_to_sort_key(x['time'])))

    hidden_cls = "" if idx == 0 else "hidden-sheet"

    html_out.append(f'''
    <!-- SHEET {idx}: {t_name_esc} -->
    <div class="page-sheet {hidden_cls}" 
         id="sheet-{idx}"
         data-idx="{idx}"
         data-teacher="{t_name_esc}"
         data-cat="{cat_esc}">
      
      <div class="sheet-content">
        <!-- HEADER -->
        <div class="sheet-header">
          <div class="header-logo-side">
            <img class="header-logo deped-img" alt="DepEd Logo">
          </div>
          <div class="header-center-text">
            <div class="arabic-header" dir="rtl" lang="ar">المدرسة المنورة الإسلامية</div>
            <div class="school-name">AL MUNAWWARA ISLAMIC SCHOOL</div>
            <div class="form-title">TEACHER INSTRUCTIONAL ATTENDANCE & LOAD MONITORING RECORD</div>
            <div class="form-sub">Online Distance Learning (Second Shift) Modality &bull; Faculty Monitoring Form &bull; School Year 2026 - 2027</div>
          </div>
          <div class="header-logo-side">
            <img class="header-logo amis-img" alt="AMIS Logo">
          </div>
        </div>

        <!-- TEACHER META BOX -->
        <div class="meta-box">
          <div class="meta-row"><span class="meta-lbl">Teacher's Name:</span><span class="meta-val td-bold">{t_name_esc}</span></div>
          <div class="meta-row"><span class="meta-lbl">Total Weekly Loads:</span><span class="meta-val td-bold">{len(items)} ODL Classes</span></div>
          <div class="meta-row"><span class="meta-lbl">Room Assignment:</span><span class="meta-val">Virtual / Google Meet</span></div>
          <div class="meta-row" style="grid-column: span 2;"><span class="meta-lbl">Week Monitored:</span><span class="meta-val"></span></div>
          <div class="meta-row"><span class="meta-lbl">School Year:</span><span class="meta-val">2026 - 2027</span></div>
        </div>

        <!-- TABLE -->
        <table class="sheet-table">
          <thead>
            <tr>
              <th class="th-num">#</th>
              <th class="th-day">Day</th>
              <th class="th-time">Scheduled Time</th>
              <th class="th-mins">Mins</th>
              <th class="th-in">Actual In</th>
              <th class="th-out">Actual Out</th>
              <th class="th-grade">Grade / Sec</th>
              <th class="th-subject">Subject / Learning Area</th>
              <th class="th-room">Room</th>
              <th class="th-status">Instruction Status</th>
              <th class="th-remarks">Signature</th>
            </tr>
          </thead>
          <tbody>''')

    row_num = 1
    for it in items:
        subj_upper = html.escape(it['subject'].upper())
        sec_short = html.escape(it['section'])
        html_out.append(f'''
            <tr>
              <td class="td-center" style="font-weight:700;color:#64748b;">{row_num}</td>
              <td class="day-cell">{it['day_abbr']}</td>
              <td class="time-slot">{it['time']}</td>
              <td class="td-center" style="font-weight:700;color:#475569;">{it['mins']}</td>
              <td class="td-center"></td>
              <td class="td-center"></td>
              <td class="grade-cell">{sec_short}</td>
              <td class="subject-cell">{subj_upper}</td>
              <td class="td-center"></td>
              <td>
                <div class="status-grid">
                  <div class="status-cell-row">
                    <label class="status-chk"><input type="checkbox"> In</label>
                    <label class="status-chk"><input type="checkbox"> Late</label>
                  </div>
                  <div class="status-cell-row">
                    <label class="status-chk"><input type="checkbox"> Abs</label>
                    <label class="status-chk"><input type="checkbox"> Sub</label>
                  </div>
                </div>
              </td>
              <td></td>
            </tr>''')
        row_num += 1

    blank_to_add = max(2, 14 - len(items))
    for b in range(blank_to_add):
        html_out.append(f'''
            <tr class="row-blank">
              <td class="td-center" style="color:#cbd5e1;font-weight:600;">{row_num}</td>
              <td></td>
              <td></td>
              <td></td>
              <td class="td-center"></td>
              <td class="td-center"></td>
              <td></td>
              <td style="color:#cbd5e1;font-size:7pt;font-weight:700;text-transform:uppercase;">(SUBSTITUTE / REMEDIAL LOAD)</td>
              <td></td>
              <td>
                <div class="status-grid">
                  <div class="status-cell-row">
                    <label class="status-chk"><input type="checkbox"> In</label>
                    <label class="status-chk"><input type="checkbox"> Late</label>
                  </div>
                  <div class="status-cell-row">
                    <label class="status-chk"><input type="checkbox"> Abs</label>
                    <label class="status-chk"><input type="checkbox"> Sub</label>
                  </div>
                </div>
              </td>
              <td></td>
            </tr>''')
        row_num += 1

    html_out.append(f'''
          </tbody>
        </table>

        <!-- SIGNATURE BLOCK -->
        <div class="sign-row" style="margin-top: auto; padding-top: 6px; border-top: 1.5px solid #cbd5e1;">
          <div class="sign-col">
            <div class="sign-line"></div>
            <div class="sign-label">{t_name_esc}</div>
            <div class="sign-title">Teacher's Signature over Printed Name</div>
          </div>
          <div class="sign-col">
            <div class="sign-line"></div>
            <div class="sign-label">Academic Coordinator / Department Head</div>
            <div class="sign-title">Verified & Monitored By</div>
          </div>
          <div class="sign-col">
            <div class="sign-line"></div>
            <div class="sign-label">School Principal / Directress</div>
            <div class="sign-title">Approved By</div>
          </div>
        </div>
      </div>
    </div>
''')

html_out.append('''
  </div>

  <script>
    const AMIS_LOGO_SRC = "''' + amis_b64 + '''";
    const DEPED_LOGO_SRC = "''' + deped_b64 + '''";

    let activeIndex = 0;
    let currentDept = 'all';

    function populateLogos() {
      document.querySelectorAll('.amis-img').forEach(img => {
        if (!img.src || img.src === window.location.href) {
          img.src = AMIS_LOGO_SRC;
        }
      });
      document.querySelectorAll('.deped-img').forEach(img => {
        if (!img.src || img.src === window.location.href) {
          img.src = DEPED_LOGO_SRC;
        }
      });
    }

    function showTeacher(val) {
      const select = document.getElementById('teacher-select');
      if (val === 'all') {
        if (select) select.value = 'all';
        document.querySelectorAll('.page-sheet').forEach(sh => {
          const cat = sh.getAttribute('data-cat');
          let match = false;
          if (currentDept === 'all') match = true;
          else if (currentDept === 'jhs' && cat === 'jhs') match = true;
          else if (currentDept === 'shs' && cat === 'shs') match = true;
          else if (currentDept === 'elem' && cat === 'elem') match = true;
          else if (currentDept === 'isal' && cat === 'isal') match = true;

          if (match) sh.classList.remove('hidden-sheet');
          else sh.classList.add('hidden-sheet');
        });
        window.scrollTo({ top: 0, behavior: 'smooth' });
        return;
      }

      const idx = parseInt(val, 10);
      activeIndex = idx;
      if (select) select.value = idx;

      document.querySelectorAll('.page-sheet').forEach(sh => {
        const sIdx = parseInt(sh.getAttribute('data-idx'), 10);
        if (sIdx === idx) {
          sh.classList.remove('hidden-sheet');
        } else {
          sh.classList.add('hidden-sheet');
        }
      });

      window.scrollTo({ top: 0, behavior: 'smooth' });
    }

    function onTeacherSelectChange() {
      const select = document.getElementById('teacher-select');
      showTeacher(select.value);
    }

    function navigateTeacher(delta) {
      const select = document.getElementById('teacher-select');
      const visibleOptions = Array.from(select.options).filter(opt => opt.style.display !== 'none' && opt.value !== 'all');
      if (visibleOptions.length === 0) return;

      const currentVal = parseInt(select.value, 10);
      let curIndexInVisible = visibleOptions.findIndex(opt => parseInt(opt.value, 10) === currentVal);
      if (curIndexInVisible === -1) curIndexInVisible = 0;

      let nextIndexInVisible = curIndexInVisible + delta;
      if (nextIndexInVisible < 0) nextIndexInVisible = visibleOptions.length - 1;
      if (nextIndexInVisible >= visibleOptions.length) nextIndexInVisible = 0;

      showTeacher(parseInt(visibleOptions[nextIndexInVisible].value, 10));
    }

    function filterDept(dept) {
      currentDept = dept;
      document.querySelectorAll('.dept-pill').forEach(btn => {
        btn.classList.remove('active');
      });
      const activeBtn = Array.from(document.querySelectorAll('.dept-pill')).find(b => {
        if (dept === 'all' && b.innerText.includes('All')) return true;
        if (dept === 'jhs' && b.innerText.includes('Junior')) return true;
        if (dept === 'shs' && b.innerText.includes('Senior')) return true;
        if (dept === 'elem' && b.innerText.includes('Elementary')) return true;
        if (dept === 'isal' && b.innerText.includes('ISAL')) return true;
        return false;
      });
      if (activeBtn) activeBtn.classList.add('active');

      const select = document.getElementById('teacher-select');

      document.querySelectorAll('.page-sheet').forEach(sh => {
        const cat = sh.getAttribute('data-cat');
        const idx = parseInt(sh.getAttribute('data-idx'), 10);
        const opt = select.querySelector('option[value="' + idx + '"]');

        let match = false;
        if (dept === 'all') match = true;
        else if (dept === 'jhs' && cat === 'jhs') match = true;
        else if (dept === 'shs' && cat === 'shs') match = true;
        else if (dept === 'elem' && cat === 'elem') match = true;
        else if (dept === 'isal' && cat === 'isal') match = true;

        if (opt) {
          opt.style.display = match ? '' : 'none';
        }
      });

      showTeacher('all');
    }

    function printActiveTeacher() {
      const select = document.getElementById('teacher-select');
      if (select && select.value === 'all') {
        showTeacher(activeIndex >= 0 ? activeIndex : 0);
      } else {
        showTeacher(activeIndex);
      }
      setTimeout(() => {
        window.print();
      }, 200);
    }

    function printAllTeachers() {
      filterDept('all');

      document.querySelectorAll('.page-sheet').forEach(sh => {
        sh.classList.remove('hidden-sheet');
      });

      const select = document.getElementById('teacher-select');
      if (select) select.value = 'all';

      setTimeout(() => {
        window.print();
      }, 300);
    }

    window.addEventListener('afterprint', () => {
      const select = document.getElementById('teacher-select');
      if (select && select.value === 'all') {
        showTeacher('all');
      } else {
        showTeacher(activeIndex);
      }
    });

    window.addEventListener('keydown', (e) => {
      if (e.key === 'ArrowLeft') {
        navigateTeacher(-1);
      } else if (e.key === 'ArrowRight') {
        navigateTeacher(1);
      }
    });

    window.addEventListener('DOMContentLoaded', () => {
      populateLogos();
      filterDept('all');
    });
  </script>
</body>
</html>
''')

output_content = "".join(html_out)
with open(OUTPUT_FILE, 'w', encoding='utf-8') as f:
    f.write(output_content)

print(f"Successfully generated {OUTPUT_FILE} ({len(output_content)} bytes).")

with open(ROOT_OUTPUT_FILE, 'w', encoding='utf-8') as f:
    f.write(output_content)
print(f"Successfully mirrored to {ROOT_OUTPUT_FILE}")
