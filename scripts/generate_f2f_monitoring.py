import openpyxl, json, re, html, base64, os

print("=== Compiling Master Pure F2F Monitoring & Teaching Loads System ===")

REPO_DIR = os.path.abspath(os.path.join(os.path.dirname(__file__), '..'))
DATA_FILE = os.path.join(REPO_DIR, 'data', 'schedule_f2f_latest.xlsx')
AMIS_LOGO = os.path.join(REPO_DIR, 'public', 'amis_logo_opt.png')
DEPED_LOGO = os.path.join(REPO_DIR, 'public', 'deped_logo_opt.png')
OUTPUT_FILE = os.path.join(REPO_DIR, 'public', 'f2f-monitoring.html')

wb = openpyxl.load_workbook(DATA_FILE, data_only=True)

with open(AMIS_LOGO, 'rb') as f:
    amis_b64 = "data:image/png;base64," + base64.b64encode(f.read()).decode('utf-8')

with open(DEPED_LOGO, 'rb') as f:
    deped_b64 = "data:image/png;base64," + base64.b64encode(f.read()).decode('utf-8')

days = ['Sunday', 'Monday', 'Tuesday', 'Wednesday', 'Thursday']
day_abbr_map = {
    'Sunday': 'SUN',
    'Monday': 'MON',
    'Tuesday': 'TUE',
    'Wednesday': 'WED',
    'Thursday': 'THU',
    'Friday': 'FRI',
    'Saturday': 'SAT'
}

def get_grid(sheet):
    grid = {}
    for r in range(1, sheet.max_row + 1):
        for c in range(1, sheet.max_column + 1):
            val = sheet.cell(r, c).value
            grid[(r, c)] = str(val).strip() if val is not None else ''
            
    for rng in sheet.merged_cells.ranges:
        top_val = grid.get((rng.min_row, rng.min_col), '')
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                grid[(r, c)] = top_val
    return grid

elem_grid = get_grid(wb['ELEM'])
hs_new_grid = get_grid(wb['HS SCHED (NEW)'])
hs_grid = get_grid(wb['HS SCHED'])

def clean_time(t_str):
    if not t_str: return ''
    t = str(t_str).strip()
    if any(kw in t.upper() for kw in ['GRADE', 'USTADH', 'TEACHER']):
        return ''
    
    m_f2f = re.search(r'([\d:apm\.\s-]+?)\s*(?:\(F2F\)|F2F)', t, re.IGNORECASE)
    if m_f2f:
        t = m_f2f.group(1).strip()
    else:
        t = re.sub(r'\(.*?\)', '', t).strip()

    t = re.sub(r'(\d{1,2}:\d{2}):+(\d{1,2}:\d{2})', r'\1 - \2', t)
    t = re.sub(r'(\d{1,2}:\d{2}):00', r'\1', t)
    
    is_pm = bool(re.search(r'(?:p\.?m\.?|pm)', t, re.IGNORECASE))
    is_am = bool(re.search(r'(?:a\.?m\.?|am)', t, re.IGNORECASE))
    
    clean = re.sub(r'(?:a\.?m\.?|p\.?m\.?|am|pm)', '', t, flags=re.IGNORECASE).strip()
    
    parts = re.split(r'\s*[-–]\s*', clean)
    if len(parts) == 2:
        start, end = parts[0].strip(), parts[1].strip()
        start = re.sub(r'^0(\d:)', r'\1', start)
        end = re.sub(r'^0(\d:)', r'\1', end)
        
        try:
            start_hour = int(start.split(':')[0]) if ':' in start else 0
            end_hour = int(end.split(':')[0]) if ':' in end else 0
        except ValueError:
            return t
            
        if is_pm:
            if start_hour == 11 and end_hour == 12:
                return f"{start} AM - {end} PM"
            return f"{start} - {end} PM"
        elif is_am:
            return f"{start} - {end} AM"
        else:
            if start_hour in [7, 8, 9, 10, 11] and end_hour in [7, 8, 9, 10, 11]:
                return f"{start} - {end} AM"
            elif start_hour == 11 and (end_hour == 12 or end_hour <= 1):
                return f"{start} AM - {end} PM"
            else:
                return f"{start} - {end} PM"
    elif len(parts) == 1:
        val = parts[0].strip()
        val = re.sub(r'^0(\d:)', r'\1', val)
        if is_pm or val.startswith(('12:', '1:', '2:', '3:', '4:', '5:')):
            return f"{val} PM"
        else:
            return f"{val} AM"
    
    return t

def parse_mins(m_str):
    if not m_str: return 40
    m_str = str(m_str).strip()
    m = re.search(r'(\d+)', m_str)
    if m:
        return int(m.group(1))
    return 40

def is_routine_text(text):
    if not text: return False
    u = text.upper()
    return any(kw in u for kw in ['GENERAL ASSEMBLY', 'RECESS', 'LUNCH', 'SALAH', 'DEPARTURE', 'SHORT BREAK', 'TRANSITION', 'HOMEROOM', 'ENTRANCE EXAM REVIEW', 'RESEARCH CONSULTATION'])

def clean_subject_name(raw_text):
    raw = raw_text.strip()
    if not raw: return ''
    clean = re.sub(r'[-–]?\s*(?:Tchr\.?|Teacher|Sir|Ustadh|Ust\.?|Alim|Tr\.?)\s+[A-Za-z]+', '', raw, flags=re.IGNORECASE).strip()
    clean = re.sub(r'\s+', ' ', clean).strip(' -–')
    return clean or raw

def extract_grade_cohort(raw_text):
    m = re.search(r'(?:Grade\s*\d+|G\d+|\d+&\d+|\b\d+\b|Kinder\s*\d+|K\d+)', raw_text, re.IGNORECASE)
    if m: return m.group(0)
    for coh in ['Abu Bakr', 'Umar', 'Uthman', 'Ali', 'Anas', 'Muadh', 'Nuaym', 'Saad', 'Usama', 'Utbah', 'Abu Sufyan', 'Abu Dharr', 'Abu Hurayrah', 'Abu Ayyub', 'Girls', 'Boys', 'Mix']:
        if coh.lower() in raw_text.lower(): return coh
    return ''

def normalize_teacher(name):
    name = re.sub(r"\s+", " ", name.strip()).strip("-– ")
    if re.match(r"^Teacher\b", name, re.IGNORECASE):
        name = "Tchr. " + name[7:].strip()
    elif re.match(r"^Tr\.?\b", name, re.IGNORECASE):
        name = "Tchr. " + name[3:].strip()
    elif re.match(r"^Tchr\.?\b", name, re.IGNORECASE):
        name = "Tchr. " + re.sub(r"^Tchr\.?\s*", "", name, flags=re.IGNORECASE)
    elif re.match(r"^Ustadha\b", name, re.IGNORECASE):
        name = "Ustadha " + name[7:].strip()
    elif re.match(r"^Ustadh\b", name, re.IGNORECASE):
        name = "Ustadh " + name[6:].strip()
    elif re.match(r"^Ust\.?\b", name, re.IGNORECASE):
        name = "Ust. " + re.sub(r"^Ust\.?\s*", "", name, flags=re.IGNORECASE)
        
    name = name.strip()
    if name in ["Tchr.", "Tchr", "Ust.", "Ust", "Teacher", "Tr.", "Tr", ""]: return "TBA"
    if name.upper() in ["TCHR. AHMAD", "SIR AHMAD"]: return "Tchr. Ahmad"
    if name in ["Ustadh Jaisam", "Ust. Jaisam"]: return "Ustadh Jaisam"
    if name == "Tchr. Kat": return "Tchr. Katrina"
    if name in ["Ust. Ubaydah", "Ust. Obaydah"]: return "Ust. Obaydah"
    if name in ["Ust. Silfah", "Ustadha Silfa", "Ust. Silfa"]: return "Ustadha Silfa"
    if name in ["Ustadha Saliha", "Ust. Saliha"]: return "Ustadha Saliha"
    if name in ["Tchr. Junaisa", "Tchr. Junaisah"]: return "Tchr. Junaisah"
    if name == "Tchr. Jairah": return "Tchr. Jayra"
    if name in ["Tchr. Moh", "Sir Mohaymen", "Sir Moh"]: return "Sir Moh"
    if name in ["Tchr. Shi", "Tchr. Shirehan"]: return "Tchr. Shirehan"
    if name in ["Tchr. Zara", "Tchr. Franchette"]: return "Tchr. Franchette"
    if name in ["Ust. Abdi", "Ust. Abdiraheem", "Ustadh Abdi", "Ustadh Abdiraheem", "Ustd. Abdi", "Ustd. Abdiraheem"]: return "Ust. Abdiraheem"
    if name in ["Ust. Ali", "Ustadh Ali", "Ustadh Muh Ali", "Ustdh ali", "Ustdh. Ali", "Ust. Muh Ali", "Ustadh Muh. Ali"]: return "Ustadh Muh Ali"
    return name

def clean_parse(text):
    text = text.strip()
    if not text: return '', '', 'EMPTY'
    upper = text.upper()
    for r in ['GENERAL ASSEMBLY', 'LUNCH AND SALAH', 'SALAH & DEPARTURE', 'DEPARTURE', 'RECESS', 'SHORT BREAK', 'TRANSITION', 'BREAK', 'HOMEROOM', 'ENTRANCE EXAM REVIEW', 'RESEARCH CONSULTATION']:
        if upper == r or (upper.startswith(r) and not any(k in upper for k in ['TCHR', 'UST', 'SIR', 'ALIM', 'TEACHER', 'TR'])):
            return text, '', 'ROUTINE'
            
    if 'WRAP-UP TIME' in upper:
        m = re.search(r'Wrap-Up Time\s*[-–]?\s*(.*)', text, re.IGNORECASE)
        t = normalize_teacher(m.group(1).strip()) if m else 'Tchr. Keychell'
        return 'Wrap-Up Time', t, 'CLASS'
        
    t_match = re.search(r'[-–]?\s*(Tchr\.?|Ust\.?|Sir|Alim|Teacher|Ustadha|Ustadh|Tr\.?)\s+(.*)$', text, re.IGNORECASE)
    if t_match:
        tchr_full = normalize_teacher((t_match.group(1) + ' ' + t_match.group(2)).strip())
        subj_part = text[:t_match.start()].strip().rstrip('-–').strip()
        return subj_part or text, tchr_full, 'CLASS'
        
    for p in ['Tchr.', 'Teacher', 'Sir', 'Ustadh', 'Ustadha', 'Ust.', 'Alim', 'Tr.']:
        if p.lower() in text.lower():
            idx = text.lower().find(p.lower())
            tchr_full = normalize_teacher(text[idx:].strip())
            subj_part = text[:idx].strip().rstrip('-–').strip()
            return subj_part or text, tchr_full, 'CLASS'
            
    return text, 'TBA', 'CLASS'

def time_to_sort_key(t_str):
    m = re.search(r'(\d+):(\d+)', t_str)
    if not m: return 9999
    h, mins = int(m.group(1)), int(m.group(2))
    if 'p.m.' in t_str.lower() or 'pm' in t_str.lower():
        if h < 12: h += 12
    elif 'a.m.' in t_str.lower() or 'am' in t_str.lower():
        if h == 12: h = 0
    else:
        if h in [1, 2, 3, 4, 5, 6]: h += 12
    return h * 60 + mins

# Sections
sections = [
    {'id': 'k1', 'code': 'K1', 'name': 'Kindergarten 1', 'dept_code': 'elem', 'dept_label': 'Kindergarten & Elementary Department', 'sheet': 'ELEM', 'r_start': 16, 'r_end': 21, 'has_mins': True, 'room': ''},
    {'id': 'k2', 'code': 'K2', 'name': 'Kindergarten 2', 'dept_code': 'elem', 'dept_label': 'Kindergarten & Elementary Department', 'sheet': 'ELEM', 'r_start': 5, 'r_end': 12, 'has_mins': True, 'room': ''},
    {'id': 'g1', 'code': 'G1', 'name': 'Grade 1', 'dept_code': 'elem', 'dept_label': 'Kindergarten & Elementary Department', 'sheet': 'ELEM', 'r_start': 26, 'r_end': 37, 'has_mins': True, 'room': ''},
    {'id': 'g2', 'code': 'G2', 'name': 'Grade 2', 'dept_code': 'elem', 'dept_label': 'Kindergarten & Elementary Department', 'sheet': 'ELEM', 'r_start': 41, 'r_end': 52, 'has_mins': True, 'room': ''},
    {'id': 'g3', 'code': 'G3', 'name': 'Grade 3', 'dept_code': 'elem', 'dept_label': 'Kindergarten & Elementary Department', 'sheet': 'ELEM', 'r_start': 56, 'r_end': 71, 'has_mins': True, 'room': ''},
    {'id': 'g4', 'code': 'G4', 'name': 'Grade 4', 'dept_code': 'elem', 'dept_label': 'Kindergarten & Elementary Department', 'sheet': 'ELEM', 'r_start': 76, 'r_end': 86, 'has_mins': True, 'room': ''},
    {'id': 'g5', 'code': 'G5', 'name': 'Grade 5', 'dept_code': 'elem', 'dept_label': 'Kindergarten & Elementary Department', 'sheet': 'ELEM', 'r_start': 91, 'r_end': 101, 'has_mins': True, 'room': ''},
    {'id': 'g6', 'code': 'G6', 'name': 'Grade 6', 'dept_code': 'elem', 'dept_label': 'Kindergarten & Elementary Department', 'sheet': 'ELEM', 'r_start': 106, 'r_end': 116, 'has_mins': True, 'room': ''},
    {'id': 'g78g', 'code': 'G78-G', 'name': 'Grade 7 & 8 Girls', 'dept_code': 'jhs', 'dept_label': 'Junior High School Department', 'sheet': 'HS SCHED (NEW)', 'r_start': 7, 'r_end': 18, 'has_mins': True, 'room': ''},
    {'id': 'g78b', 'code': 'G78-B', 'name': 'Grade 7 & 8 Boys', 'dept_code': 'jhs', 'dept_label': 'Junior High School Department', 'sheet': 'HS SCHED (NEW)', 'r_start': 23, 'r_end': 34, 'has_mins': True, 'room': ''},
    {'id': 'g910g', 'code': 'G910-G', 'name': 'Grade 9 & 10 Girls', 'dept_code': 'jhs', 'dept_label': 'Junior High School Department', 'sheet': 'HS SCHED (NEW)', 'r_start': 39, 'r_end': 50, 'has_mins': True, 'room': ''},
    {'id': 'g910b', 'code': 'G910-B', 'name': 'Grade 9 & 10 Boys', 'dept_code': 'jhs', 'dept_label': 'Junior High School Department', 'sheet': 'HS SCHED (NEW)', 'r_start': 55, 'r_end': 66, 'has_mins': True, 'room': ''},
    {'id': 'g11', 'code': 'G11', 'name': 'Grade 11', 'dept_code': 'shs', 'dept_label': 'Senior High School Department', 'sheet': 'HS SCHED (NEW)', 'r_start': 71, 'r_end': 82, 'has_mins': True, 'room': ''},
    {'id': 'g12', 'code': 'G12', 'name': 'Grade 12', 'dept_code': 'shs', 'dept_label': 'Senior High School Department', 'sheet': 'HS SCHED (NEW)', 'r_start': 89, 'r_end': 100, 'has_mins': True, 'room': ''}
]

# Build Header
def build_header_html(form_title, form_subtitle=None):
    sub_markup = f'<div class="form-sub">{form_subtitle}</div>' if form_subtitle else '<div class="form-sub">Pure Face-to-Face & Blended Modality &bull; School Year 2026 - 2027</div>'
    return f'''
        <div class="sheet-header">
          <div class="header-logo-side">
            <img class="header-logo deped-img" alt="Department of Education Logo">
          </div>
          <div class="header-center-text">
            <div class="arabic-header" dir="rtl" lang="ar">المدرسة المنورة الإسلامية</div>
            <div class="school-name">AL MUNAWWARA ISLAMIC SCHOOL</div>
            <div class="form-title">{form_title}</div>
            {sub_markup}
          </div>
          <div class="header-logo-side">
            <img class="header-logo amis-img" alt="Al Munawwara Islamic School Official Seal">
          </div>
        </div>
    '''

# 1. PARSE HS LOADS (15 Teachers)
ws_loads = wb['HS LOADS']
hs_coords = [
    ('Teacher Jayra', 2, 2, 'Junior & Senior High School Faculty', 'hs_loads'),
    ('Teacher Aniah', 2, 10, 'Junior & Senior High School Faculty', 'hs_loads'),
    ('Teacher Halnaisa', 2, 18, 'Junior & Senior High School Faculty', 'hs_loads'),
    ('Teacher Shanen', 22, 2, 'Junior & Senior High School Faculty', 'hs_loads'),
    ('Teacher Shirehan', 22, 10, 'Junior & Senior High School Faculty', 'hs_loads'),
    ('Teacher Abegail', 22, 18, 'Junior & Senior High School Faculty', 'hs_loads'),
    ('Teacher Rowena', 42, 2, 'Junior & Senior High School Faculty', 'hs_loads'),
    ('Teacher Nof', 42, 10, 'Junior & Senior High School Faculty', 'hs_loads'),
    ('Teacher Thea', 42, 18, 'Junior & Senior High School Faculty', 'hs_loads'),
    ('Teacher Nadzra', 62, 2, 'Junior & Senior High School Faculty', 'hs_loads'),
    ('Teacher Sophia', 62, 10, 'Junior & Senior High School Faculty', 'hs_loads'),
    ('Teacher Wardah', 62, 18, 'Junior & Senior High School Faculty', 'hs_loads'),
    ('Teacher Ethel', 82, 2, 'Junior & Senior High School Faculty', 'hs_loads'),
    ('Sir Mohaymen', 82, 10, 'Junior & Senior High School Faculty', 'hs_loads'),
    ('Teacher Marie', 82, 18, 'Junior & Senior High School Faculty', 'hs_loads'),
]

teacher_loads_data = []

for default_name, r, c, dept, cat in hs_coords:
    t_name = str(ws_loads.cell(r, c).value or default_name).strip()
    grid_rows = []
    summary_map = {}
    total_periods = 0
    total_mins = 0

    for row in range(r+2, r+19):
        time_slot = clean_time(str(ws_loads.cell(row, c).value or ''))
        mins_val = parse_mins(str(ws_loads.cell(row, c+1).value or ''))
        row_cells = [str(ws_loads.cell(row, c+2+d).value or '').strip() for d in range(5)]
        
        is_routine = False
        routine_label = ''
        first_content = next((rc for rc in row_cells if rc), '')
        if first_content and is_routine_text(first_content):
            is_routine = True
            routine_label = first_content
            
        day_entries = {}
        for d_idx, day in enumerate(days):
            cell_val = row_cells[d_idx]
            if is_routine:
                day_entries[day] = {'text': routine_label, 'is_class': False}
            elif cell_val and not is_routine_text(cell_val):
                clean_subj = clean_subject_name(cell_val)
                cohort = extract_grade_cohort(cell_val)
                day_entries[day] = {'text': clean_subj, 'cohort': cohort, 'raw': cell_val, 'is_class': True}
                total_periods += 1
                total_mins += mins_val
                
                sum_key = (clean_subj, time_slot)
                if sum_key not in summary_map:
                    summary_map[sum_key] = {
                        'subject': clean_subj,
                        'cohort': cohort,
                        'days': [],
                        'time': time_slot,
                        'mins_per_session': mins_val,
                        'sessions': 0,
                        'total_mins': 0
                    }
                if day not in summary_map[sum_key]['days']:
                    summary_map[sum_key]['days'].append(day)
                summary_map[sum_key]['sessions'] += 1
                summary_map[sum_key]['total_mins'] += mins_val
            else:
                day_entries[day] = {'text': '', 'is_class': False}

        grid_rows.append({
            'time': time_slot,
            'mins': mins_val,
            'is_routine': is_routine,
            'routine_label': routine_label,
            'days': day_entries
        })

    summary_list = list(summary_map.values())
    summary_list.sort(key=lambda x: -x['total_mins'])

    teacher_loads_data.append({
        'name': t_name,
        'department': dept,
        'category': cat,
        'grid_rows': grid_rows,
        'summary': summary_list,
        'total_periods': total_periods,
        'total_mins': total_mins,
        'total_hours': round(total_mins / 60.0, 1)
    })

print(f"Loaded {len(teacher_loads_data)} teachers from HS LOADS.")

# 2. ALSO PARSE ELEM TEACHERS TO LOAD DATA
elem_grade_sections = [
    ('Kindergarten 1', 14, 21),
    ('Kindergarten 2', 3, 12),
    ('Grade 1', 24, 37),
    ('Grade 2', 39, 52),
    ('Grade 3', 54, 71),
    ('Grade 4', 74, 86),
    ('Grade 5', 89, 101),
    ('Grade 6', 104, 116),
]

elem_teachers_slots = {}
for gname, r_start, r_end in elem_grade_sections:
    for r in range(r_start+2, r_end+1):
        t_val = clean_time(elem_grid.get((r, 2), ''))
        m_val = parse_mins(elem_grid.get((r, 3), ''))
        if not t_val: continue
        for d_idx, day in enumerate(days):
            cell_val = elem_grid.get((r, 4 + d_idx), '')
            if not cell_val: continue
            
            m = re.search(r'[-–]?\s*(Tchr\.?|Ust\.?|Sir|Teacher|Ustadha|Ustadh|Alim)\s+([A-Za-z]+)', cell_val, re.I)
            if m:
                t_raw = f"{m.group(1)} {m.group(2)}"
                subj = cell_val[:m.start()].strip(' -–')
            else:
                m2 = re.search(r'(Tchr\.?|Teacher|Sir|Ust\.?|Ustadh|Alim)\s+([A-Za-z]+)', cell_val, re.I)
                if m2:
                    t_raw = f"{m2.group(1)} {m2.group(2)}"
                    subj = cell_val.replace(m2.group(0), '').strip(' -–')
                else:
                    continue
            
            norm_t = normalize_teacher(t_raw)
            if not norm_t or norm_t == 'TBA': continue
            
            if norm_t not in elem_teachers_slots:
                elem_teachers_slots[norm_t] = []
                
            elem_teachers_slots[norm_t].append({
                'grade': gname,
                'day': day,
                'time': t_val,
                'mins': m_val,
                'subject': subj or cell_val,
                'raw': cell_val
            })

elem_standard_times = [
    ('7:30-7:40 a.m.', 10, True, 'GENERAL ASSEMBLY (F2F)'),
    ('7:40-8:25 a.m.', 45, False, ''),
    ('8:25-9:05 a.m.', 40, False, ''),
    ('9:05-9:45 a.m.', 40, False, ''),
    ('9:45-10:00 a.m.', 15, True, 'RECESS'),
    ('10:00-10:45 a.m.', 45, False, ''),
    ('10:45-11:30 a.m.', 45, False, ''),
    ('11:30-12:40 p.m.', 70, True, 'LUNCH & SALAH'),
    ('12:40-1:25 p.m.', 45, False, ''),
    ('1:25-2:10 p.m.', 45, False, ''),
    ('2:15-3:00 p.m.', 45, False, ''),
    ('3:00-3:30 p.m.', 30, True, 'SALAH & DEPARTURE'),
]

hs_t_names = {t['name'] for t in teacher_loads_data}

for t_name, slots in sorted(elem_teachers_slots.items()):
    if any(h.lower() in t_name.lower() or t_name.lower() in h.lower() for h in hs_t_names):
        continue
    
    is_isal = any(t_name.startswith(p) for p in ['Ust.', 'Ustadh', 'Ustadha', 'Alim'])
    dept = 'ISAL Department Faculty' if is_isal else 'Elementary Department Faculty'
    cat = 'isal' if is_isal else 'elem'

    grid_rows = []
    summary_map = {}
    total_periods = len(slots)
    total_mins = sum(s['mins'] for s in slots)

    for time_slot, default_m, is_routine, routine_lbl in elem_standard_times:
        day_entries = {}
        for day in days:
            if is_routine:
                day_entries[day] = {'text': routine_lbl, 'is_class': False}
            else:
                matching = [s for s in slots if s['day'] == day and (s['time'] in time_slot or time_slot in s['time'])]
                if matching:
                    m_item = matching[0]
                    day_entries[day] = {
                        'text': m_item['subject'],
                        'cohort': m_item['grade'],
                        'raw': m_item['raw'],
                        'is_class': True
                    }
                else:
                    day_entries[day] = {'text': '', 'is_class': False}

        grid_rows.append({
            'time': time_slot,
            'mins': default_m,
            'is_routine': is_routine,
            'routine_label': routine_lbl,
            'days': day_entries
        })

    for s in slots:
        clean_subj = s['subject']
        sum_key = (clean_subj, s['grade'])
        if sum_key not in summary_map:
            summary_map[sum_key] = {
                'subject': clean_subj,
                'cohort': s['grade'],
                'days': [],
                'time': s['time'],
                'mins_per_session': s['mins'],
                'sessions': 0,
                'total_mins': 0
            }
        if s['day'] not in summary_map[sum_key]['days']:
            summary_map[sum_key]['days'].append(s['day'])
        summary_map[sum_key]['sessions'] += 1
        summary_map[sum_key]['total_mins'] += s['mins']

    summary_list = list(summary_map.values())
    summary_list.sort(key=lambda x: -x['total_mins'])

    teacher_loads_data.append({
        'name': t_name,
        'department': dept,
        'category': cat,
        'grid_rows': grid_rows,
        'summary': summary_list,
        'total_periods': total_periods,
        'total_mins': total_mins,
        'total_hours': round(total_mins / 60.0, 1)
    })

print(f"Total Per-Teacher Load sheets compiled: {len(teacher_loads_data)}")

# Aggregation for Attendance View
teacher_schedule = {}
for sec in sections:
    s_grid = elem_grid if sec['sheet'] == 'ELEM' else (hs_new_grid if sec['sheet'] == 'HS SCHED (NEW)' else hs_grid)
    for r in range(sec['r_start'], sec['r_end'] + 1):
        t_slot = clean_time(s_grid.get((r, 2), ''))
        if not t_slot: continue
        
        for d_idx, day in enumerate(days):
            c_val = s_grid.get((r, 4 + d_idx), '').strip()
            if not c_val: continue
            
            subj, tchr, kind = clean_parse(c_val)
            if kind == 'CLASS' and tchr and tchr != 'TBA':
                if tchr not in teacher_schedule:
                    teacher_schedule[tchr] = []
                teacher_schedule[tchr].append({
                    'day': day,
                    'time': t_slot,
                    'section': sec['name'],
                    'subject': subj,
                    'dept_code': sec['dept_code'],
                    'dept_label': sec['dept_label']
                })

# START HTML GENERATION
html_out = []
html_out.append('''<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8">
  <meta name="viewport" content="width=device-width, initial-scale=1.0">
  <link rel="preconnect" href="https://fonts.googleapis.com">
  <link rel="preconnect" href="https://fonts.gstatic.com" crossorigin>
  <link href="https://fonts.googleapis.com/css2?family=Amiri:wght@700&display=swap" rel="stylesheet">
  <title>Pure F2F Daily Teaching Load Monitoring & Schedule Forms - Al Munawwara Islamic School</title>
  <style>
    @page {
      size: A4 portrait;
      margin: 6mm 8mm 6mm 8mm;
    }
    * {
      box-sizing: border-box;
      font-family: -apple-system, BlinkMacSystemFont, "Segoe UI", Roboto, Arial, sans-serif;
    }
    body {
      margin: 0;
      padding: 0;
      background: #f1f5f9;
      color: #0f172a;
      -webkit-print-color-adjust: exact;
      print-color-adjust: exact;
    }
    .toolbar {
      background: #ffffff;
      color: #0f172a;
      padding: 12px 24px;
      position: sticky;
      top: 0;
      z-index: 999;
      box-shadow: 0 2px 10px rgba(0,0,0,0.06);
      border-bottom: 2px solid #059669;
    }
    .toolbar-header {
      display: flex;
      justify-content: space-between;
      align-items: center;
      margin-bottom: 12px;
      flex-wrap: wrap;
      gap: 12px;
    }
    .toolbar-brand {
      display: flex;
      align-items: center;
      gap: 12px;
    }
    .toolbar-logo {
      width: 44px;
      height: 44px;
      object-fit: contain;
    }
    .toolbar-title {
      font-size: 16px;
      font-weight: 800;
      color: #064e3b;
      letter-spacing: 0.3px;
      display: flex;
      align-items: center;
      gap: 10px;
    }
    .badge-f2f {
      background: #d1fae5;
      color: #065f46;
      border: 1px solid #10b981;
      font-size: 11px;
      padding: 2px 8px;
      border-radius: 9999px;
      font-weight: 800;
      text-transform: uppercase;
    }
    .modality-bar {
      display: flex;
      align-items: center;
      gap: 8px;
      margin-bottom: 12px;
      padding-bottom: 10px;
      border-bottom: 1px solid #e2e8f0;
      flex-wrap: wrap;
    }
    .modality-label {
      font-size: 11px;
      font-weight: 800;
      color: #64748b;
      text-transform: uppercase;
      letter-spacing: 0.4px;
      margin-right: 4px;
    }
    .modality-pill {
      display: inline-flex;
      align-items: center;
      gap: 6px;
      padding: 5px 12px;
      border-radius: 6px;
      font-size: 11.5px;
      font-weight: 700;
      text-decoration: none;
      transition: all 0.15s ease;
      border: 1px solid #cbd5e1;
      background: #f8fafc;
      color: #334155;
    }
    .modality-pill:hover {
      background: #f1f5f9;
      color: #0f172a;
      border-color: #94a3b8;
    }
    .modality-pill.active {
      background: #064e3b;
      color: #ffffff;
      border-color: #064e3b;
    }
    .modality-dot {
      width: 7px;
      height: 7px;
      border-radius: 50%;
      display: inline-block;
    }
    .dot-active {
      background: #10b981;
      box-shadow: 0 0 0 2px rgba(16, 185, 129, 0.3);
    }
    .dot-dev {
      background: #f59e0b;
    }
    .badge-dev {
      background: #fef3c7;
      color: #92400e;
      border: 1px solid #fde68a;
      font-size: 9.5px;
      padding: 1px 6px;
      border-radius: 9999px;
      font-weight: 800;
      letter-spacing: 0.2px;
      text-transform: uppercase;
    }
    .toolbar-actions {
      display: flex;
      gap: 8px;
      align-items: center;
      flex-wrap: wrap;
    }
    .btn-print {
      background: #0284c7;
      color: white;
      border: none;
      padding: 8px 18px;
      font-size: 13px;
      font-weight: 700;
      border-radius: 6px;
      cursor: pointer;
      display: inline-flex;
      align-items: center;
      gap: 6px;
      transition: background 0.15s ease;
    }
    .btn-print:hover { background: #0369a1; }
    .btn-print-green {
      background: #059669;
    }
    .btn-print-green:hover { background: #047857; }
    .view-tabs {
      display: flex;
      gap: 8px;
      margin-bottom: 12px;
      border-bottom: 1px solid #e2e8f0;
      padding-bottom: 8px;
      overflow-x: auto;
    }
    .view-tab {
      background: #f8fafc;
      color: #475569;
      border: 1px solid #cbd5e1;
      padding: 7px 16px;
      border-radius: 6px;
      font-size: 12px;
      font-weight: 600;
      cursor: pointer;
      white-space: nowrap;
      transition: all 0.15s;
    }
    .view-tab:hover {
      background: #f1f5f9;
      color: #0f172a;
    }
    .view-tab.active {
      background: #059669;
      color: #ffffff;
      border-color: #059669;
      font-weight: 700;
      box-shadow: 0 2px 4px rgba(5, 150, 105, 0.2);
    }
    .view-tab-highlight {
      border-color: #059669;
      color: #064e3b;
      font-weight: 700;
    }
    .filter-grid {
      display: grid;
      grid-template-columns: repeat(auto-fit, minmax(180px, 1fr));
      gap: 12px;
      align-items: center;
    }
    .filter-group {
      display: flex;
      flex-direction: column;
      gap: 4px;
    }
    .filter-group label {
      font-size: 11px;
      font-weight: 700;
      color: #475569;
      text-transform: uppercase;
      letter-spacing: 0.3px;
    }
    .filter-group select, .filter-group input {
      background: #ffffff;
      border: 1px solid #cbd5e1;
      color: #0f172a;
      padding: 7px 10px;
      font-size: 12px;
      border-radius: 6px;
      outline: none;
      transition: border-color 0.15s;
    }
    .filter-group select:focus, .filter-group input:focus {
      border-color: #059669;
      box-shadow: 0 0 0 2px rgba(5, 150, 105, 0.15);
    }
    .status-summary {
      margin-top: 10px;
      font-size: 11.5px;
      color: #64748b;
      display: flex;
      justify-content: space-between;
      align-items: center;
    }
    .sheet-wrapper {
      padding: 24px 10px 60px 10px;
      display: flex;
      flex-direction: column;
      align-items: center;
      gap: 24px;
    }
    .page-sheet {
      width: 210mm;
      min-height: 290mm;
      padding: 6mm 8mm;
      background: #ffffff;
      box-shadow: 0 4px 15px rgba(0,0,0,0.08);
      border: 1px solid #e2e8f0;
      position: relative;
      display: flex;
      flex-direction: column;
      justify-content: space-between;
      box-sizing: border-box;
      page-break-after: always;
      break-after: page;
    }
    .sheet-content {
      flex: 1;
      display: flex;
      flex-direction: column;
    }
    .sheet-header {
      display: flex;
      align-items: center;
      justify-content: space-between;
      border-bottom: 2px solid #0f172a;
      padding-bottom: 5px;
      margin-bottom: 5px;
      gap: 10px;
    }
    .header-logo-side {
      width: 50px;
      display: flex;
      justify-content: center;
      align-items: center;
      flex-shrink: 0;
    }
    .header-logo {
      width: 50px;
      height: 50px;
      object-fit: contain;
    }
    .header-center-text {
      flex: 1;
      text-align: center;
    }
    .arabic-header {
      font-family: 'Amiri', 'Traditional Arabic', 'Noto Naskh Arabic', 'Scheherazade New', 'Times New Roman', serif;
      font-size: 15pt;
      font-weight: 700;
      color: #064e3b;
      direction: rtl;
      line-height: 1.2;
      margin-bottom: 1px;
      text-align: center;
    }
    .school-name {
      font-size: 11pt;
      font-weight: 900;
      color: #0f172a;
      letter-spacing: 0.5px;
      margin: 1px 0 2px 0;
      text-transform: uppercase;
      line-height: 1.15;
    }
    .form-title {
      font-size: 9pt;
      font-weight: 800;
      color: #064e3b;
      text-transform: uppercase;
      letter-spacing: 0.5px;
      line-height: 1.15;
    }
    .form-sub {
      font-size: 7pt;
      color: #475569;
      font-weight: 600;
      margin-top: 1px;
    }
    .meta-box {
      border: 1px solid #94a3b8;
      background: #f8fafc;
      border-radius: 3px;
      padding: 4px 8px;
      margin-bottom: 5px;
      display: grid;
      grid-template-columns: 1.3fr 1.2fr 1fr;
      gap: 3px 12px;
      font-size: 7.5pt;
    }
    .meta-row {
      display: flex;
      align-items: center;
      gap: 4px;
    }
    .meta-lbl {
      font-weight: 700;
      color: #334155;
      white-space: nowrap;
    }
    .meta-val {
      font-weight: 600;
      color: #0f172a;
      border-bottom: 1px dotted #94a3b8;
      flex: 1;
      min-height: 12px;
    }
    .td-bold {
      font-weight: 800;
      color: #064e3b;
    }
    .section-title {
      font-size: 7.6pt;
      font-weight: 800;
      color: #064e3b;
      text-transform: uppercase;
      letter-spacing: 0.3px;
      margin: 3px 0 2px 0;
      display: flex;
      align-items: center;
      gap: 6px;
    }
    .section-title::after {
      content: '';
      flex: 1;
      height: 1px;
      background: #cbd5e1;
    }
    table.sheet-table {
      width: 100%;
      border-collapse: collapse;
      font-size: 7.4pt;
      line-height: 1.15;
    }
    table.sheet-table th, table.sheet-table td {
      border: 1px solid #94a3b8;
      padding: 3px 5px;
      vertical-align: middle;
    }
    table.sheet-table thead th {
      background: #e2e8f0;
      color: #0f172a;
      font-weight: 800;
      text-align: center;
      font-size: 7.5pt;
      padding: 4px 5px;
    }
    table.schedule-table {
      width: 100%;
      border-collapse: collapse;
      font-size: 6.9pt;
      line-height: 1.15;
      margin-bottom: 5px;
    }
    table.schedule-table th, table.schedule-table td {
      border: 1px solid #94a3b8;
      padding: 2px 4px;
      vertical-align: middle;
    }
    table.schedule-table thead th {
      background: #e2e8f0;
      color: #0f172a;
      font-weight: 800;
      text-align: center;
      font-size: 7pt;
      padding: 3px 4px;
    }
    .time-slot-col {
      width: 17%;
      text-align: center;
      font-weight: 700;
      background: #f8fafc;
      white-space: nowrap;
    }
    .mins-col {
      width: 5%;
      text-align: center;
      font-weight: 700;
      background: #f8fafc;
    }
    .day-col {
      width: 15.6%;
      text-align: center;
    }
    .routine-row-cell {
      background: #f1f5f9;
      color: #475569;
      font-weight: 700;
      text-align: center;
      font-size: 6.8pt;
      letter-spacing: 0.4px;
      padding: 2px 4px;
    }
    .class-cell {
      background: #ecfdf5;
      border: 1px solid #6ee7b7 !important;
      text-align: center;
      padding: 2px 3px;
    }
    .cell-subj {
      font-weight: 800;
      color: #064e3b;
      font-size: 6.9pt;
      line-height: 1.1;
    }
    .cell-cohort {
      font-size: 6.2pt;
      color: #0f172a;
      font-weight: 600;
      margin-top: 1px;
    }
    .free-cell {
      color: #cbd5e1;
      text-align: center;
      font-size: 6.5pt;
    }
    table.summary-table {
      width: 100%;
      border-collapse: collapse;
      font-size: 6.8pt;
      line-height: 1.15;
      margin-bottom: 4px;
    }
    table.summary-table th, table.summary-table td {
      border: 1px solid #cbd5e1;
      padding: 2.5px 5px;
      vertical-align: middle;
    }
    table.summary-table thead th {
      background: #e2e8f0;
      color: #0f172a;
      font-weight: 800;
      text-align: center;
      font-size: 6.9pt;
    }
    table.summary-table tfoot th {
      background: #f8fafc;
      color: #064e3b;
      font-weight: 800;
      border-top: 1.5px solid #0f172a;
      font-size: 7.2pt;
    }
    .sign-block {
      margin-top: 6px;
      display: grid;
      grid-template-columns: 1fr 1fr;
      gap: 24px;
      font-size: 7.2pt;
      padding-top: 4px;
      border-top: 1px dashed #cbd5e1;
    }
    .sign-col {
      display: flex;
      flex-direction: column;
    }
    .sign-title {
      color: #475569;
      font-weight: 700;
      margin-bottom: 16px;
    }
    .sign-name {
      font-weight: 800;
      color: #0f172a;
      border-top: 1px solid #0f172a;
      padding-top: 2px;
      text-transform: uppercase;
      letter-spacing: 0.3px;
    }
    .sign-role {
      color: #64748b;
      font-size: 6.5pt;
    }
    .row-routine {
      background: #f1f5f9;
      color: #475569;
      font-weight: 700;
      text-align: center;
      font-style: italic;
    }
    .row-blank {
      background: #ffffff;
      height: 24px;
    }
    .time-slot {
      font-weight: 700;
      color: #0f172a;
      text-align: center;
      white-space: nowrap;
    }
    .time-box {
      border: 1px solid #94a3b8;
      padding: 1px 6px;
      font-family: monospace;
      font-size: 7pt;
      color: #64748b;
      display: inline-block;
      min-width: 48px;
      text-align: center;
      background: #ffffff;
    }
    .status-boxes {
      display: flex;
      gap: 4px;
      font-size: 6.5pt;
      justify-content: center;
      white-space: nowrap;
    }
    .status-boxes span {
      color: #334155;
    }
    .status-grid {
      display: flex;
      flex-direction: column;
      gap: 1.5px;
      font-size: 6.5pt;
      font-weight: 700;
      color: #334155;
      line-height: 1.15;
    }
    .status-row {
      display: flex;
      justify-content: space-between;
      padding: 0 2px;
      white-space: nowrap;
    }
    .th-num { width: 3%; }
    .th-time { width: 14%; }
    .th-mins { width: 5%; }
    .th-in { width: 8.5%; }
    .th-out { width: 8.5%; }
    .th-subject { width: 23%; }
    .th-teacher { width: 18%; }
    .th-room { width: 6%; }
    .th-status { width: 10%; }
    .th-remarks { width: 13%; }
    .td-center { text-align: center; }
    .hidden-sheet {
      display: none !important;
    }
    @media print {
      body {
        background: #ffffff !important;
      }
      .toolbar {
        display: none !important;
      }
      .sheet-wrapper {
        padding: 0 !important;
        gap: 0 !important;
      }}
      .page-sheet {
        box-shadow: none !important;
        border: none !important;
        margin: 0 !important;
        padding: 5mm 6mm !important;
        width: 100% !important;
        min-height: auto !important;
        page-break-after: always !important;
        break-after: page !important;
      }
      .hidden-sheet {
        display: none !important;
      }
    }
  </style>
</head>
<body>

  <!-- LIGHT MODE TOOLBAR -->
  <div class="toolbar">
    <!-- MODALITY NAVIGATION -->
    <div class="modality-bar">
      <span class="modality-label">Learning Modality:</span>
      <a href="/f2f-monitoring.html" class="modality-pill active" title="Active S.Y. 2026 - 2027">
        <span class="modality-dot dot-active"></span>
         Face-to-Face (F2F)
      </a>
      <a href="/odl-first-shift.html" class="modality-pill" title="Online Distance Learning First Shift">
        <span class="modality-dot dot-active"></span>
        ODL First Shift
      </a>
      <a href="/odl-second-shift.html" class="modality-pill" title="Online Distance Learning Second Shift">
        <span class="modality-dot dot-active"></span>
        ODL Second Shift
      </a>
    </div>

    <div class="toolbar-header">
      <div class="toolbar-brand">
        <img class="toolbar-logo amis-img" alt="AMIS Logo">
        <div>
          <div class="toolbar-title">
            AL MUNAWWARA ISLAMIC SCHOOL
            <span class="badge-f2f">Pure F2F Portal</span>
          </div>
          <div style="font-size:11.5px;color:#475569;font-weight:500;">
            Instructional Teaching Load Monitoring System &bull; Pure Face-to-Face & Blended Modality (S.Y. 2026 - 2027)
          </div>
        </div>
      </div>
      <div class="toolbar-actions">
        <a href="/teacher-monitoring.html" class="btn-print" style="text-decoration:none;background:#059669;" title="Open Individual Teacher Load & Attendance Monitoring Sheets">
          ‍ Per-Teacher Forms
        </a>
        <button class="btn-print" onclick="window.print()">
          ️ Print Active Sheet
        </button>
        <button class="btn-print btn-print-green" onclick="showAllAndPrint()">
           Print All Filtered Sheets
        </button>
      </div>
    </div>

    <!-- VIEW TABS -->
    <div class="view-tabs">
      <button class="view-tab" onclick="switchView(this, 'daily')"> Classroom Daily Forms (1 Page/Grade/Day)</button>
      <button class="view-tab" onclick="switchView(this, 'walkthrough')"> Department Walkthrough Sheets (Cluster View)</button>
      <button class="view-tab" onclick="switchView(this, 'weekly')"> Weekly Timetable Matrix (Sunday-Thursday)</button>
      <button class="view-tab active view-tab-highlight" onclick="switchView(this, 'teacher_loads')">‍ Per-Teacher Weekly Loads (HS LOADS)</button>
      <button class="view-tab" onclick="switchView(this, 'teachers')"> Teacher Attendance & Monitoring Log</button>
    </div>

    <!-- FILTERS -->
    <div class="filter-grid">
      <div class="filter-group" id="day-filter-group" style="display:none;">
        <label>Day of Week</label>
        <select id="filter-day" onchange="applyFilters()">
          <option value="all">All Days (Sunday - Thursday)</option>
          <option value="Sunday">Sunday</option>
          <option value="Monday">Monday</option>
          <option value="Tuesday">Tuesday</option>
          <option value="Wednesday">Wednesday</option>
          <option value="Thursday">Thursday</option>
        </select>
      </div>

      <div class="filter-group" id="dept-filter-group">
        <label>Department / Category</label>
        <select id="filter-dept" onchange="applyFilters()">
          <option value="all">All Departments / Faculty</option>
          <option value="hs_loads">⭐ High School Faculty (HS LOADS - 15 Teachers)</option>
          <option value="elem">Kindergarten & Elementary Faculty</option>
          <option value="isal">ISAL & Islamic Studies Department</option>
        </select>
      </div>

      <div class="filter-group" id="grade-filter-group" style="display:none;">
        <label>Grade Level / Cohort</label>
        <select id="filter-grade" onchange="applyFilters()">
          <option value="all">All Grades (Pure F2F)</option>
          <option value="k1">Kindergarten 1 (F2F)</option>
          <option value="k2">Kindergarten 2 (F2F)</option>
          <option value="g1">Grade 1 (F2F)</option>
          <option value="g2">Grade 2 (F2F)</option>
          <option value="g3">Grade 3 (F2F)</option>
          <option value="g4">Grade 4 (F2F)</option>
          <option value="g5">Grade 5 (F2F)</option>
          <option value="g6">Grade 6 (F2F)</option>
          <option value="g78g">Grade 7 & 8 Girls (F2F)</option>
          <option value="g78b">Grade 7 & 8 Boys (F2F)</option>
          <option value="g910g">Grade 9 & 10 Girls (F2F)</option>
          <option value="g910b">Grade 9 & 10 Boys (F2F)</option>
          <option value="g11">Grade 11 (F2F)</option>
          <option value="g12">Grade 12 (F2F)</option>
        </select>
      </div>

      <div class="filter-group" id="teacher-filter-group">
        <label>Select Teacher</label>
        <select id="filter-teacher" onchange="applyFilters()">
          <option value="all">All Teachers (''')

html_out.append(str(len(teacher_loads_data)) + ' Instructors)</option>\n')
html_out.append('          <optgroup label="⭐ High School Faculty (From HS LOADS)">\n')

for t in [x for x in teacher_loads_data if x['category'] == 'hs_loads']:
    html_out.append(f'            <option value="{html.escape(t["name"])}">{html.escape(t["name"])} ({t["total_periods"]} periods - {t["total_hours"]} hrs)</option>\n')

html_out.append('''          </optgroup>
          <optgroup label="Elementary Faculty (From ELEM)">\n''')

for t in [x for x in teacher_loads_data if x['category'] == 'elem']:
    html_out.append(f'            <option value="{html.escape(t["name"])}">{html.escape(t["name"])} ({t["total_periods"]} periods - {t["total_hours"]} hrs)</option>\n')

html_out.append('''          </optgroup>
          <optgroup label="ISAL & Arabic Studies Faculty">\n''')

for t in [x for x in teacher_loads_data if x['category'] == 'isal']:
    html_out.append(f'            <option value="{html.escape(t["name"])}">{html.escape(t["name"])} ({t["total_periods"]} periods - {t["total_hours"]} hrs)</option>\n')

html_out.append('''          </optgroup>
        </select>
      </div>
    </div>

    <div class="status-summary">
      <span id="active-count-text">Displaying monitoring forms...</span>
      <span>Theme: <b>Light Mode</b> &bull; Orientation: <b>Portrait (A4)</b></span>
    </div>
  </div>

  <!-- ALL SHEETS WRAPPER -->
  <div class="sheet-wrapper">
''')

# 1. CLASSROOM DAILY MONITORING SHEETS
for d in days:
    for sec in sections:
        header_markup = build_header_html(
            "DAILY CLASSROOM INSTRUCTIONAL LOAD MONITORING FORM",
            f"Pure Face-to-Face Modality &bull; {sec['name']} &bull; S.Y. 2026 - 2027"
        )
        s_grid = elem_grid if sec['sheet'] == 'ELEM' else (hs_new_grid if sec['sheet'] == 'HS SCHED (NEW)' else hs_grid)
        d_idx = days.index(d)
        
        html_out.append(f'''
    <!-- DAILY SHEET: {d} - {sec['name']} -->
    <div class="page-sheet" 
         data-view="daily"
         data-day="{d}"
         data-dept="{sec['dept_code']}"
         data-grade="{sec['id']}">
      
      <div class="sheet-content">
        {header_markup}

        <!-- META BOX -->
        <div class="meta-box">
          <div class="meta-row"><span class="meta-lbl">Grade / Cohort:</span><span class="meta-val td-bold">{sec['name']} (Pure F2F)</span></div>
          <div class="meta-row"><span class="meta-lbl">Day of Week:</span><span class="meta-val td-bold">{d}</span></div>
          <div class="meta-row"><span class="meta-lbl">Room:</span><span class="meta-val"></span></div>
          <div class="meta-row"><span class="meta-lbl">Department:</span><span class="meta-val">{sec['dept_label']}</span></div>
          <div class="meta-row"><span class="meta-lbl">Date of Monitoring:</span><span class="meta-val"></span></div>
          <div class="meta-row"><span class="meta-lbl">Modality:</span><span class="meta-val">Pure Face-to-Face (F2F)</span></div>
        </div>

        <!-- TABLE -->
        <table class="sheet-table">
          <thead>
            <tr>
              <th class="th-num">#</th>
              <th class="th-time">Scheduled Time</th>
              <th class="th-mins">Mins</th>
              <th class="th-in">Actual In</th>
              <th class="th-out">Actual Out</th>
              <th class="th-subject">Subject / Learning Area</th>
              <th class="th-teacher">Assigned Subject Teacher</th>
              <th class="th-room">Room</th>
              <th class="th-status">Instruction Status</th>
              <th class="th-remarks">Remarks / Initial</th>
            </tr>
          </thead>
          <tbody>''')
        
        row_num = 1
        class_count = 0
        for r in range(sec['r_start'], sec['r_end'] + 1):
            t_slot = clean_time(s_grid.get((r, 2), ''))
            m_slot = s_grid.get((r, 3), '').strip()
            if not t_slot: continue
            
            c_val = s_grid.get((r, 4 + d_idx), '').strip()
            subj, tchr, kind = clean_parse(c_val)
            
            if kind == 'ROUTINE':
                html_out.append(f'''
            <tr class="row-routine">
              <td class="td-center">{row_num}</td>
              <td class="time-slot">{t_slot}</td>
              <td class="td-center">{m_slot}</td>
              <td colspan="7" style="text-align:center;padding:3px 8px;letter-spacing:0.5px;">
                <b>{html.escape(subj)}</b>
              </td>
            </tr>''')
            elif kind == 'CLASS':
                class_count += 1
                tchr_display = html.escape(tchr) if tchr != 'TBA' else '<span style="color:#64748b;font-weight:normal;">[ TBA ]</span>'
                html_out.append(f'''
            <tr>
              <td class="td-center td-bold">{row_num}</td>
              <td class="time-slot">{t_slot}</td>
              <td class="td-center">{m_slot}</td>
              <td class="td-center"></td>
              <td class="td-center"></td>
              <td><b>{html.escape(subj)}</b></td>
              <td><b>{tchr_display}</b></td>
              <td class="td-center"></td>
              <td>
                <div class="status-grid">
                  <div class="status-row"><span>[ ] In</span><span>[ ] Late</span></div>
                  <div class="status-row"><span>[ ] Abs</span><span>[ ] Sub</span></div>
                </div>
              </td>
              <td></td>
            </tr>''')
            else:
                html_out.append(f'''
            <tr class="row-blank">
              <td class="td-center" style="color:#cbd5e1;">{row_num}</td>
              <td class="time-slot">{t_slot}</td>
              <td class="td-center">{m_slot}</td>
              <td class="td-center"></td>
              <td class="td-center"></td>
              <td></td>
              <td></td>
              <td class="td-center"></td>
              <td></td>
              <td></td>
            </tr>''')
            row_num += 1

        blank_to_add = max(1, 14 - row_num)
        for b in range(blank_to_add):
            html_out.append(f'''
            <tr class="row-blank">
              <td class="td-center" style="color:#cbd5e1;">{row_num}</td>
              <td></td>
              <td></td>
              <td class="td-center"></td>
              <td class="td-center"></td>
              <td><span style="color:#cbd5e1;font-size:6.5pt;">(Extra / Remedial Period)</span></td>
              <td></td>
              <td></td>
              <td></td>
              <td></td>
            </tr>''')
            row_num += 1

        html_out.append('''
          </tbody>
        </table>
      </div>
    </div>
''')

# 2. WALKTHROUGH SHEETS
clusters = [
    {'id': 'elem', 'name': 'Elementary Cluster Walkthrough', 'sections': [s for s in sections if s['dept_code'] == 'elem']},
    {'id': 'jhs', 'name': 'Junior High School Walkthrough', 'sections': [s for s in sections if s['dept_code'] == 'jhs']},
    {'id': 'shs', 'name': 'Senior High School Walkthrough', 'sections': [s for s in sections if s['dept_code'] == 'shs']},
]

for d in days:
    for cl in clusters:
        header_markup = build_header_html(
            "DAILY CLASSROOM WALKTHROUGH & MONITORING SUMMARY",
            f"Department Cluster Walkthrough &bull; {cl['name']} &bull; {d}"
        )
        d_idx = days.index(d)
        
        html_out.append(f'''
    <!-- WALKTHROUGH SHEET: {d} - {cl['name']} -->
    <div class="page-sheet" 
         data-view="walkthrough"
         data-day="{d}"
         data-dept="{cl['id']}">
      
      <div class="sheet-content">
        {header_markup}

        <!-- META BOX -->
        <div class="meta-box">
          <div class="meta-row"><span class="meta-lbl">Cluster:</span><span class="meta-val td-bold">{cl['name']}</span></div>
          <div class="meta-row"><span class="meta-lbl">Day of Week:</span><span class="meta-val td-bold">{d}</span></div>
          <div class="meta-row"><span class="meta-lbl">Date:</span><span class="meta-val"></span></div>
          <div class="meta-row"><span class="meta-lbl">Monitored Sections:</span><span class="meta-val">{len(cl['sections'])} Pure F2F Classes</span></div>
          <div class="meta-row"><span class="meta-lbl">Monitoring Purpose:</span><span class="meta-val">Classroom Attendance & Instructional Verification</span></div>
          <div class="meta-row"><span class="meta-lbl">School Year:</span><span class="meta-val">2026 - 2027</span></div>
        </div>

        <!-- TABLE -->
        <table class="sheet-table">
          <thead>
            <tr>
              <th style="width:14%;">Grade / Section</th>
              <th class="th-time">Time Slot</th>
              <th style="width:24%;">Subject / Learning Area</th>
              <th style="width:20%;">Assigned Teacher</th>
              <th class="th-room">Room</th>
              <th style="width:14%;">Teacher Status</th>
              <th style="width:16%;">Observer Initials</th>
            </tr>
          </thead>
          <tbody>''')
        
        for sec in cl['sections']:
            s_grid = elem_grid if sec['sheet'] == 'ELEM' else (hs_new_grid if sec['sheet'] == 'HS SCHED (NEW)' else hs_grid)
            classes_in_sec = []
            for r in range(sec['r_start'], sec['r_end'] + 1):
                t_slot = clean_time(s_grid.get((r, 2), ''))
                c_val = s_grid.get((r, 4 + d_idx), '').strip()
                if not t_slot or not c_val: continue
                subj, tchr, kind = clean_parse(c_val)
                if kind == 'CLASS':
                    classes_in_sec.append((t_slot, subj, tchr))
                    
            for idx, (t_slot, subj, tchr) in enumerate(classes_in_sec):
                sec_td = f'<td rowspan="{len(classes_in_sec)}" class="td-center td-bold" style="background:#f8fafc;">{sec["name"]}</td>' if idx == 0 else ''
                tchr_display = html.escape(tchr) if tchr != 'TBA' else '<span style="color:#64748b;">[ TBA ]</span>'
                html_out.append(f'''
            <tr>
              {sec_td}
              <td class="time-slot">{t_slot}</td>
              <td><b>{html.escape(subj)}</b></td>
              <td><b>{tchr_display}</b></td>
              <td class="td-center"></td>
              <td>
                <div class="status-boxes">
                  <span>[ ] In</span>
                  <span>[ ] Late</span>
                  <span>[ ] Sub</span>
                </div>
              </td>
              <td></td>
            </tr>''')

        html_out.append('''
          </tbody>
        </table>
      </div>
    </div>
''')

# 3. WEEKLY TIMETABLE MATRIX SHEETS
for sec in sections:
    header_markup = build_header_html(
        "WEEKLY CLASS TIMETABLE & INSTRUCTIONAL LOAD MATRIX",
        f"Master Weekly Schedule &bull; {sec['name']} &bull; S.Y. 2026 - 2027"
    )
    s_grid = elem_grid if sec['sheet'] == 'ELEM' else (hs_new_grid if sec['sheet'] == 'HS SCHED (NEW)' else hs_grid)
    
    html_out.append(f'''
    <!-- WEEKLY SHEET: {sec['name']} -->
    <div class="page-sheet" 
         data-view="weekly"
         data-dept="{sec['dept_code']}"
         data-grade="{sec['id']}">
      
      <div class="sheet-content">
        {header_markup}

        <!-- META BOX -->
        <div class="meta-box">
          <div class="meta-row"><span class="meta-lbl">Grade / Cohort:</span><span class="meta-val td-bold">{sec['name']} (Pure F2F)</span></div>
          <div class="meta-row"><span class="meta-lbl">Weekly Cycle:</span><span class="meta-val td-bold">Sunday to Thursday (5 Days)</span></div>
          <div class="meta-row"><span class="meta-lbl">Room:</span><span class="meta-val"></span></div>
          <div class="meta-row"><span class="meta-lbl">Department:</span><span class="meta-val">{sec['dept_label']}</span></div>
          <div class="meta-row"><span class="meta-lbl">Modality:</span><span class="meta-val">Pure Face-to-Face</span></div>
          <div class="meta-row"><span class="meta-lbl">School Year:</span><span class="meta-val">2026 - 2027</span></div>
        </div>

        <!-- TABLE -->
        <table class="sheet-table" style="font-size:6.8pt;">
          <thead>
            <tr>
              <th class="th-time" style="width:13%;">Time Slot</th>
              <th class="th-mins" style="width:5%;">Mins</th>
              <th style="width:16.4%;">Sunday</th>
              <th style="width:16.4%;">Monday</th>
              <th style="width:16.4%;">Tuesday</th>
              <th style="width:16.4%;">Wednesday</th>
              <th style="width:16.4%;">Thursday</th>
            </tr>
          </thead>
          <tbody>''')
    
    for r in range(sec['r_start'], sec['r_end'] + 1):
        t_slot = clean_time(s_grid.get((r, 2), ''))
        m_slot = s_grid.get((r, 3), '').strip()
        if not t_slot: continue
        
        row_vals = [s_grid.get((r, 4 + d_idx), '').strip() for d_idx in range(5)]
        first_c = next((v for v in row_vals if v), '')
        is_routine = False
        if first_c:
            for kw in ['GENERAL ASSEMBLY', 'RECESS', 'LUNCH', 'SALAH', 'DEPARTURE', 'SHORT BREAK']:
                if kw in first_c.upper(): is_routine = True; break
                
        if is_routine:
            html_out.append(f'''
            <tr class="row-routine">
              <td class="time-slot">{t_slot}</td>
              <td class="td-center">{m_slot}</td>
              <td colspan="5" style="text-align:center;font-weight:700;letter-spacing:0.5px;">{html.escape(first_c)}</td>
            </tr>''')
        else:
            html_out.append(f'''
            <tr>
              <td class="time-slot">{t_slot}</td>
              <td class="td-center">{m_slot}</td>''')
            for c_val in row_vals:
                subj, tchr, kind = clean_parse(c_val)
                if not c_val:
                    html_out.append('<td></td>')
                elif any(kw in c_val.upper() for kw in ['RECESS', 'LUNCH', 'SALAH', 'ASSEMBLY', 'BREAK', 'DEPARTURE', 'HOMEROOM', 'ENTRANCE EXAM REVIEW', 'RESEARCH CONSULTATION']):
                    html_out.append(f'<td class="row-routine" style="font-size:6.8pt;">{html.escape(c_val)}</td>')
                else:
                    html_out.append(f'''
              <td style="font-size:7pt;padding:2px 3px;">
                <b>{html.escape(subj)}</b><br>
                <span style="color:#0f172a;font-size:6.5pt;">{html.escape(tchr)}</span>
              </td>''')
            html_out.append('</tr>')

    html_out.append('''
          </tbody>
        </table>
      </div>
    </div>
''')

# 4. NEW: INDIVIDUAL TEACHER'S TEACHING LOAD & WEEKLY SCHEDULE (FROM HS LOADS & ELEM)
for t in teacher_loads_data:
    t_name_esc = html.escape(t['name'])
    dept_esc = html.escape(t['department'])
    cat_esc = html.escape(t['category'])
    header_markup = build_header_html(
        "INDIVIDUAL TEACHER'S TEACHING LOAD & WEEKLY SCHEDULE",
        f"Master Weekly Faculty Program &bull; {t_name_esc} &bull; S.Y. 2026 - 2027"
    )

    html_out.append(f'''
    <!-- TEACHER LOAD SHEET: {t_name_esc} -->
    <div class="page-sheet sheet-teacher-load" 
         data-view="teacher_loads"
         data-teacher="{t_name_esc}"
         data-dept="{cat_esc}">
      
      <div class="sheet-content">
        {header_markup}

        <!-- TEACHER META BOX -->
        <div class="meta-box">
          <div class="meta-row"><span class="meta-lbl">Teacher's Name:</span><span class="meta-val td-bold">{t_name_esc}</span></div>
          <div class="meta-row"><span class="meta-lbl">Total Weekly Load:</span><span class="meta-val td-bold">{t['total_periods']} Classes ({t['total_mins']} mins / {t['total_hours']} hrs/wk)</span></div>
          <div class="meta-row"><span class="meta-lbl">Room Assignment:</span><span class="meta-val"></span></div>
          <div class="meta-row"><span class="meta-lbl">Department / Faculty:</span><span class="meta-val">{dept_esc}</span></div>
          <div class="meta-row"><span class="meta-lbl">Instructional Mode:</span><span class="meta-val">Pure Face-to-Face & Blended</span></div>
          <div class="meta-row"><span class="meta-lbl">School Year:</span><span class="meta-val">2026 - 2027</span></div>
        </div>

        <!-- SECTION 1: WEEKLY SCHEDULE MATRIX -->
        <div class="section-title"> Part I: Weekly Timetable Matrix (Sunday to Thursday)</div>
        <table class="schedule-table">
          <thead>
            <tr>
              <th class="time-slot-col">Scheduled Time</th>
              <th class="mins-col">Mins</th>
              <th class="day-col">Sunday</th>
              <th class="day-col">Monday</th>
              <th class="day-col">Tuesday</th>
              <th class="day-col">Wednesday</th>
              <th class="day-col">Thursday</th>
            </tr>
          </thead>
          <tbody>''')

    for row in t['grid_rows']:
        t_slot = row['time']
        m_val = row['mins']
        
        if row['is_routine']:
            r_lbl = row['routine_label']
            html_out.append(f'''
            <tr>
              <td class="time-slot-col">{t_slot}</td>
              <td class="mins-col">{m_val}</td>
              <td colspan="5" class="routine-row-cell">{html.escape(r_lbl)}</td>
            </tr>''')
        else:
            html_out.append(f'''
            <tr>
              <td class="time-slot-col">{t_slot}</td>
              <td class="mins-col">{m_val}</td>''')
            for day in days:
                d_entry = row['days'].get(day, {})
                if d_entry.get('is_class'):
                    txt = html.escape(d_entry['text'])
                    coh = html.escape(d_entry.get('cohort', ''))
                    html_out.append(f'''
              <td class="class-cell">
                <div class="cell-subj">{txt}</div>
                {f'<div class="cell-cohort">{coh}</div>' if coh and coh not in txt else ''}
              </td>''')
                else:
                    html_out.append('<td class="free-cell">—</td>')
            html_out.append('</tr>')

    html_out.append('''
          </tbody>
        </table>

        <!-- SECTION 2: LOAD SUMMARY BREAKDOWN -->
        <div class="section-title"> Part II: Teaching Load & Instructional Course Summary</div>
        <table class="summary-table">
          <thead>
            <tr>
              <th style="width:4%;">#</th>
              <th style="width:28%;">Subject / Course Description</th>
              <th style="width:20%;">Grade / Cohort / Section</th>
              <th style="width:20%;">Days Taught</th>
              <th style="width:16%;">Time Slot</th>
              <th style="width:12%;">Weekly Mins</th>
            </tr>
          </thead>
          <tbody>''')

    s_num = 1
    for sm in t['summary']:
        subj_lbl = html.escape(sm['subject'])
        coh_lbl = html.escape(sm['cohort'])
        days_str = ", ".join([d[:3] for d in sm['days']])
        time_str = html.escape(sm['time'])
        tot_m = sm['total_mins']
        sess = sm['sessions']

        html_out.append(f'''
            <tr>
              <td style="text-align:center;font-weight:700;">{s_num}</td>
              <td style="font-weight:700;color:#064e3b;">{subj_lbl}</td>
              <td style="text-align:center;">{coh_lbl}</td>
              <td style="text-align:center;"><b>{days_str}</b> ({sess}x/wk)</td>
              <td style="text-align:center;">{time_str}</td>
              <td style="text-align:center;font-weight:700;">{tot_m} mins</td>
            </tr>''')
        s_num += 1

    if not t['summary']:
        html_out.append('''
            <tr>
              <td colspan="6" style="text-align:center;color:#94a3b8;font-style:italic;">No regular teaching assignments recorded</td>
            </tr>''')

    html_out.append(f'''
          </tbody>
          <tfoot>
            <tr>
              <th colspan="4" style="text-align:right;padding-right:8px;">TOTAL INSTRUCTIONAL TEACHING LOAD:</th>
              <th colspan="2" style="text-align:center;color:#064e3b;">
                {t['total_periods']} Class Periods &bull; {t['total_mins']} Minutes/Week ({t['total_hours']} Hours/Week)
              </th>
            </tr>
          </tfoot>
        </table>

        <!-- SIGN-OFF BLOCK -->
        <div class="sign-block">
          <div class="sign-col">
            <div class="sign-title">Prepared & Conformed by:</div>
            <div class="sign-name">{t_name_esc}</div>
            <div class="sign-role">Faculty Member / Subject Teacher &bull; S.Y. 2026 - 2027</div>
          </div>
          <div class="sign-col">
            <div class="sign-title">Attested & Approved by:</div>
            <div class="sign-name">ACADEMIC COORDINATOR / PRINCIPAL</div>
            <div class="sign-role">Al Munawwara Islamic School</div>
          </div>
        </div>

      </div>
    </div>
''')

# 5. TEACHER ATTENDANCE & MONITORING LOG
for t_name in sorted(teacher_schedule.keys()):
    t_items = teacher_schedule[t_name]
    t_items.sort(key=lambda x: (days.index(x['day']), time_to_sort_key(x['time'])))
    t_dept = t_items[0]['dept_label'] if t_items else 'Integrated Faculty'
    header_markup = build_header_html(
        "TEACHER ATTENDANCE & INSTRUCTIONAL LOAD LOG",
        "Faculty Daily Attendance Sign-In & Verification Record"
    )
    
    html_out.append(f'''
    <!-- TEACHER ATTENDANCE SHEET: {t_name} -->
    <div class="page-sheet sheet-teacher-attendance" 
         data-view="teachers"
         data-teacher="{html.escape(t_name)}">
      
      <div class="sheet-content">
        {header_markup}

        <!-- META BOX -->
        <div class="meta-box">
          <div class="meta-row"><span class="meta-lbl">Teacher's Name:</span><span class="meta-val td-bold">{t_name}</span></div>
          <div class="meta-row"><span class="meta-lbl">Total Weekly Loads:</span><span class="meta-val td-bold">{len(t_items)} Pure F2F Classes</span></div>
          <div class="meta-row"><span class="meta-lbl">Department:</span><span class="meta-val">{t_dept}</span></div>
          <div class="meta-row"><span class="meta-lbl">Instructional Mode:</span><span class="meta-val">Pure Face-to-Face</span></div>
          <div class="meta-row"><span class="meta-lbl">School Year:</span><span class="meta-val">2026 - 2027</span></div>
          <div class="meta-row"><span class="meta-lbl">Faculty Status:</span><span class="meta-val">Active Faculty Member</span></div>
        </div>

        <!-- TABLE -->
        <table class="sheet-table">
          <thead>
            <tr>
              <th class="th-num">#</th>
              <th style="width:11%;">Day</th>
              <th class="th-time">Scheduled Time</th>
              <th class="th-in">Actual In</th>
              <th class="th-out">Actual Out</th>
              <th style="width:16%;">Grade / Cohort</th>
              <th class="th-subject">Subject / Learning Area</th>
              <th class="th-room">Room</th>
              <th class="th-status">Status</th>
              <th class="th-remarks">Teacher Signature</th>
            </tr>
          </thead>
          <tbody>''')
    
    t_num = 1
    for it in t_items:
        html_out.append(f'''
            <tr>
              <td class="td-center td-bold">{t_num}</td>
              <td class="td-center td-bold">{day_abbr_map.get(it['day'], it['day'].upper()[:3])}</td>
              <td class="time-slot">{it['time']}</td>
              <td class="td-center"></td>
              <td class="td-center"></td>
              <td><b>{it['section']}</b></td>
              <td><b>{it['subject']}</b></td>
              <td class="td-center"></td>
              <td>
                <div class="status-grid">
                  <div class="status-row"><span>[ ] In</span><span>[ ] Late</span></div>
                  <div class="status-row"><span>[ ] Abs</span><span>[ ] Sub</span></div>
                </div>
              </td>
              <td></td>
            </tr>''')
        t_num += 1

    blank_to_add = max(2, 14 - len(t_items))
    for b in range(blank_to_add):
        html_out.append(f'''
            <tr class="row-blank">
              <td class="td-center" style="color:#cbd5e1;">{t_num}</td>
              <td></td>
              <td></td>
              <td class="td-center"></td>
              <td class="td-center"></td>
              <td></td>
              <td><span style="color:#cbd5e1;font-size:6.5pt;">(Substitute / Extra Load)</span></td>
              <td></td>
              <td>
                <div class="status-grid">
                  <div class="status-row"><span>[ ] In</span><span>[ ] Late</span></div>
                  <div class="status-row"><span>[ ] Abs</span><span>[ ] Sub</span></div>
                </div>
              </td>
              <td></td>
            </tr>''')
        t_num += 1

    html_out.append('''
          </tbody>
        </table>
      </div>
    </div>
''')

# JAVASCRIPT LOGIC
html_out.append('''
  </div>

  <script>
    const AMIS_LOGO_SRC = "''' + amis_b64 + '''";
    const DEPED_LOGO_SRC = "''' + deped_b64 + '''";

    function populateLogos() {
      document.querySelectorAll('.amis-img').forEach(img => {
        img.src = AMIS_LOGO_SRC;
      });
      document.querySelectorAll('.deped-img').forEach(img => {
        img.src = DEPED_LOGO_SRC;
      });
    }

    let currentView = 'teacher_loads';

    function switchView(btn, viewName) {
      currentView = viewName;
      document.querySelectorAll('.view-tab').forEach(tab => {
        tab.classList.remove('active');
        tab.classList.remove('view-tab-highlight');
      });
      if (btn) {
        btn.classList.add('active');
        if (viewName === 'teacher_loads') btn.classList.add('view-tab-highlight');
      }

      const dayGroup = document.getElementById('day-filter-group');
      const deptGroup = document.getElementById('dept-filter-group');
      const gradeGroup = document.getElementById('grade-filter-group');
      const teacherGroup = document.getElementById('teacher-filter-group');

      if (viewName === 'daily') {
        dayGroup.style.display = 'flex';
        deptGroup.style.display = 'flex';
        gradeGroup.style.display = 'flex';
        teacherGroup.style.display = 'none';
      } else if (viewName === 'walkthrough') {
        dayGroup.style.display = 'flex';
        deptGroup.style.display = 'flex';
        gradeGroup.style.display = 'none';
        teacherGroup.style.display = 'none';
      } else if (viewName === 'weekly') {
        dayGroup.style.display = 'none';
        deptGroup.style.display = 'flex';
        gradeGroup.style.display = 'flex';
        teacherGroup.style.display = 'none';
      } else if (viewName === 'teacher_loads') {
        dayGroup.style.display = 'none';
        deptGroup.style.display = 'flex';
        gradeGroup.style.display = 'none';
        teacherGroup.style.display = 'flex';
      } else if (viewName === 'teachers') {
        dayGroup.style.display = 'none';
        deptGroup.style.display = 'none';
        gradeGroup.style.display = 'none';
        teacherGroup.style.display = 'flex';
      }

      applyFilters();
    }

    function applyFilters() {
      const selectedDay = document.getElementById('filter-day').value;
      const selectedDept = document.getElementById('filter-dept').value;
      const selectedGrade = document.getElementById('filter-grade').value;
      const selectedTeacher = document.getElementById('filter-teacher').value;

      let visibleCount = 0;
      const sheets = document.querySelectorAll('.page-sheet');

      sheets.forEach(sheet => {
        const view = sheet.getAttribute('data-view');
        const day = sheet.getAttribute('data-day');
        const dept = sheet.getAttribute('data-dept');
        const grade = sheet.getAttribute('data-grade');
        const teacher = sheet.getAttribute('data-teacher');

        let show = false;

        if (view === currentView) {
          show = true;

          if (currentView === 'daily' || currentView === 'walkthrough') {
            if (selectedDay !== 'all' && day !== selectedDay) {
              show = false;
            }
          }

          if (currentView === 'daily' || currentView === 'walkthrough' || currentView === 'weekly') {
            if (selectedDept !== 'all' && dept !== selectedDept) {
              show = false;
            }
          }

          if (currentView === 'daily' || currentView === 'weekly') {
            if (selectedGrade !== 'all' && grade !== selectedGrade) {
              show = false;
            }
          }

          if (currentView === 'teacher_loads') {
            if (selectedDept !== 'all' && dept !== selectedDept) {
              show = false;
            }
            if (selectedTeacher !== 'all' && teacher !== selectedTeacher) {
              show = false;
            }
          }

          if (currentView === 'teachers') {
            if (selectedTeacher !== 'all' && teacher !== selectedTeacher) {
              show = false;
            }
          }
        }

        if (show) {
          sheet.classList.remove('hidden-sheet');
          visibleCount++;
        } else {
          sheet.classList.add('hidden-sheet');
        }
      });

      const countText = document.getElementById('active-count-text');
      countText.innerHTML = 'Showing <b>' + visibleCount + '</b> active printable sheet' + (visibleCount === 1 ? '' : 's') + ' for current view.';
    }

    function showAllAndPrint() {
      document.querySelectorAll('.page-sheet[data-view="' + currentView + '"]').forEach(sh => {
        sh.classList.remove('hidden-sheet');
      });
      window.print();
      applyFilters();
    }

    window.addEventListener('DOMContentLoaded', () => {
      populateLogos();
      // Default to Teacher Loads view!
      switchView(document.querySelector('.view-tab-highlight'), 'teacher_loads');
    });
  </script>
</body>
</html>
''')

target_file = OUTPUT_FILE
with open(target_file, 'w', encoding='utf-8') as f:
    f.write(''.join(html_out))

print(f"Generated {target_file} ({os.path.getsize(target_file):,} bytes).")
